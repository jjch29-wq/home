# -*- coding: utf-8 -*-
"""독립 실행형 PAUT 검사설정표 생성기."""

from __future__ import annotations

import argparse
import json
import math
import os
import tempfile
import tkinter as tk
from dataclasses import asdict, dataclass, field, fields
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

import openpyxl
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


APP_TITLE = "PAUT 검사설정표 생성기"
APP_VERSION = "1.11.0"
CONFIG_PATH = Path(__file__).with_name("paut_scan_plan_settings.json")
PROBE_SOURCE = "https://ims.evidentscientific.com/en/products/u8330072/u8330072"


@dataclass
class ApertureConfig:
    probe_id: str = "P1"
    beamset_id: str = "A1"
    scan_side: str = "양측"
    active_elements: int = 32
    first_element: int = 1
    focal_depth_mm: float = 15.87
    index_offsets: str = "15, 30, 40"
    left_index_offsets: str = ""
    right_index_offsets: str = ""
    beam_exit_measurements: str = ""
    left_beam_exit_measurements: str = ""
    right_beam_exit_measurements: str = ""

    @staticmethod
    def _parse_indices(value: str) -> list[float]:
        values = []
        for token in value.replace("/", ",").split(","):
            token = token.strip()
            if token:
                values.append(float(token))
        return values

    def index_values(self) -> list[float]:
        return [
            value
            for source in (self.index_offsets, self.left_index_offsets, self.right_index_offsets)
            for value in self._parse_indices(source)
        ]

    def index_entries(self) -> list[tuple[str, float]]:
        entries = []
        for side, source in (
            (self.scan_side, self.index_offsets),
            ("좌측", self.left_index_offsets),
            ("우측", self.right_index_offsets),
        ):
            entries.extend((side, value) for value in self._parse_indices(source))
        return entries

    def index_text_for_side(self, side: str) -> str:
        side_values = {
            "양측": self._parse_indices(self.index_offsets),
            "좌측": self._parse_indices(self.left_index_offsets),
            "우측": self._parse_indices(self.right_index_offsets),
        }
        values = side_values.get(side) or side_values["양측"]
        return " / ".join(f"{value:g}" for value in values)

    def index_display_values(self) -> list[str]:
        displays = []
        for side, value in self.index_entries():
            if side == "좌측":
                displays.append(f"L {value:g} (-{value:g})")
            elif side == "우측":
                displays.append(f"R {value:g} (+{value:g})")
            else:
                displays.append(f"{value:g}")
        return displays

    def index_values_for_side(self, side_sign: int) -> list[float]:
        requested_side = "좌측" if side_sign < 0 else "우측"
        if self.scan_side != "양측" and self.scan_side != requested_side:
            return []
        common = self._parse_indices(self.index_offsets)
        specific = self._parse_indices(
            self.left_index_offsets if side_sign < 0 else self.right_index_offsets
        )
        return specific or common

    @staticmethod
    def _parse_measurements(value: str) -> list[tuple[float, float]]:
        measurements = []
        for token in value.replace("/", ",").split(","):
            token = token.strip()
            if not token:
                continue
            parts = token.replace("=", ":").split(":")
            if len(parts) != 2:
                raise ValueError(
                    "각도별 BeamTool 출사점→용접 중심 수평거리는 '각도:거리' 형식으로 입력하세요."
                )
            measurements.append((float(parts[0].strip()), float(parts[1].strip())))
        return measurements

    def exit_measurements(self) -> list[tuple[str, float, float]]:
        measurements = []
        for side, value in (
            (self.scan_side, self.beam_exit_measurements),
            ("좌측", self.left_beam_exit_measurements),
            ("우측", self.right_beam_exit_measurements),
        ):
            measurements.extend(
                (side, angle, distance)
                for angle, distance in self._parse_measurements(value)
            )
        return measurements


@dataclass
class ScanSettings:
    probe_model: str = "5L64-A2"
    frequency_mhz: float = 5.0
    total_elements: int = 64
    pitch_mm: float = 0.60
    elevation_mm: float = 10.0
    thickness_mm: float = 15.87
    bevel_angle_deg: float = 37.5
    bevel_tolerance_deg: float = 2.5
    weld_center_to_root_offset_mm: float = 0.0
    min_angle_deg: float = 46.0
    max_angle_deg: float = 61.0
    angle_step_deg: float = 1.0
    wave_type: str = "SW"
    law_config: str = "Sectorial"
    focus_type: str = "True depth"
    scan_direction: str = "양측 주사"
    apertures: list[ApertureConfig] = field(
        default_factory=lambda: [ApertureConfig()]
    )

    @classmethod
    def from_dict(cls, data: dict) -> "ScanSettings":
        data = dict(data)
        raw_apertures = data.pop("apertures", None)
        # 이전 단일 설정 JSON도 자동으로 다중 설정 구조로 변환한다.
        if raw_apertures is None:
            raw_apertures = [{
                "active_elements": data.pop("active_elements", 32),
                "first_element": data.pop("first_element", 1),
                "focal_depth_mm": data.pop("focal_depth_mm", data.get("thickness_mm", 15.87)),
                "index_offsets": data.pop("index_offsets", "15, 30, 40"),
            }]
        converted_apertures = []
        for item in raw_apertures:
            if isinstance(item, ApertureConfig):
                converted_apertures.append(item)
                continue
            item = dict(item)
            perpendicular_angle = 90.0 - float(data.get("bevel_angle_deg", 37.5))
            # v1.3~1.4의 출사점 거리 설정을 각도:거리 형식으로 이전한다.
            legacy_distance = item.pop("exit_to_weld_mm", None)
            legacy_min = item.pop("exit_to_weld_min_mm", None)
            legacy_center = item.pop("exit_to_weld_center_mm", legacy_distance)
            legacy_max = item.pop("exit_to_weld_max_mm", None)
            if "beam_exit_measurements" not in item:
                legacy_pairs = (
                    (perpendicular_angle - 6.0, legacy_min),
                    (perpendicular_angle, legacy_center),
                    (perpendicular_angle + 6.0, legacy_max),
                )
                item["beam_exit_measurements"] = ", ".join(
                    f"{angle:g}:{distance:g}" for angle, distance in legacy_pairs
                    if distance is not None
                )
            converted_apertures.append(ApertureConfig(**item))
        # 구버전 다중 설정은 모두 기본 A1으로 변환될 수 있으므로 중복 ID를 자동 보정한다.
        used_identifiers = set()
        for aperture in converted_apertures:
            identifier = (aperture.probe_id, aperture.beamset_id, aperture.scan_side)
            if identifier in used_identifiers:
                sequence = 1
                while (aperture.probe_id, f"A{sequence}", aperture.scan_side) in used_identifiers:
                    sequence += 1
                aperture.beamset_id = f"A{sequence}"
                identifier = (aperture.probe_id, aperture.beamset_id, aperture.scan_side)
            used_identifiers.add(identifier)
        data["apertures"] = converted_apertures
        return cls(**data)

    def validate(self) -> None:
        if not self.probe_model.strip():
            raise ValueError("탐촉자 모델을 입력하세요.")
        positive = {
            "주파수": self.frequency_mhz,
            "총 소자 수": self.total_elements,
            "Pitch": self.pitch_mm,
            "Elevation": self.elevation_mm,
            "두께": self.thickness_mm,
            "각도 간격": self.angle_step_deg,
        }
        for label, value in positive.items():
            if value <= 0:
                raise ValueError(f"{label}은(는) 0보다 커야 합니다.")
        if not 0 < self.min_angle_deg < 90 or not 0 < self.max_angle_deg < 90:
            raise ValueError("빔각은 0° 초과 90° 미만이어야 합니다.")
        if self.min_angle_deg > self.max_angle_deg:
            raise ValueError("최소 빔각은 최대 빔각보다 클 수 없습니다.")
        law_count = (self.max_angle_deg - self.min_angle_deg) / self.angle_step_deg
        if abs(law_count - round(law_count)) > 1e-8:
            raise ValueError("빔각 범위가 각도 간격으로 정확히 나누어지지 않습니다.")
        if not self.apertures:
            raise ValueError("활성소자 설정을 하나 이상 추가하세요.")
        identifiers = set()
        for number, aperture in enumerate(self.apertures, 1):
            if not aperture.probe_id.strip() or not aperture.beamset_id.strip():
                raise ValueError(f"설정 {number}: 탐촉자 ID와 Beamset ID를 입력하세요.")
            if aperture.scan_side not in {"좌측", "우측", "양측"}:
                raise ValueError(f"설정 {number}: 검사 측은 좌측, 우측 또는 양측이어야 합니다.")
            identifier = (aperture.probe_id.strip(), aperture.beamset_id.strip(), aperture.scan_side)
            if identifier in identifiers:
                raise ValueError(f"설정 {number}: 같은 탐촉자·Beamset·검사 측 ID가 중복되었습니다.")
            identifiers.add(identifier)
            if aperture.active_elements <= 0 or aperture.active_elements > self.total_elements:
                raise ValueError(f"설정 {number}: 활성 소자 수가 올바르지 않습니다.")
            if aperture.first_element < 1:
                raise ValueError(f"설정 {number}: First Element는 1 이상이어야 합니다.")
            if aperture.first_element + aperture.active_elements - 1 > self.total_elements:
                raise ValueError(f"설정 {number}: 활성 소자 범위가 총 소자 수를 초과합니다.")
            if aperture.focal_depth_mm <= 0:
                raise ValueError(f"설정 {number}: 초점 깊이는 0보다 커야 합니다.")
            for _side, angle, distance in aperture.exit_measurements():
                if not 0 < angle < 90:
                    raise ValueError(f"설정 {number}: 측정 각도는 0° 초과 90° 미만이어야 합니다.")
                if distance <= 0:
                    raise ValueError(
                        f"설정 {number}: BeamTool 출사점→용접 중심 수평거리는 0보다 커야 합니다."
                    )
                adjusted_distance = root_center_distance(
                    distance, _side, self.weld_center_to_root_offset_mm,
                )
                if adjusted_distance <= 0:
                    raise ValueError(
                        f"설정 {number}: 오프셋 보정 후 출사점→루트 중심 거리는 0보다 커야 합니다."
                    )
            if (abs(self.weld_center_to_root_offset_mm) > 1e-12
                    and aperture.scan_side == "양측"
                    and aperture.beam_exit_measurements.strip()):
                raise ValueError(
                    f"설정 {number}: 루트 중심 오프셋이 있으면 양측 공통 거리를 사용할 수 없습니다. "
                    "좌측/우측 각도:거리를 각각 입력하세요."
                )
            indices = aperture.index_values()
            if not indices or any(value < 0 for value in indices):
                raise ValueError(f"설정 {number}: Index는 0 이상의 숫자를 입력하세요.")


def reflected_bevel_depth(thickness_mm: float, bevel_angle_deg: float,
                          refracted_angle_deg: float, exit_to_weld_mm: float) -> float | None:
    """0.5 skip 반사 후 탐촉자 측 개선면을 통과하는 상면 기준 깊이."""
    beam_tangent = math.tan(math.radians(refracted_angle_deg))
    bevel_tangent = math.tan(math.radians(bevel_angle_deg))
    depth = (
        2 * thickness_mm * beam_tangent
        + thickness_mm * bevel_tangent
        - exit_to_weld_mm
    ) / (beam_tangent + bevel_tangent)
    if depth < 0 or depth > thickness_mm:
        return None
    return depth


def root_center_distance(exit_to_weld_center_mm: float, side: str,
                         weld_center_to_root_offset_mm: float) -> float:
    """BeamTool 용접 중심 거리를 루트 중심 거리로 보정한다.

    오프셋은 도면의 오른쪽을 양(+)으로 한다. 좌측 탐촉자는 오프셋을 더하고,
    우측 탐촉자는 빼야 동일한 루트 중심까지의 수평거리가 된다.
    """
    if side == "좌측":
        return exit_to_weld_center_mm + weld_center_to_root_offset_mm
    if side == "우측":
        return exit_to_weld_center_mm - weld_center_to_root_offset_mm
    return exit_to_weld_center_mm


def physical_focus_depth(thickness_mm: float, unfolded_depth_mm: float) -> tuple[float | None, str]:
    """펼친 True depth를 실제 상면 기준 깊이와 검사 구간으로 변환한다."""
    if unfolded_depth_mm <= thickness_mm:
        return unfolded_depth_mm, "직접입사 구간"
    if unfolded_depth_mm <= 2 * thickness_mm:
        return 2 * thickness_mm - unfolded_depth_mm, "0.5~1 skip 반사 구간"
    return None, "1 skip 초과"


def _font(size: int = 14, bold: bool = False):
    candidates = [
        r"C:\Windows\Fonts\malgun.ttf",
        r"C:\Windows\Fonts\malgunbd.ttf" if bold else "",
        r"C:\Windows\Fonts\arial.ttf",
    ]
    for candidate in candidates:
        if candidate and os.path.exists(candidate):
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def make_beam_diagram(settings: ScanSettings, output_path: str, width=1120, height=360) -> None:
    """입력값을 기반으로 개략적인 양측 주사/반사 경로 그림을 만든다."""
    settings.validate()
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = _font(14)
    small = _font(12)

    left, right = 70, width - 70
    # Keep the drawing vertically compact.  The previous large blank band above
    # the plate pushed the lower beam paths outside Excel's printable image area
    # for some printer/PDF drivers.
    top, bottom = 110, 270
    center = (left + right) / 2
    plate_height = bottom - top
    draw.rectangle((left, top, right, bottom), outline="#263238", width=2, fill="#FAFAFA")

    bevel_px = max(18, min(75, plate_height / math.tan(math.radians(settings.bevel_angle_deg))))
    draw.polygon(
        [(center - bevel_px, top), (center, bottom), (center + bevel_px, top)],
        fill="#ECEFF1", outline="#455A64",
    )

    scale = plate_height / settings.thickness_mm
    colors = ["#00A651", "#F4C20D", "#D81B60", "#1565C0", "#EF6C00", "#7C3AED"]
    probe_y = top - 48

    for side in (-1, 1):
        for config_index, aperture in enumerate(settings.apertures):
            for idx_pos, index_mm in enumerate(aperture.index_values_for_side(side)):
                probe_x = center + side * (bevel_px + index_mm * scale * 0.72)
                probe_y_offset = probe_y - config_index * 46
                probe_w, probe_h = 64, 34
                draw.rectangle(
                    (probe_x - probe_w / 2, probe_y_offset,
                     probe_x + probe_w / 2, probe_y_offset + probe_h),
                    fill="#A7A4FF", outline=colors[config_index % len(colors)], width=2,
                )
                label = f"{aperture.probe_id}-{aperture.beamset_id} I{index_mm:g}"
                draw.text((probe_x - 28, probe_y_offset - 17), label, fill="#212121", font=small)

                for angle_index, angle in enumerate((settings.min_angle_deg, settings.max_angle_deg)):
                    color = colors[(config_index * 2 + angle_index) % len(colors)]
                    horizontal = settings.thickness_mm * math.tan(math.radians(angle)) * scale
                    start = (probe_x, top)
                    hit_x = probe_x - side * horizontal
                    hit = (hit_x, bottom)
                    draw.line((start, hit), fill=color, width=3)
                    reflected_x = hit_x - side * horizontal
                    draw.line((hit, (reflected_x, top)), fill=color, width=2)

    draw.line((left, top, left - 22, top), fill="#2E7D32", width=2)
    draw.line((left, bottom, left - 22, bottom), fill="#2E7D32", width=2)
    draw.line((left - 15, top, left - 15, bottom), fill="#2E7D32", width=2)
    draw.text((10, (top + bottom) / 2 - 8), f"t={settings.thickness_mm:g} mm", fill="#2E7D32", font=font)
    title = (
        f"{settings.probe_model} | {settings.wave_type} | "
        f"{settings.min_angle_deg:g}°~{settings.max_angle_deg:g}° | {settings.scan_direction}"
    )
    draw.text((width / 2, 20), title, fill="#111827", font=_font(18, True), anchor="ma")
    draw.text(
        (width / 2, height - 24),
        "개략도: 실제 적용 전 웨지 출사점, 용접부 형상 및 교정시험편으로 커버리지를 확인할 것",
        fill="#B91C1C", font=small, anchor="ma",
    )
    legend = " | ".join(
        f"{item.probe_id}-{item.beamset_id}: {item.active_elements}el, "
        f"First {item.first_element}, TD {item.focal_depth_mm:g}"
        for item in settings.apertures
    )
    draw.text((width / 2, 48), legend, fill="#334155", font=small, anchor="ma")
    image.save(output_path, format="PNG")


def create_workbook(settings: ScanSettings, output_path: str) -> str:
    settings.validate()
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PAUT 설정표"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A8"
    ws.print_area = "A1:K42"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = False

    navy = "1F4E78"
    blue = "D9EAF7"
    pale = "EAF3F8"
    input_fill = "FFF2CC"
    white = "FFFFFF"
    thin = Side(style="thin", color="8091A5")
    medium = Side(style="medium", color="1F2937")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    outer = Border(left=medium, right=medium, top=medium, bottom=medium)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = [18, 15, 16, 18, 22, 20, 18, 18, 18, 25, 34]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.merge_cells("A1:K2")
    ws["A1"] = "PAUT 검사 설정 및 빔 방향 검토표"
    ws["A1"].font = Font(name="맑은 고딕", size=18, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 30

    def section(row: int, text: str):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
        cell = ws.cell(row, 1, text)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(name="맑은 고딕", size=12, bold=True, color="17365D")
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in range(1, 12):
            ws.cell(row, col).border = outer

    section(4, "1. 기본 입력조건")
    basic = [
        ("부재 두께 (mm)", settings.thickness_mm, "공칭 개선각 (°)", settings.bevel_angle_deg,
         "개선각 공차 (±°)", settings.bevel_tolerance_deg, "Angle step (°)", settings.angle_step_deg),
        ("최소 빔각 (°)", settings.min_angle_deg, "최대 빔각 (°)", settings.max_angle_deg,
         "검사 방향", settings.scan_direction, "Wave type", settings.wave_type),
        ("용접 중심→루트 중심 오프셋 (mm)", settings.weld_center_to_root_offset_mm,
         "오프셋 부호", "우측 + / 좌측 -", "거리 출처", "BeamTool: Beam Exit to Weld",
         "보정 기준", "좌측 +오프셋 / 우측 -오프셋"),
    ]
    # Keep legacy correction metadata in the settings model, but omit its
    # explanatory row from the field setup workbook.
    for r, row_data in enumerate(basic[:2], start=5):
        for label_col, value_col, label, value in (
            (1, 2, row_data[0], row_data[1]), (4, 5, row_data[2], row_data[3]),
            (7, 8, row_data[4], row_data[5]), (10, 11, row_data[6], row_data[7]),
        ):
            ws.cell(r, label_col, label)
            ws.cell(r, value_col, value)
            ws.cell(r, label_col).fill = PatternFill("solid", fgColor=pale)
            ws.cell(r, value_col).fill = PatternFill("solid", fgColor=input_fill)
            for col in (label_col, value_col):
                ws.cell(r, col).border = border
                ws.cell(r, col).alignment = center
                ws.cell(r, col).font = Font(name="맑은 고딕", size=10, bold=(col == label_col))

    ws.row_dimensions[7].height = 0
    ws.row_dimensions[7].hidden = True

    max_index_count = max(len(aperture.index_display_values()) for aperture in settings.apertures)
    headers = ["Probe", "Wave type", "Law config.", "Focus type", "Active aperture",
               "Sweep angle range", "Law count", "True depth\n(unfolded)"]
    for col, value in enumerate(headers, 1):
        ws.cell(8, col, value)
    if max_index_count == 1:
        ws.merge_cells("I8:K8")
        ws["I8"] = "Index"
    elif max_index_count == 2:
        ws["I8"] = "Index 1"
        ws.merge_cells("J8:K8")
        ws["J8"] = "Index 2"
    else:
        ws["I8"], ws["J8"], ws["K8"] = "Index 1", "Index 2", "Index 3+"
    for col in range(1, 12):
        cell = ws.cell(8, col)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="맑은 고딕", size=9, bold=True, color=white)
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[8].height = 31
    for offset, aperture in enumerate(settings.apertures):
        row = 9 + offset
        indices = aperture.index_display_values()
        values = [f"{settings.probe_model}\n{aperture.probe_id}-{aperture.beamset_id} ({aperture.scan_side})",
                  settings.wave_type,
                  settings.law_config, settings.focus_type,
                  None, f"{settings.min_angle_deg:g}°~{settings.max_angle_deg:g}°", None,
                  aperture.focal_depth_mm]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        if max_index_count == 1:
            ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=11)
            ws.cell(row, 9, indices[0])
        elif max_index_count == 2:
            ws.cell(row, 9, indices[0] if indices else None)
            ws.merge_cells(start_row=row, start_column=10, end_row=row, end_column=11)
            ws.cell(row, 10, indices[1] if len(indices) > 1 else None)
        else:
            ws.cell(row, 9, indices[0] if indices else None)
            ws.cell(row, 10, indices[1] if len(indices) > 1 else None)
            ws.cell(row, 11, " / ".join(indices[2:]))
        for col in range(1, 12):
            ws.cell(row, col).alignment = center
            ws.cell(row, col).border = border
            ws.cell(row, col).font = Font(name="맑은 고딕", size=10)
        ws.cell(row, 7, "=INT((E6-B6)/K5)+1")
        ws.cell(row, 5).number_format = '0.00" mm"'
        ws.cell(row, 7).number_format = '0" laws"'
        ws.row_dimensions[row].height = 32

    probe_section_row = max(12, 10 + len(settings.apertures))
    section(probe_section_row, "2. 탐촉자 및 Focal Law 정보")
    probe_header_row = probe_section_row + 1
    probe_data_start = probe_header_row + 1
    probe_headers = ["탐촉자 형식", "주파수 (MHz)", "진동자 수", "Total aperture (mm)",
                     "Pitch (mm)", "Active elements", "First Element", "Focusing Type",
                     "True depth (unfolded)", "Index", "비고"]
    for col, value in enumerate(probe_headers, 1):
        cell = ws.cell(probe_header_row, col, value)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="맑은 고딕", size=9, bold=True, color=white)
        cell.alignment = center
        cell.border = border
    for offset, aperture in enumerate(settings.apertures):
        row = probe_data_start + offset
        probe_values = [f"{settings.probe_model}\n{aperture.probe_id}-{aperture.beamset_id} ({aperture.scan_side})",
                        settings.frequency_mhz, settings.total_elements, None,
                        settings.pitch_mm, aperture.active_elements, aperture.first_element,
                        settings.focus_type, aperture.focal_depth_mm,
                        " | ".join(aperture.index_display_values()), None]
        for col, value in enumerate(probe_values, 1):
            ws.cell(row, col, value)
            ws.cell(row, col).alignment = center
            ws.cell(row, col).border = border
            ws.cell(row, col).font = Font(name="맑은 고딕", size=10)
        ws.cell(row, 4, f"=C{row}*E{row}")
        ws.cell(row, 11, f'=F{row}&" elements × "&TEXT(E{row},"0.00")&" mm = "&TEXT(F{row}*E{row},"0.00")&" mm"')
        ws.cell(row, 4).number_format = '0.00" mm"'
        ws.cell(row, 1).comment = Comment(f"Manufacturer source: {PROBE_SOURCE}", "User")
        ws.row_dimensions[row].height = 34

    # Summary active-aperture formulas point to the corresponding Focal Law row.
    for offset in range(len(settings.apertures)):
        ws.cell(9 + offset, 5, f"=F{probe_data_start + offset}*E{probe_data_start + offset}")

    scan_section_row = probe_data_start + len(settings.apertures) + 1
    section(scan_section_row, "3. PAUT Scan Plan")
    diagram_start_row = scan_section_row + 1
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as temp_file:
        diagram_path = temp_file.name
    try:
        make_beam_diagram(settings, diagram_path)
        diagram = XLImage(diagram_path)
        diagram.width = 1000
        # Preserve the source aspect ratio so the whole drawing is visible and
        # the lower beam paths are not cropped or vertically distorted.
        diagram.height = round(diagram.width * 360 / 1120)
        ws.add_image(diagram, f"A{diagram_start_row}")
        # Excel row heights are points while image dimensions are pixels.
        # Reserve the converted image height plus a small print-driver margin.
        diagram_row_count = math.ceil((diagram.height * 0.75) / 16) + 3
        diagram_end_row = diagram_start_row + diagram_row_count
        for row in range(diagram_start_row, diagram_end_row):
            ws.row_dimensions[row].height = 16

        # The workbook is a field setup sheet, so stop after the scan-plan
        # diagram.  Beam-path, true-depth, and bevel-intersection calculations
        # remain available in the application preview and validation logic but
        # are intentionally omitted from the exported deliverable.
        # Floating images do not extend Excel's print area when users move them.
        # Keep several blank rows inside the print area so the diagram can be
        # moved downward without its lower portion being clipped in print/PDF.
        diagram_move_margin_rows = 10
        print_end_row = diagram_end_row + diagram_move_margin_rows
        for row in range(diagram_end_row, print_end_row + 1):
            ws.row_dimensions[row].height = 16
        ws.print_area = f"A1:K{print_end_row}"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.save(output)
        return str(output)

        calc_section_row = diagram_end_row + 1
        section(calc_section_row, "4. 빔 경로 계산 및 검토사항")
        calc_header_row = calc_section_row + 1
        calc_headers = ["항목", "최소각", "최대각"]
        for col, value in enumerate(calc_headers, 1):
            ws.cell(calc_header_row, col, value)
        ws.merge_cells(start_row=calc_header_row, start_column=4, end_row=calc_header_row, end_column=11)
        ws.cell(calc_header_row, 4, "검토")
        for col in range(1, 12):
            ws.cell(calc_header_row, col).fill = PatternFill("solid", fgColor=navy)
            ws.cell(calc_header_row, col).font = Font(name="맑은 고딕", size=9, bold=True, color=white)
            ws.cell(calc_header_row, col).alignment = center
            ws.cell(calc_header_row, col).border = border
        leg_row, skip_row = calc_header_row + 1, calc_header_row + 2
        ws.cell(leg_row, 1, "1-leg 표면거리")
        ws.cell(leg_row, 2, "=B5*TAN(RADIANS(B6))")
        ws.cell(leg_row, 3, "=B5*TAN(RADIANS(E6))")
        ws.cell(leg_row, 4, "내면(ID) 도달 시 표면 투영거리")
        ws.merge_cells(start_row=leg_row, start_column=4, end_row=leg_row, end_column=11)
        ws.cell(skip_row, 1, "Full-skip 거리")
        ws.cell(skip_row, 2, f"=2*B{leg_row}")
        ws.cell(skip_row, 3, f"=2*C{leg_row}")
        ws.cell(skip_row, 4, "1회 반사 후 상면 도달거리")
        ws.merge_cells(start_row=skip_row, start_column=4, end_row=skip_row, end_column=11)

        focus_title_row = skip_row + 1
        ws.merge_cells(start_row=focus_title_row, start_column=1,
                       end_row=focus_title_row, end_column=11)
        ws.cell(focus_title_row, 1, "True depth (unfolded) 기준 실제 초점 위치")
        focus_header_row = focus_title_row + 1
        focus_headers = ["설정", "True depth (unfolded)", "실제 초점 깊이"]
        for col, value in enumerate(focus_headers, 1):
            ws.cell(focus_header_row, col, value)
        ws.merge_cells(start_row=focus_header_row, start_column=4,
                       end_row=focus_header_row, end_column=11)
        ws.cell(focus_header_row, 4, "검사 구간")
        for col in range(1, 12):
            title_cell = ws.cell(focus_title_row, col)
            title_cell.fill = PatternFill("solid", fgColor=blue)
            title_cell.font = Font(name="맑은 고딕", size=10, bold=True, color="17365D")
            title_cell.alignment = center
            title_cell.border = border
            header_cell = ws.cell(focus_header_row, col)
            header_cell.fill = PatternFill("solid", fgColor=navy)
            header_cell.font = Font(name="맑은 고딕", size=9, bold=True, color=white)
            header_cell.alignment = center
            header_cell.border = border
        focus_end_row = focus_header_row
        for offset, aperture in enumerate(settings.apertures):
            row = focus_header_row + 1 + offset
            focus_end_row = row
            ws.cell(row, 1, f"{aperture.probe_id}-{aperture.beamset_id}")
            ws.cell(row, 2, aperture.focal_depth_mm)
            ws.cell(
                row, 3,
                f'=IF(B{row}<=$B$5,B{row},IF(B{row}<=2*$B$5,2*$B$5-B{row},NA()))',
            )
            ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=11)
            ws.cell(
                row, 4,
                f'=IF(B{row}<=$B$5,"직접입사 구간",'
                f'IF(B{row}<=2*$B$5,"0.5~1 skip 반사 구간","1 skip 초과"))',
            )
            for col in range(1, 12):
                cell = ws.cell(row, col)
                cell.border = border
                cell.alignment = center
                cell.font = Font(name="맑은 고딕", size=9)
            ws.cell(row, 2).number_format = '0.00" mm"'
            ws.cell(row, 3).number_format = '0.00" mm"'

        measurements = [
            (number, side, angle, distance)
            for number, aperture in enumerate(settings.apertures, 1)
            for side, angle, distance in aperture.exit_measurements()
        ]
        measurement_rows = []
        if measurements:
            measurement_title_row = focus_end_row + 1
            ws.merge_cells(start_row=measurement_title_row, start_column=1,
                           end_row=measurement_title_row, end_column=11)
            ws.cell(measurement_title_row, 1, "0.5 skip 반사 후 개선면 통과 깊이")
            measurement_header_row = measurement_title_row + 1
            measurement_headers = ["설정", "검사 측", "Index", "실제 위치", "실제 굴절각",
                                   "BeamTool 출사점→용접 중심\n수평거리", "루트 오프셋 보정",
                                   "출사점→루트 중심\n보정거리", "통과 깊이"]
            for col, value in enumerate(measurement_headers, 1):
                ws.cell(measurement_header_row, col, value)
            ws.merge_cells(start_row=measurement_header_row, start_column=10,
                           end_row=measurement_header_row, end_column=11)
            ws.cell(measurement_header_row, 10, "개선면 수직 기준 ±6° 검토")
            for col in range(1, 12):
                cell = ws.cell(measurement_title_row, col)
                cell.fill = PatternFill("solid", fgColor=blue)
                cell.font = Font(name="맑은 고딕", size=10, bold=True, color="17365D")
                cell.alignment = center
                cell.border = border
                cell = ws.cell(measurement_header_row, col)
                cell.fill = PatternFill("solid", fgColor=navy)
                cell.font = Font(name="맑은 고딕", size=9, bold=True, color=white)
                cell.alignment = center
                cell.border = border

            for offset, (number, side, angle, distance) in enumerate(measurements):
                row = measurement_header_row + 1 + offset
                measurement_rows.append(row)
                aperture = settings.apertures[number - 1]
                index_text = aperture.index_text_for_side(side)
                ws.cell(row, 1, f"{aperture.probe_id}-{aperture.beamset_id}")
                ws.cell(row, 2, side)
                ws.cell(row, 3, index_text)
                if index_text and "/" not in index_text:
                    index_value = float(index_text)
                    if side == "양측":
                        ws.cell(row, 4, f"±{index_value:g}")
                    else:
                        position = -index_value if side == "좌측" else index_value
                        ws.cell(row, 4, position)
                else:
                    ws.cell(row, 4, index_text)
                ws.cell(row, 5, angle)
                ws.cell(row, 6, distance)
                signed_offset = (
                    settings.weld_center_to_root_offset_mm
                    if side == "좌측"
                    else -settings.weld_center_to_root_offset_mm if side == "우측" else 0.0
                )
                ws.cell(row, 7, signed_offset)
                ws.cell(row, 8, f"=F{row}+G{row}")
                ws.cell(
                    row, 9,
                    f"=(2*$B$5*TAN(RADIANS(E{row}))+$B$5*TAN(RADIANS($E$5))-H{row})/"
                    f"(TAN(RADIANS(E{row}))+TAN(RADIANS($E$5)))",
                )
                ws.merge_cells(start_row=row, start_column=10, end_row=row, end_column=11)
                ws.cell(
                    row, 10,
                    f'=IF(OR(I{row}<0,I{row}>$B$5),"교차 없음",'
                    f'IF(ABS(E{row}-(90-$E$5))<=6,"±6° 충족",'
                    f'"±6° 미충족 (편차 "&TEXT(ABS(E{row}-(90-$E$5)),"0.00")&"°)"))',
                )
                for col in range(1, 12):
                    cell = ws.cell(row, col)
                    cell.border = border
                    cell.alignment = center
                    cell.font = Font(name="맑은 고딕", size=9)
                ws.cell(row, 3).number_format = '0.00" mm"'
                ws.cell(row, 4).number_format = '+0.00;-0.00;0.00" mm"'
                ws.cell(row, 5).number_format = '0.00"°"'
                ws.cell(row, 6).number_format = '0.00" mm"'
                ws.cell(row, 7).number_format = '0.00" mm"'
                ws.cell(row, 8).number_format = '0.00" mm"'
                ws.cell(row, 9).number_format = '0.00" mm"'
            note_row = measurement_header_row + 1 + len(measurements)
        else:
            note_row = focus_end_row + 1

        ws.cell(note_row, 1, "주의")
        ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=11)
        ws.cell(note_row, 2, "실제 적용 전 웨지 출사점, 용접부 폭·캡·HAZ 및 교정시험편으로 커버리지를 확인할 것")
        for row in (*range(leg_row, skip_row + 1), note_row):
            for col in range(1, 12):
                cell = ws.cell(row, col)
                cell.border = border
                cell.alignment = center
                cell.font = Font(name="맑은 고딕", size=9, color="B91C1C" if row == note_row else "000000")
        for cell in (ws.cell(leg_row, 2), ws.cell(leg_row, 3), ws.cell(skip_row, 2), ws.cell(skip_row, 3)):
            cell.number_format = '0.00" mm"'

        ws.print_area = f"A1:K{note_row}"

        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.save(output)
    finally:
        try:
            os.unlink(diagram_path)
        except OSError:
            pass
    return str(output)


class PautScanPlanApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(f"{APP_TITLE} v{APP_VERSION}")
        self.geometry("1080x760")
        self.minsize(980, 700)
        self.vars: dict[str, tk.StringVar] = {}
        self.aperture_vars: dict[str, tk.StringVar] = {}
        self.apertures: list[ApertureConfig] = []
        self._build_ui()
        self._set_values(self._load_default_settings())
        self.after(150, self.refresh_preview)

    def _build_ui(self):
        style = ttk.Style(self)
        style.configure("Title.TLabel", font=("맑은 고딕", 16, "bold"))
        style.configure("Section.TLabelframe.Label", font=("맑은 고딕", 10, "bold"))

        header = ttk.Frame(self, padding=(16, 12))
        header.pack(fill="x")
        ttk.Label(header, text="PAUT 검사설정표 생성기", style="Title.TLabel").pack(side="left")
        ttk.Label(header, text="입력 → 검증 → 빔 경로 확인 → Excel 생성").pack(side="left", padx=18)

        body = ttk.Panedwindow(self, orient="horizontal")
        body.pack(fill="both", expand=True, padx=12, pady=(0, 10))
        form_outer = ttk.Frame(body, padding=4)
        preview_outer = ttk.Frame(body, padding=4)
        body.add(form_outer, weight=2)
        body.add(preview_outer, weight=3)

        canvas = tk.Canvas(form_outer, highlightthickness=0)
        scrollbar = ttk.Scrollbar(form_outer, orient="vertical", command=canvas.yview)
        form = ttk.Frame(canvas)
        form.bind("<Configure>", lambda _e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=form, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        groups = [
            ("기본 조건", [
                ("thickness_mm", "부재 두께 (mm)"), ("bevel_angle_deg", "공칭 개선각 (°)"),
                ("bevel_tolerance_deg", "개선각 공차 (±°)"), ("scan_direction", "검사 방향"),
                ("weld_center_to_root_offset_mm", "용접 중심→루트 중심 오프셋 (mm)"),
            ]),
            ("탐촉자", [
                ("probe_model", "탐촉자 모델"), ("frequency_mhz", "주파수 (MHz)"),
                ("total_elements", "총 소자 수"), ("pitch_mm", "Pitch (mm)"),
                ("elevation_mm", "Elevation (mm)"),
            ]),
            ("Focal Law", [
                ("min_angle_deg", "최소 빔각 (°)"), ("max_angle_deg", "최대 빔각 (°)"),
                ("angle_step_deg", "Angle step (°)"), ("wave_type", "Wave type"),
                ("law_config", "Law config."), ("focus_type", "Focus type"),
            ]),
        ]
        combo_values = {
            "scan_direction": ("양측 주사", "단측 주사"),
            "wave_type": ("SW", "LW"),
            "law_config": ("Sectorial", "Linear"),
            "focus_type": ("True depth", "Sound path", "Projection"),
        }
        for title, items in groups:
            box = ttk.LabelFrame(form, text=title, style="Section.TLabelframe", padding=10)
            box.pack(fill="x", pady=5)
            visible_items = [
                item for item in items
                if item[0] != "weld_center_to_root_offset_mm"
            ]
            for row, (key, label) in enumerate(visible_items):
                ttk.Label(box, text=label, width=20).grid(row=row, column=0, sticky="w", pady=3)
                var = tk.StringVar()
                self.vars[key] = var
                if key in combo_values:
                    widget = ttk.Combobox(box, textvariable=var, values=combo_values[key], state="readonly", width=22)
                else:
                    widget = ttk.Entry(box, textvariable=var, width=25)
                widget.grid(row=row, column=1, sticky="ew", pady=3)
                widget.bind("<FocusOut>", lambda _e: self.refresh_preview())
                widget.bind("<Return>", lambda _e: self.refresh_preview())
            box.columnconfigure(1, weight=1)

        aperture_box = ttk.LabelFrame(form, text="활성소자 설정", style="Section.TLabelframe", padding=10)
        aperture_box.pack(fill="x", pady=5)
        self.vars["weld_center_to_root_offset_mm"] = tk.StringVar(value="0")
        columns = ("probe", "beamset", "side", "active", "first", "focal", "indices", "exit")
        self.aperture_tree = ttk.Treeview(aperture_box, columns=columns, show="headings", height=4)
        headings = (("probe", "탐촉자 ID"), ("beamset", "Beamset"),
                    ("side", "검사 측"), ("active", "활성 소자"), ("first", "First"),
                    ("focal", "초점 깊이"), ("indices", "Index"),
                    ("exit", "좌/우 각도:BeamTool 용접 중심거리"))
        for key, label in headings:
            self.aperture_tree.heading(key, text=label)
            self.aperture_tree.column(key, width=72 if key not in {"indices", "exit"} else 145, anchor="center")
        self.aperture_tree.grid(row=0, column=0, columnspan=4, sticky="ew", pady=(0, 7))
        aperture_scroll = ttk.Scrollbar(aperture_box, orient="horizontal", command=self.aperture_tree.xview)
        aperture_scroll.grid(row=1, column=0, columnspan=4, sticky="ew")
        self.aperture_tree.configure(xscrollcommand=aperture_scroll.set)
        self.aperture_tree.configure(displaycolumns=columns[:-1])
        self.aperture_tree.bind("<<TreeviewSelect>>", self._select_aperture)

        editors = (("probe_id", "탐촉자 ID"), ("beamset_id", "Beamset ID"),
                   ("scan_side", "검사 측 (좌측/우측/양측)"),
                   ("active_elements", "활성 소자"), ("first_element", "First Element"),
                   ("focal_depth_mm", "초점 깊이"), ("index_offsets", "공통 Index"),
                   ("left_index_offsets", "좌측 Index"),
                   ("right_index_offsets", "우측 Index"),
                   ("beam_exit_measurements", "공통 각도:BeamTool 용접 중심거리"),
                   ("left_beam_exit_measurements", "좌측 각도:BeamTool 용접 중심거리"),
                   ("right_beam_exit_measurements", "우측 각도:BeamTool 용접 중심거리"))
        hidden_editor_keys = {
            "beam_exit_measurements",
            "left_beam_exit_measurements",
            "right_beam_exit_measurements",
        }
        visible_editors = [
            item for item in editors if item[0] not in hidden_editor_keys
        ]
        for row, (key, label) in enumerate(visible_editors, 2):
            ttk.Label(aperture_box, text=label).grid(row=row, column=0, sticky="w", pady=2)
            var = tk.StringVar()
            self.aperture_vars[key] = var
            if key == "scan_side":
                widget = ttk.Combobox(aperture_box, textvariable=var,
                                      values=("좌측", "우측", "양측"), state="readonly", width=20)
            else:
                widget = ttk.Entry(aperture_box, textvariable=var, width=22)
            widget.grid(row=row, column=1, columnspan=3, sticky="ew", pady=2)
        for key in hidden_editor_keys:
            self.aperture_vars[key] = tk.StringVar(value="")
        button_row = len(visible_editors) + 2
        ttk.Button(aperture_box, text="추가", command=self.add_aperture).grid(row=button_row, column=0, pady=7)
        ttk.Button(aperture_box, text="선택 수정", command=self.update_aperture).grid(row=button_row, column=1, pady=7)
        ttk.Button(aperture_box, text="선택 삭제", command=self.delete_aperture).grid(row=button_row, column=2, pady=7)
        ttk.Button(aperture_box, text="입력 지우기", command=self._clear_aperture_editor).grid(row=button_row, column=3, pady=7)
        aperture_box.columnconfigure(3, weight=1)

        action = ttk.Frame(form, padding=(0, 10))
        action.pack(fill="x")
        ttk.Button(action, text="설정 불러오기", command=self.load_settings).pack(side="left", padx=3)
        ttk.Button(action, text="설정 저장", command=self.save_settings).pack(side="left", padx=3)
        ttk.Button(action, text="기본값", command=self.reset_defaults).pack(side="left", padx=3)

        preview_box = ttk.LabelFrame(preview_outer, text="빔 경로 미리보기", padding=8)
        preview_box.pack(fill="both", expand=True)
        self.preview = tk.Canvas(preview_box, bg="white", highlightthickness=1, highlightbackground="#CBD5E1")
        self.preview.pack(fill="both", expand=True)
        self.preview.bind("<Configure>", lambda _e: self.refresh_preview())

        calc_box = ttk.LabelFrame(preview_outer, text="빔 경로 계산", padding=8)
        calc_box.pack_forget()
        ttk.Label(calc_box, text="항목", anchor="center").grid(row=0, column=0, sticky="ew")
        ttk.Label(calc_box, text="최소각", anchor="center").grid(row=0, column=1, sticky="ew")
        ttk.Label(calc_box, text="최대각", anchor="center").grid(row=0, column=2, sticky="ew")
        self.path_calc_vars = {
            "min_angle": tk.StringVar(value="-"),
            "max_angle": tk.StringVar(value="-"),
            "min_leg": tk.StringVar(value="-"),
            "max_leg": tk.StringVar(value="-"),
            "min_skip": tk.StringVar(value="-"),
            "max_skip": tk.StringVar(value="-"),
        }
        ttk.Label(calc_box, text="굴절각").grid(row=1, column=0, sticky="w", pady=2)
        ttk.Label(calc_box, textvariable=self.path_calc_vars["min_angle"], anchor="center").grid(row=1, column=1, sticky="ew")
        ttk.Label(calc_box, textvariable=self.path_calc_vars["max_angle"], anchor="center").grid(row=1, column=2, sticky="ew")
        ttk.Label(calc_box, text="1-leg 표면거리").grid(row=2, column=0, sticky="w", pady=2)
        ttk.Label(calc_box, textvariable=self.path_calc_vars["min_leg"], anchor="center").grid(row=2, column=1, sticky="ew")
        ttk.Label(calc_box, textvariable=self.path_calc_vars["max_leg"], anchor="center").grid(row=2, column=2, sticky="ew")
        ttk.Label(calc_box, text="Full-skip 거리").grid(row=3, column=0, sticky="w", pady=2)
        ttk.Label(calc_box, textvariable=self.path_calc_vars["min_skip"], anchor="center").grid(row=3, column=1, sticky="ew")
        ttk.Label(calc_box, textvariable=self.path_calc_vars["max_skip"], anchor="center").grid(row=3, column=2, sticky="ew")
        ttk.Label(calc_box, text="계산 기준: 두께 × tan(굴절각), 웨지 내 경로 제외",
                  foreground="#64748B").grid(row=4, column=0, columnspan=3, sticky="w", pady=(5, 0))
        self.bevel_depth_var = tk.StringVar(
            value="BeamTool 출사점→용접 중심 거리와 루트 중심 오프셋으로 개선면 통과 깊이를 계산합니다."
        )
        ttk.Label(calc_box, textvariable=self.bevel_depth_var, foreground="#0F766E",
                  wraplength=760).grid(row=5, column=0, columnspan=3, sticky="w", pady=(4, 0))
        self.focus_depth_var = tk.StringVar(value="")
        ttk.Label(calc_box, textvariable=self.focus_depth_var, foreground="#1D4ED8",
                  wraplength=760).grid(row=6, column=0, columnspan=3, sticky="w", pady=(4, 0))
        for column in range(3):
            calc_box.columnconfigure(column, weight=1)

        self.status = tk.StringVar(value="입력값을 확인하세요.")
        ttk.Label(preview_outer, textvariable=self.status, foreground="#334155").pack(fill="x", pady=6)

        footer = ttk.Frame(self, padding=(12, 0, 12, 12))
        footer.pack(fill="x")
        ttk.Button(footer, text="Excel 생성", command=self.export_excel).pack(side="right", ipadx=18, ipady=6)
        ttk.Button(footer, text="미리보기 갱신", command=self.refresh_preview).pack(side="right", padx=8, ipady=6)

    def _load_default_settings(self) -> ScanSettings:
        if CONFIG_PATH.exists():
            try:
                return ScanSettings.from_dict(json.loads(CONFIG_PATH.read_text(encoding="utf-8")))
            except Exception:
                pass
        return ScanSettings()

    def _set_values(self, settings: ScanSettings):
        for item in fields(settings):
            if item.name != "apertures":
                value = (
                    0 if item.name == "weld_center_to_root_offset_mm"
                    else getattr(settings, item.name)
                )
                self.vars[item.name].set(str(value))
        self.apertures = []
        for item in settings.apertures:
            aperture_data = asdict(item)
            for key in (
                "beam_exit_measurements",
                "left_beam_exit_measurements",
                "right_beam_exit_measurements",
            ):
                aperture_data[key] = ""
            self.apertures.append(ApertureConfig(**aperture_data))
        self._refresh_aperture_tree()
        self._clear_aperture_editor()

    def _get_settings(self) -> ScanSettings:
        return self._settings_with_apertures(self.apertures)

    def _settings_with_apertures(self, apertures: list[ApertureConfig]) -> ScanSettings:
        numeric_float = {
            "frequency_mhz", "pitch_mm", "elevation_mm", "thickness_mm",
            "bevel_angle_deg", "bevel_tolerance_deg", "min_angle_deg",
            "max_angle_deg", "angle_step_deg", "weld_center_to_root_offset_mm",
        }
        numeric_int = {"total_elements"}
        data = {}
        for item in fields(ScanSettings):
            if item.name == "apertures":
                continue
            value = self.vars[item.name].get().strip()
            if item.name in numeric_float:
                data[item.name] = float(value)
            elif item.name in numeric_int:
                data[item.name] = int(float(value))
            else:
                data[item.name] = value
        data["apertures"] = [ApertureConfig(**asdict(item)) for item in apertures]
        settings = ScanSettings(**data)
        settings.validate()
        return settings

    def _aperture_from_editor(self) -> ApertureConfig:
        try:
            return ApertureConfig(
                probe_id=self.aperture_vars["probe_id"].get().strip(),
                beamset_id=self.aperture_vars["beamset_id"].get().strip(),
                scan_side=self.aperture_vars["scan_side"].get().strip(),
                active_elements=int(float(self.aperture_vars["active_elements"].get().strip())),
                first_element=int(float(self.aperture_vars["first_element"].get().strip())),
                focal_depth_mm=float(self.aperture_vars["focal_depth_mm"].get().strip()),
                index_offsets=self.aperture_vars["index_offsets"].get().strip(),
                left_index_offsets=self.aperture_vars["left_index_offsets"].get().strip(),
                right_index_offsets=self.aperture_vars["right_index_offsets"].get().strip(),
                beam_exit_measurements=self.aperture_vars["beam_exit_measurements"].get().strip(),
                left_beam_exit_measurements=self.aperture_vars["left_beam_exit_measurements"].get().strip(),
                right_beam_exit_measurements=self.aperture_vars["right_beam_exit_measurements"].get().strip(),
            )
        except ValueError as exc:
            raise ValueError("활성소자 설정의 숫자 입력값을 확인하세요.") from exc

    def _refresh_aperture_tree(self):
        self.aperture_tree.delete(*self.aperture_tree.get_children())
        for number, item in enumerate(self.apertures, 1):
            self.aperture_tree.insert("", "end", iid=str(number - 1), values=(
                item.probe_id, item.beamset_id, item.scan_side,
                item.active_elements, item.first_element, f"{item.focal_depth_mm:g}",
                " | ".join(part for part in (
                    f"공통 {item.index_offsets}" if item.index_offsets else "",
                    f"L {item.left_index_offsets}" if item.left_index_offsets else "",
                    f"R {item.right_index_offsets}" if item.right_index_offsets else "",
                ) if part),
                " | ".join(part for part in (
                    f"공통 {item.beam_exit_measurements}" if item.beam_exit_measurements else "",
                    f"L {item.left_beam_exit_measurements}" if item.left_beam_exit_measurements else "",
                    f"R {item.right_beam_exit_measurements}" if item.right_beam_exit_measurements else "",
                ) if part),
            ))

    def _clear_aperture_editor(self):
        defaults = ApertureConfig()
        used_ids = {
            item.beamset_id for item in self.apertures
            if item.probe_id == defaults.probe_id and item.scan_side == defaults.scan_side
        }
        sequence = 1
        while f"A{sequence}" in used_ids:
            sequence += 1
        defaults.beamset_id = f"A{sequence}"
        for key in self.aperture_vars:
            value = getattr(defaults, key)
            self.aperture_vars[key].set("" if value is None else str(value))
        self.aperture_tree.selection_remove(*self.aperture_tree.selection())

    def _select_aperture(self, _event=None):
        selected = self.aperture_tree.selection()
        if not selected:
            return
        item = self.apertures[int(selected[0])]
        for key in self.aperture_vars:
            value = getattr(item, key)
            self.aperture_vars[key].set("" if value is None else str(value))

    def add_aperture(self):
        try:
            item = self._aperture_from_editor()
            self._settings_with_apertures([*self.apertures, item])
        except Exception as exc:
            messagebox.showerror("입력 오류", str(exc), parent=self)
            return
        self.apertures.append(item)
        self._refresh_aperture_tree()
        self._clear_aperture_editor()
        self.refresh_preview()

    def update_aperture(self):
        selected = self.aperture_tree.selection()
        if not selected:
            messagebox.showinfo("선택 필요", "수정할 설정을 선택하세요.", parent=self)
            return
        try:
            item = self._aperture_from_editor()
            trial = list(self.apertures)
            trial[int(selected[0])] = item
            self._settings_with_apertures(trial)
        except Exception as exc:
            messagebox.showerror("입력 오류", str(exc), parent=self)
            return
        self.apertures[int(selected[0])] = item
        self._refresh_aperture_tree()
        self.refresh_preview()

    def delete_aperture(self):
        selected = self.aperture_tree.selection()
        if not selected:
            messagebox.showinfo("선택 필요", "삭제할 설정을 선택하세요.", parent=self)
            return
        if len(self.apertures) == 1:
            messagebox.showinfo("삭제 불가", "활성소자 설정은 하나 이상 필요합니다.", parent=self)
            return
        del self.apertures[int(selected[0])]
        self._refresh_aperture_tree()
        self._clear_aperture_editor()
        self.refresh_preview()

    def refresh_preview(self):
        self.preview.delete("all")
        try:
            settings = self._get_settings()
        except Exception as exc:
            self.status.set(f"입력 오류: {exc}")
            self.preview.create_text(20, 20, text=str(exc), anchor="nw", fill="#B91C1C")
            for var in self.path_calc_vars.values():
                var.set("-")
            self.bevel_depth_var.set("개선면 통과 깊이: 계산 불가")
            self.focus_depth_var.set("실제 초점 깊이: 계산 불가")
            return
        width = max(400, self.preview.winfo_width())
        height = max(300, self.preview.winfo_height())
        margin = 45
        top, bottom = height * 0.42, height * 0.75
        center = width / 2
        self.preview.create_rectangle(margin, top, width - margin, bottom, fill="#F8FAFC", outline="#334155", width=2)
        plate_h = bottom - top
        bevel = max(15, min(60, plate_h / math.tan(math.radians(settings.bevel_angle_deg))))
        self.preview.create_polygon(center - bevel, top, center, bottom, center + bevel, top,
                                    fill="#E2E8F0", outline="#475569")
        scale = plate_h / settings.thickness_mm
        colors = ("#16A34A", "#D97706", "#DB2777", "#2563EB")
        for side in (-1, 1):
            for config_pos, aperture in enumerate(settings.apertures):
                for index in aperture.index_values_for_side(side):
                    x = center + side * (bevel + index * scale * 0.55)
                    y = top - 34 - config_pos * 24
                    color = colors[config_pos % len(colors)]
                    self.preview.create_rectangle(x - 24, y, x + 24, y + 22, fill="#A5B4FC", outline=color)
                    self.preview.create_text(
                        x, y - 7, text=f"{aperture.probe_id}-{aperture.beamset_id}", fill=color
                    )
                    for angle_pos, angle in enumerate((settings.min_angle_deg, settings.max_angle_deg)):
                        dx = settings.thickness_mm * math.tan(math.radians(angle)) * scale
                        hit_x = x - side * dx
                        beam_color = colors[(config_pos * 2 + angle_pos) % len(colors)]
                        self.preview.create_line(x, top, hit_x, bottom, fill=beam_color, width=2)
                        self.preview.create_line(hit_x, bottom, hit_x - side * dx, top, fill=beam_color, width=2)
        law_count = int(round((settings.max_angle_deg - settings.min_angle_deg) / settings.angle_step_deg)) + 1
        min_leg = settings.thickness_mm * math.tan(math.radians(settings.min_angle_deg))
        max_leg = settings.thickness_mm * math.tan(math.radians(settings.max_angle_deg))
        self.path_calc_vars["min_angle"].set(f"{settings.min_angle_deg:g}°")
        self.path_calc_vars["max_angle"].set(f"{settings.max_angle_deg:g}°")
        self.path_calc_vars["min_leg"].set(f"{min_leg:.2f} mm")
        self.path_calc_vars["max_leg"].set(f"{max_leg:.2f} mm")
        self.path_calc_vars["min_skip"].set(f"{2 * min_leg:.2f} mm")
        self.path_calc_vars["max_skip"].set(f"{2 * max_leg:.2f} mm")
        perpendicular_angle = 90.0 - settings.bevel_angle_deg
        depth_results = []
        for number, aperture in enumerate(settings.apertures, 1):
            angle_distance_pairs = aperture.exit_measurements()
            if not angle_distance_pairs:
                depth_results.append(
                    f"{aperture.probe_id}-{aperture.beamset_id}: "
                    "각도별 BeamTool 출사점→용접 중심 수평거리 미입력"
                )
                continue
            aperture_results = []
            for side, angle, distance in angle_distance_pairs:
                root_distance = root_center_distance(
                    distance, side, settings.weld_center_to_root_offset_mm,
                )
                depth = reflected_bevel_depth(
                    settings.thickness_mm, settings.bevel_angle_deg, angle, root_distance,
                )
                depth_text = "교차 없음" if depth is None else f"{depth:.2f} mm"
                deviation = abs(angle - perpendicular_angle)
                compliance = "±6° 충족" if deviation <= 6.0 else f"편차 {deviation:.2f}°"
                aperture_results.append(
                    f"{side} {angle:g}° BeamTool={distance:g}, "
                    f"루트보정={root_distance:g} → {depth_text} ({compliance})"
                )
            depth_results.append(
                f"{aperture.probe_id}-{aperture.beamset_id}: " + " / ".join(aperture_results)
            )
        self.bevel_depth_var.set("개선면 통과 깊이 | " + "   ".join(depth_results))
        focus_results = []
        for number, aperture in enumerate(settings.apertures, 1):
            actual_depth, region = physical_focus_depth(
                settings.thickness_mm, aperture.focal_depth_mm,
            )
            actual_text = "계산 범위 밖" if actual_depth is None else f"{actual_depth:.2f} mm"
            focus_results.append(
                f"{aperture.probe_id}-{aperture.beamset_id}: unfolded "
                f"{aperture.focal_depth_mm:g} mm → 실제 {actual_text} ({region})"
            )
        self.focus_depth_var.set("초점 위치 | " + "   ".join(focus_results))
        apertures = ", ".join(
            f"{item.probe_id}-{item.beamset_id} {item.active_elements * settings.pitch_mm:.2f} mm"
            for item in settings.apertures
        )
        self.status.set(f"검증 완료 | 설정 {len(settings.apertures)}개 | Focal laws {law_count}개 | {apertures}")

    def reset_defaults(self):
        self._set_values(ScanSettings())
        self.refresh_preview()

    def save_settings(self):
        try:
            settings = self._get_settings()
        except Exception as exc:
            messagebox.showerror("입력 오류", str(exc), parent=self)
            return
        path = filedialog.asksaveasfilename(
            parent=self, title="설정 저장", defaultextension=".json",
            initialfile="PAUT_검사설정.json", filetypes=[("JSON", "*.json")],
        )
        if path:
            Path(path).write_text(json.dumps(asdict(settings), ensure_ascii=False, indent=2), encoding="utf-8")

    def load_settings(self):
        path = filedialog.askopenfilename(parent=self, title="설정 불러오기", filetypes=[("JSON", "*.json")])
        if not path:
            return
        try:
            settings = ScanSettings.from_dict(json.loads(Path(path).read_text(encoding="utf-8")))
            settings.validate()
            self._set_values(settings)
            self.refresh_preview()
        except Exception as exc:
            messagebox.showerror("불러오기 오류", str(exc), parent=self)

    def export_excel(self):
        try:
            settings = self._get_settings()
        except Exception as exc:
            messagebox.showerror("입력 오류", str(exc), parent=self)
            return
        default_name = (
            f"PAUT_검사설정_{settings.probe_model}_{settings.thickness_mm:g}mm_"
            f"{settings.min_angle_deg:g}-{settings.max_angle_deg:g}도.xlsx"
        )
        path = filedialog.asksaveasfilename(
            parent=self, title="Excel 저장", defaultextension=".xlsx",
            initialfile=default_name, filetypes=[("Excel", "*.xlsx")],
        )
        if not path:
            return
        try:
            create_workbook(settings, path)
            CONFIG_PATH.write_text(json.dumps(asdict(settings), ensure_ascii=False, indent=2), encoding="utf-8")
            messagebox.showinfo("완료", f"검사설정표를 생성했습니다.\n{path}", parent=self)
        except Exception as exc:
            messagebox.showerror("생성 오류", str(exc), parent=self)


def self_test(output_path: str) -> None:
    focus_20, region_20 = physical_focus_depth(15.87, 20.0)
    focus_30, region_30 = physical_focus_depth(15.87, 30.0)
    assert abs(focus_20 - 11.74) < 1e-8 and region_20 == "0.5~1 skip 반사 구간"
    assert abs(focus_30 - 1.74) < 1e-8 and region_30 == "0.5~1 skip 반사 구간"
    assert abs(reflected_bevel_depth(15.87, 37.5, 46.5, 27.45) - 9.98) < 0.02
    assert abs(reflected_bevel_depth(15.87, 37.5, 58.5, 27.45) - 15.22) < 0.02
    assert abs(reflected_bevel_depth(15.87, 37.5, 56.09, 39.94) - 8.63) < 0.02
    assert reflected_bevel_depth(15.87, 37.5, 70.0, 27.45) is None
    assert root_center_distance(75.63, "좌측", 2.0) == 77.63
    assert root_center_distance(75.63, "우측", 2.0) == 73.63
    legacy = ScanSettings.from_dict({"apertures": [{
        "active_elements": 32, "first_element": 5, "focal_depth_mm": 15.87,
        "index_offsets": "15", "exit_to_weld_mm": 39.94,
    }]})
    assert legacy.apertures[0].exit_measurements() == [("양측", 52.5, 39.94)]
    migrated = ScanSettings.from_dict({"apertures": [
        {"active_elements": 32, "first_element": 5, "focal_depth_mm": 20, "index_offsets": "15"},
        {"active_elements": 32, "first_element": 27, "focal_depth_mm": 30, "index_offsets": "15"},
    ]})
    assert [item.beamset_id for item in migrated.apertures] == ["A1", "A2"]
    measured = ApertureConfig(
        left_beam_exit_measurements="46.58:27.45",
        right_beam_exit_measurements="56.09=39.94",
    )
    assert measured.exit_measurements() == [
        ("좌측", 46.58, 27.45), ("우측", 56.09, 39.94),
    ]
    settings = ScanSettings(weld_center_to_root_offset_mm=2.0, apertures=[
        ApertureConfig(probe_id="P1", beamset_id="A1", scan_side="좌측",
                       active_elements=32, first_element=1,
                       focal_depth_mm=15.87, index_offsets="42",
                       beam_exit_measurements="46.58:27.45,56.09:39.94"),
        ApertureConfig(probe_id="P1", beamset_id="A2", scan_side="우측",
                       active_elements=16, first_element=33,
                       focal_depth_mm=12.0, index_offsets="48",
                       beam_exit_measurements="58.5:40.38"),
    ])
    created = create_workbook(settings, output_path)
    check = openpyxl.load_workbook(created, data_only=False)
    sheet = check["PAUT 설정표"]
    assert sheet["B5"].value == 15.87
    assert sheet["E9"].value == "=F14*E14"
    assert sheet["E10"].value == "=F15*E15"
    assert sheet["G9"].value == "=INT((E6-B6)/K5)+1"
    assert sheet["F14"].value == 32
    assert sheet["F15"].value == 16
    assert sheet["D15"].value == "=C15*E15"
    assert sheet.column_dimensions["K"].width == 34
    assert sheet.row_dimensions[14].height == 34
    assert sheet["K14"].alignment.wrap_text is True
    assert sheet["A39"].value is None
    assert sheet.print_options.horizontalCentered is True
    assert sheet.page_margins.left == sheet.page_margins.right == 0.25
    assert len(sheet._images) == 1
    assert sheet.print_area.endswith("$A$1:$K$47")
    check.close()
    return
    assert "D40:K40" in {str(cell_range) for cell_range in sheet.merged_cells.ranges}
    assert "D41:K41" in {str(cell_range) for cell_range in sheet.merged_cells.ranges}
    assert "D42:K42" in {str(cell_range) for cell_range in sheet.merged_cells.ranges}
    assert sheet["A43"].value == "True depth (unfolded) 기준 실제 초점 위치"
    assert sheet["B45"].value == 15.87
    assert sheet["C45"].value.startswith("=IF(B45<=$B$5")
    assert sheet["A47"].value == "0.5 skip 반사 후 개선면 통과 깊이"
    assert sheet["A49"].value == "P1-A1"
    assert sheet["B49"].value == "좌측"
    assert sheet["C49"].value == "42"
    assert sheet["D49"].value == -42
    assert sheet["E49"].value == 46.58
    assert sheet["F49"].value == 27.45
    assert sheet["G49"].value == 2.0
    assert sheet["H49"].value == "=F49+G49"
    assert sheet["I49"].value.startswith("=(2*$B$5*TAN")
    assert "ABS(E49-(90-$E$5))<=6" in sheet["J49"].value
    assert sheet.print_options.horizontalCentered is True
    assert sheet.page_margins.left == sheet.page_margins.right == 0.25
    assert len(sheet._images) == 1
    assert sheet.print_area == "'PAUT 설정표'!$A$1:$K$52"
    check.close()


def main():
    parser = argparse.ArgumentParser(description=APP_TITLE)
    parser.add_argument("--self-test", metavar="OUTPUT", help="GUI 없이 샘플 Excel을 생성하고 검증")
    args = parser.parse_args()
    if args.self_test:
        self_test(args.self_test)
        print(f"SELF TEST OK: {args.self_test}")
        return
    app = PautScanPlanApp()
    app.mainloop()


if __name__ == "__main__":
    main()
