"""Central branch submission metadata; source records and files are read only.

Run through 중앙지사_속성정보_생성.py. No changes to either business app.
"""
from __future__ import annotations

import datetime as dt
import copy
import hashlib
import json
import os
import re
import tempfile
from pathlib import Path
from urllib.parse import quote

METHODS = ('RT', 'PAUT', 'MT', 'PT')
FIELDS = ('공사명', '시공사', '하도사', '용역사', '도면번호', '라인번호',
          '용접부번호', '용접부특성', '검사일자', '결함종류', '보고서종류', '보고서 매수')
LINKS = ('보고서', '필름영상화', '데이터')
IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.tif', '.tiff', '.bmp', '.dcm', '.dicom', '.pdf'}
POINT_PATTERN = re.compile(r'(?<![0-9A-Za-z])(\d+\s*-\s*\d+)(?![0-9A-Za-z])')
CACHE_NAME = '.central_attribute_connections.json'
REPORT_EXTENSIONS = {'.pdf', '.xlsx', '.xls', '.xlsm', '.doc', '.docx', '.hwp', '.hwpx'}
SUPPLEMENT_FIELDS = (*FIELDS, *LINKS, '검사결과', '보고서번호', 'RT영상', '촬영포인트', '검사차수', '자동연결영상')


def parse_points(value):
    values = value if isinstance(value, list) else re.split(r'[,;/\n]+', text(value))
    values = [text(v).replace(' ', '') for v in values if text(v)]
    return list(dict.fromkeys(values)) if values and all(re.fullmatch(r'\d+-\d+', v) for v in values) else []


def positive_quantity(value):
    try:
        return float(text(value).replace(',', '')) > 0
    except ValueError:
        return False


def text(value):
    return '' if value is None else str(value).strip()


def first(record, *keys):
    return next((text(record.get(k)) for k in keys if text(record.get(k))), '')


def date_text(value):
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime('%Y-%m-%d')
    value = text(value)
    for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d', '%Y%m%d', '%Y-%m-%d %H:%M:%S'):
        try:
            return dt.datetime.strptime(value, fmt).strftime('%Y-%m-%d')
        except ValueError:
            pass
    return value


def normalize(record, method='', date='', source=''):
    row = {key: text(record.get(key)) for key in FIELDS}
    row.update({
        '검사일자': date_text(first(record, '검사일자', 'Date') or date),
        '보고서종류': first(record, '보고서종류', '검사방법') or method,
        '라인번호': first(record, '라인번호', 'Line No.', 'Line'),
        '도면번호': first(record, '도면번호', 'Dwg', 'ISO'),
        '용접부번호': first(record, '용접부번호', 'Joint No.', 'Joint No', 'Joint'),
        '결함종류': first(record, '결함종류', 'Type of Flaw'),
        '용접부특성': first(record, '용접부특성'),
        '출처': source,
        '확인사항': [],
        '검사결과': first(record, '검사결과', '결과', 'Evaluation', 'Result'),
        '보고서번호': first(record, '보고서번호', '보고서 번호', 'Report No.', 'Report No', 'ReportNo', 'report_no'),
        'RT영상': record.get('RT영상', {}),
        '촬영포인트': parse_points(record.get('촬영포인트') or first(record, '구간정보', 'Loc')) if (first(record, '보고서종류', '검사방법') or method).upper() == 'RT' else [],
        '검사차수': first(record, '검사차수', '재검사차수'),
    })
    # 업체 is deliberately not assigned to contractor/subcontractor without a rule.
    if record.get('업체'):
        row['확인사항'].append('원본 업체: ' + text(record['업체']) + ' (시공사/하도사 구분 확인)')
    if not row['용접부특성'] and (positive_quantity(record.get('RT_OR')) or positive_quantity(record.get('RT_RE'))):
        ori = positive_quantity(record.get('RT_OR'))
        if ori:
            row['용접부특성'] = 'O'
        else:
            row['확인사항'].append('RT_RE는 불량매수: 전체 RT_OR 및 용접부특성 확인')
    for key in LINKS:
        row[key] = first(record, key + '경로', key + '파일')
    return row


def identity(row):
    # Do not equate drawings and line numbers, or remove identifier punctuation.
    return tuple(text(row.get(k)).upper() for k in
                 ('검사일자', '보고서종류', '라인번호', '용접부번호', '용접부특성'))


def load_history(path):
    with open(path, encoding='utf-8-sig') as stream:
        history = json.load(stream)
    if not isinstance(history, dict):
        raise ValueError('중앙지사 이력은 날짜별 객체 형식이어야 합니다.')
    rows = []
    for date, day in history.items():
        if not isinstance(day, dict):
            raise ValueError(f'잘못된 일자 데이터: {date}')
        for index, item in enumerate(day.get('ndt_results', [])):
            if not isinstance(item, dict):
                raise ValueError(f'{date} 검사 기록 형식 오류')
            if text(item.get('검사방법')).upper() in METHODS:
                row = normalize(item, date=date, source=f'중앙지사 {date} #{index + 1}')
                row['보고서종류'] = row['보고서종류'].upper()
                row['원본값'] = {key: copy.deepcopy(row.get(key, '')) for key in SUPPLEMENT_FIELDS}
                row['id'] = hashlib.sha256((str(Path(path).resolve()) + '|' + date + '|' + str(index)).encode()).hexdigest()[:20]
                rows.append(row)
    return rows


def load_report_data(path, method):
    if method not in METHODS:
        raise ValueError('RT, PAUT, MT, PT 중 검사종류를 선택하세요.')
    if Path(path).suffix.lower() == '.json':
        with open(path, encoding='utf-8-sig') as stream:
            data = json.load(stream)
    else:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = book.active
            values = iter(sheet.values)
            headers = [text(v) for v in next(values)]
            data = [dict(zip(headers, values)) for values in values if any(v is not None for v in values)]
        finally:
            book.close()
    if not isinstance(data, list) or any(not isinstance(item, dict) for item in data):
        raise ValueError('보고서 앱의 데이터 저장(JSON) 또는 표 형태의 속성정보 XLSX를 선택하세요.')
    rows = []
    for index, item in enumerate(data):
        if item.get('selected') is False or item.get('date_filtered') is False:
            continue
        row = normalize(item, method=method, source=f'{Path(path).name} #{index + 1}')
        if not row['용접부번호']:
            raise ValueError(f'{index + 1}번째 보고서 데이터에 용접부번호가 없습니다.')
        if row['보고서종류'].upper() != method:
            raise ValueError(f'{index + 1}번째 데이터의 검사종류가 선택한 {method}와 다릅니다.')
        row['보고서종류'] = method
        row['id'] = hashlib.sha256((str(Path(path).resolve()) + '|' + str(index)).encode()).hexdigest()[:20]
        rows.append(row)
    return rows


def merge_rows(history, reports, unmatched=None):
    """Central history owns the row list; unmatched reports are diagnostics only."""
    rows = copy.deepcopy(history)
    groups = {}
    for row in rows:
        groups.setdefault(identity(row), []).append(row)
    for group in groups.values():
        if len(group) > 1:
            for row in group:
                row['확인사항'].append('중앙지사 대조 키 중복: 자료 자동 병합·연결 복원 제외')
    report_groups = {}
    for row in reports:
        report_groups.setdefault(identity(row), []).append(row)
    for key, group in report_groups.items():
        matches = groups.get(key, [])
        complete = all(key[:4])  # Empty weld characteristic can match empty only.
        if complete and len(matches) == len(group) == 1:
            target, report = matches[0], group[0]
            for field in (*FIELDS, *LINKS, '검사결과', '보고서번호'):
                new, old = text(report.get(field)), text(target.get(field))
                if new and old and new != old:
                    target['확인사항'].append(f'{field} 불일치: 중앙지사={old}, 보고서={new}')
                elif new:
                    target[field] = new
            target['출처'] += ' / ' + report['출처']
            target['확인사항'].extend(report['확인사항'])
            if report.get('RT영상'):
                if target.get('RT영상') and target['RT영상'] != report['RT영상']:
                    target['확인사항'].append('RT 포인트 영상 연결 불일치')
                else:
                    target['RT영상'] = report['RT영상']
            for field in ('촬영포인트', '검사차수'):
                if not target.get(field) and report.get(field):
                    target[field] = copy.deepcopy(report[field])
                elif report.get(field) and target.get(field) != report[field]:
                    target['확인사항'].append(field + ' 보고서 불일치')
        else:
            reason = '대조 키 누락' if not complete else ('중복 후보' if matches else '중앙지사 일치 기록 없음')
            for report in group:
                if unmatched is not None:
                    unmatched.append({**copy.deepcopy(report), '대조결과': reason})
            for target in matches:
                target['확인사항'].append('보고서 대조 확인: ' + reason)
    return rows


def restore_connections(current, saved):
    """Restore supplements only to a unique unchanged source identity.

    Source-filled values always win. Old rows never become current inspection rows.
    Duplicate or incomplete identities cannot safely recover file attachments.
    """
    result = copy.deepcopy(current)
    current_groups, saved_groups = {}, {}
    for row in result:
        current_groups.setdefault(identity(row.get('원본값', row)), []).append(row)
    for row in saved:
        saved_groups.setdefault(identity(row.get('원본값', row)), []).append(row)
    for key, group in current_groups.items():
        previous = saved_groups.get(key, [])
        if not all(key[:4]) or len(group) != 1 or len(previous) != 1:
            continue
        target, old = group[0], previous[0]
        # Saved warnings remain until their underlying discrepancy is reviewed.
        target['확인사항'] = list(dict.fromkeys([*target.get('확인사항', []), *old.get('확인사항', [])]))
        source = target.get('원본값', target)
        previous_source = old.get('원본값', {})
        for field in ('검사결과', '도면번호', '촬영포인트', '검사차수'):
            if field in previous_source and previous_source[field] != source.get(field):
                target['확인사항'].append(f'중앙지사 {field} 변경: 보고서/영상 연결 재확인')
        for field in SUPPLEMENT_FIELDS:
            if not source.get(field) and old.get(field):
                target[field] = copy.deepcopy(old[field])
    return result


def read_connections(path):
    with open(path, encoding='utf-8-sig') as stream:
        data = json.load(stream)
    rows = data.get('rows') if isinstance(data, dict) and data.get('version') in (1, 2, 3) else None
    if isinstance(rows, list):
        pending = data.get('unresolved', [])
        if not isinstance(pending, list):
            raise ValueError('연결정보 복원보류 목록 형식 오류')
        rows = rows + pending
    if not isinstance(rows, list) or any(not isinstance(r, dict) or not all(k in r for k in (*FIELDS, 'id', '확인사항')) for r in rows):
        raise ValueError('올바른 속성정보 연결정보 파일이 아닙니다.')
    return rows


def write_connections(path, rows, root):
    if not root:
        raise ValueError('제출 폴더를 먼저 선택하세요.')
    path = Path(path)
    previous = read_connections(path) if path.exists() else []
    saved = copy.deepcopy(rows)
    for row in saved:
        for field in LINKS:
            target = contained_path(root, row.get(field))
            if target:
                row[field] = target.relative_to(Path(root).resolve()).as_posix()
        if row.get('보고서종류') == 'RT':
            # Save explicit attachments; folder-derived points are rescanned next time.
            mapping = row.get('RT영상') or {}
            row['RT영상'] = {point: [contained_path(root, value).relative_to(Path(root).resolve()).as_posix()
                                    for value in values] for point, values in point_images({**row, '필름영상화': ''}, root).items()} if mapping else {}
    current_groups, previous_groups = {}, {}
    for row in saved:
        current_groups.setdefault(identity(row.get('원본값', row)), []).append(row)
    for row in previous:
        previous_groups.setdefault(identity(row.get('원본값', row)), []).append(row)
    unresolved = []
    seen = set()
    for key, group in previous_groups.items():
        if all(key[:4]) and len(group) == 1 and len(current_groups.get(key, [])) == 1:
            continue
        for row in group:
            source = row.get('원본값', {})
            if not any(row.get(field) and row.get(field) != source.get(field) for field in SUPPLEMENT_FIELDS):
                continue
            signature = json.dumps(row, ensure_ascii=False, sort_keys=True)
            if signature not in seen:
                unresolved.append(copy.deepcopy(row))
                seen.add(signature)
    descriptor, temporary = tempfile.mkstemp(prefix='.connections_', suffix='.json', dir=path.parent)
    os.close(descriptor)
    try:
        Path(temporary).write_text(json.dumps({'version': 3, 'rows': saved, 'unresolved': unresolved}, ensure_ascii=False, indent=2), encoding='utf-8')
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)


def contained_path(root, value):
    if not value:
        return None
    root = Path(root).resolve()
    path = Path(value)
    path = (path if path.is_absolute() else root / path).resolve()
    if not path.is_relative_to(root):
        raise ValueError('링크 대상은 제출 폴더 안에 있어야 합니다: ' + str(path))
    return path


def inventory(root):
    root = Path(root).resolve()
    return sorted(path for method in METHODS if (root / method).is_dir()
                  for path in (root / method).rglob('*')
                  if path.is_file() and path.resolve().is_relative_to(root) and not path.name.startswith('~$'))


def folder_component(value):
    value = text(value)
    if not value:
        raise ValueError('폴더 생성에 필요한 라인번호·용접부번호·검사일자가 없습니다.')
    clean = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', value).rstrip(' .')
    reserved = re.fullmatch(r'(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?', clean, re.I)
    if clean != value or not clean or len(clean) > 90 or reserved:
        clean = (clean[:75] or '항목') + '_' + hashlib.sha256(value.encode()).hexdigest()[:10]
    return clean


def submission_plan(rows):
    plans, grouped = [], {}
    for row in rows:
        method = row.get('보고서종류')
        if method not in METHODS:
            raise ValueError('지원하지 않는 검사종류: ' + text(method))
        date = dt.date.fromisoformat(row['검사일자']).isoformat()
        kind = {'O': '최초', 'REPAIR': '보수'}.get(row.get('용접부특성'), '검사구분미확인')
        if row.get('검사차수'):
            kind += '_' + folder_component(row['검사차수'])
        path = Path(method) / folder_component(row['라인번호']) / folder_component(row['용접부번호']) / (date + '_' + kind)
        grouped.setdefault(path.as_posix().casefold(), []).append((row, path))
    for group in grouped.values():
        for row, path in group:
            if len(group) > 1:
                token = hashlib.sha256(json.dumps([row['id'], row.get('원본값', identity(row))], sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:12]
                path = path.with_name(path.name + '_기록-' + token)
            plans.append((row, path))
    return plans


def create_submission_folders(rows, root):
    if not rows:
        raise ValueError('중앙지사 이력을 먼저 불러오세요.')
    root = Path(root).resolve()
    plans = submission_plan(rows)  # Validate the entire plan before creating anything.
    targets = []
    for method in METHODS:
        targets.append(root / method)
        targets.append(root / method / '보고서')  # Shared reports need only one copy.
    for row, path in plans:
        targets.append(root / path / '보고서')
        if row['보고서종류'] == 'RT':
            targets.append(root / path / '필름영상화')
        elif row['보고서종류'] == 'PAUT':
            targets.append(root / path / '데이터')
    for target in targets:
        contained_path(root, str(target))
        if target.exists() and not target.is_dir():
            raise ValueError('폴더 위치에 파일이 있습니다: ' + str(target))
    for target in targets:
        target.mkdir(parents=True, exist_ok=True)
    return len(plans)


def auto_connect(rows, root):
    """Resolve only the generated record folder or an exact shared report number."""
    root = Path(root).resolve()
    result = copy.deepcopy(rows)
    for row, relative in submission_plan(result):
        directory = contained_path(root, relative)
        notes = row.setdefault('확인사항', [])
        if not row.get('보고서'):
            own = directory / '보고서'
            candidates = [p for p in own.iterdir() if p.is_file() and p.suffix.lower() in REPORT_EXTENSIONS and not p.name.startswith('~$')] if own.is_dir() else []
            if not candidates and row.get('보고서번호'):
                shared = contained_path(root, Path(row['보고서종류']) / '보고서')
                candidates = [p for p in shared.iterdir() if p.is_file() and p.suffix.lower() in REPORT_EXTENSIONS
                              and p.stem.casefold() == row['보고서번호'].casefold()] if shared.is_dir() else []
            if row.get('보고서번호'):
                candidates = [p for p in candidates if p.stem.casefold() == row['보고서번호'].casefold()]
            if len(candidates) == 1:
                target = contained_path(root, str(candidates[0]))
                row['보고서'] = target.relative_to(root).as_posix()
                if not row.get('보고서번호'):
                    row['보고서번호'] = target.stem
            elif len(candidates) > 1:
                message = '자동 연결 보류: 보고서 후보가 여러 개입니다.'
                if message not in notes:
                    notes.append(message)
        if row['보고서종류'] == 'RT' and (row.get('자동연결영상') or (not row.get('RT영상') and not row.get('필름영상화'))):
            images = directory / '필름영상화'
            if images.is_dir():
                # This exact path was derived from this record, not discovered by point alone.
                mapping = {}
                for target in sorted(images.iterdir()):
                    if not target.is_file() or target.suffix.lower() not in IMAGE_EXTENSIONS:
                        continue
                    matches = POINT_PATTERN.findall(target.stem)
                    if len(matches) == 1:
                        safe = contained_path(root, str(target))
                        mapping.setdefault(matches[0].replace(' ', ''), []).append(safe.relative_to(root).as_posix())
                if mapping:
                    row['RT영상'] = mapping
                    row['필름영상화'] = images.relative_to(root).as_posix()
                    row['자동연결영상'] = True
        elif row['보고서종류'] == 'PAUT' and not row.get('데이터'):
            data = directory / '데이터'
            if data.is_dir() and any(p.is_file() for p in data.rglob('*')):
                row['데이터'] = contained_path(root, str(data)).relative_to(root).as_posix()
    return result


def point_images(row, root):
    """Use explicit per-point paths or scan only the record's assigned image folder.

    Never search all RT images by point number: every weld can have a 1-2 image.
    Repeated point images are retained as distinct links rather than picked once.
    """
    if row.get('보고서종류') != 'RT':
        return {}
    mapping = row.get('RT영상') or {}
    if not isinstance(mapping, dict):
        raise ValueError('RT영상은 {"1-2": ["상대경로"]} 형식이어야 합니다.')
    if mapping:
        result = {}
        for label, values in mapping.items():
            label = text(label).replace(' ', '')
            if not re.fullmatch(r'\d+-\d+', label):
                raise ValueError('RT 포인트 표기 확인: ' + label)
            values = [values] if isinstance(values, str) else values
            if not isinstance(values, list) or any(not isinstance(v, str) for v in values):
                raise ValueError('RT 포인트별 영상 경로는 문자열 목록이어야 합니다.')
            result[label] = list(dict.fromkeys(values))
        return result
    folder = contained_path(root, row.get('필름영상화'))
    if not folder or not folder.exists():
        return {}
    if folder.is_dir():
        relative = folder.relative_to(Path(root).resolve())
        joint = text(row.get('용접부번호'))
        line = text(row.get('라인번호'))
        # Shared method/image roots cannot establish ownership of bare 1-2 files.
        # Only a line/weld-specific folder may be scanned without explicit mapping.
        if not joint or not line or line.casefold() not in [p.casefold() for p in relative.parts] or joint.casefold() not in [p.casefold() for p in relative.parts]:
            raise ValueError('RT 자동 검색은 라인번호/용접부번호가 경로에 포함된 전용 폴더만 가능합니다. 포인트별 영상 파일을 직접 지정하세요.')
        # Avoid recursively collecting other welds or repeat inspections below it.
        paths = sorted(folder.iterdir())
    else:
        raise ValueError('RT 영상 파일은 포인트별 영상 연결에서 해당 포인트를 직접 지정하세요.')
    result = {}
    for path in paths:
        if not path.is_file() or path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue
        matches = POINT_PATTERN.findall(path.stem)
        # Multiple numeric ranges could be dates or identifiers; require manual mapping.
        if len(matches) == 1:
            point = matches[0].replace(' ', '')
            resolved = contained_path(root, str(path))
            result.setdefault(point, []).append(resolved.relative_to(Path(root).resolve()).as_posix())
    return result


def issues(row, root):
    result = list(row.get('확인사항', []))
    if row.get('보고서종류') not in METHODS:
        result.append('보고서종류 확인: RT/PAUT/MT/PT 중 선택')
    if not row.get('검사결과'):
        result.append('검사결과 미입력')
    if not row.get('보고서번호'):
        result.append('보고서번호 미입력')
    for field in ('공사명', '시공사', '용역사', '도면번호', '라인번호', '용접부번호', '용접부특성', '검사일자'):
        if not text(row.get(field)):
            result.append(field + ' 미입력')
    try:
        dt.date.fromisoformat(text(row.get('검사일자')))
    except ValueError:
        result.append('검사일자 형식 확인')
    if row.get('용접부특성') not in ('O', 'REPAIR'):
        result.append('용접부특성 확인')
    pages = text(row.get('보고서 매수'))
    if not pages.isdigit() or int(pages) < 1:
        result.append('보고서 매수 확인')
    required = ['보고서'] + (['데이터'] if row['보고서종류'] == 'PAUT' else [])
    for field in required:
        try:
            target = contained_path(root, row.get(field))
            if target is None:
                result.append(field + ' 미연결')
            elif not target.exists():
                result.append(field + ' 파일/폴더 없음')
            elif not target.is_relative_to(Path(root).resolve() / row['보고서종류']):
                result.append(field + ' 검사종류 폴더 불일치')
            elif target.is_dir() and not any(p.is_file() for p in target.rglob('*')):
                result.append(field + ' 폴더 비어 있음')
        except ValueError as error:
            result.append(str(error))
    if row['보고서종류'] == 'RT':
        try:
            images = point_images(row, root)
            for point in row.get('촬영포인트', []):
                if not images.get(point):
                    result.append(f'RT {point} 영상 미연결 (중앙지사 촬영포인트)')
            if not images or not any(images.values()):
                result.append('RT 포인트 영상 미연결: 1-2, 2-3 등의 영상 파일을 지정하세요')
            for point, values in images.items():
                if not values:
                    result.append(f'RT {point} 영상 미연결')
                for value in values:
                    target = contained_path(root, value)
                    if target is None or not target.is_file():
                        result.append(f'RT {point} 영상 파일 없음: {value}')
                    elif not target.is_relative_to(Path(root).resolve() / 'RT'):
                        result.append(f'RT {point} 영상이 RT 폴더 밖에 있음')
                    elif target.suffix.lower() not in IMAGE_EXTENSIONS:
                        result.append(f'RT {point} 영상 형식 확인: {target.suffix}')
        except ValueError as error:
            result.append(str(error))
    return list(dict.fromkeys(result))


def export_workbook(rows, root, destination, unmatched=None, pending=None):
    """Native Excel hyperlinks; never generate executable formulas from source text."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    root, destination = Path(root).resolve(), Path(destination).resolve()
    if destination.parent != root:
        raise ValueError('속성정보 엑셀은 제출 폴더 최상위에 저장하세요.')
    if not rows:
        raise ValueError('생성할 검사 기록이 없습니다.')
    if destination.exists():
        raise FileExistsError('같은 이름의 파일이 있습니다. 다른 이름을 선택하세요.')
    book = Workbook()
    sheet = book.active
    sheet.title = '속성정보'
    mappings = []
    for record in rows:
        try:
            mappings.append(point_images(record, root))
        except ValueError:
            mappings.append({})
    point_counts = {}
    for record, mapping in zip(rows, mappings):
        for point, values in mapping.items():
            point_counts[point] = max(point_counts.get(point, 0), len(values))
        for point in record.get('촬영포인트', []):
            point_counts[point] = max(point_counts.get(point, 0), 1)
    point_columns = [(point, index) for point in sorted(point_counts, key=lambda p: tuple(map(int, p.split('-'))))
                     for index in range(point_counts[point])]
    headers = (*FIELDS, '보고서번호 (클릭하여 열기)', '영상 보관 위치', 'PAUT 데이터 열기', '확인사항', '데이터 출처', '기록ID',
               '검사결과', 'RT 영상 수', *(f'RT {p}' + (f' ({i + 1})' if i else '') for p, i in point_columns))
    sheet.append(headers)

    def put(ws, row, col, value):
        cell = ws.cell(row, col, value)
        if isinstance(value, str):
            cell.data_type = 's'  # Source values beginning with '=' remain literal.
        return cell

    def link(ws, row, col, path, label):
        cell = put(ws, row, col, label)
        # Encode #, %, spaces, and non-ASCII names for file-relative URI targets.
        cell.hyperlink = quote(path.relative_to(root).as_posix(), safe='/')
        cell.style = 'Hyperlink'

    for index, record in enumerate(rows, 2):
        for col, field in enumerate(FIELDS, 1):
            value = record.get(field, '')
            if field == '검사일자':
                try:
                    value = dt.date.fromisoformat(value)
                except (ValueError, TypeError):
                    pass
            if field == '보고서 매수' and text(value).isdigit():
                value = int(value)
            cell = put(sheet, index, col, value)
            if field == '검사일자':
                cell.number_format = 'yyyy-mm-dd'
        for col, field in enumerate(LINKS, 13):
            if field == '보고서':
                put(sheet, index, col, record.get('보고서번호', ''))
            try:
                path = contained_path(root, record.get(field))
                if path and path.exists():
                    label = (record.get('보고서번호') or '번호 미입력 · 보고서 열기') if field == '보고서' else field + ' 열기'
                    link(sheet, index, col, path, label)
            except ValueError:
                pass
        put(sheet, index, 16, '\n'.join(issues(record, root)) or '입력 및 연결 확인 완료')
        put(sheet, index, 17, record.get('출처', ''))
        put(sheet, index, 18, record['id'])
        put(sheet, index, 19, record.get('검사결과', ''))
        mapping = mappings[index - 2]
        put(sheet, index, 20, sum(len(values) for values in mapping.values()) if record['보고서종류'] == 'RT' else '')
        for col, (point, image_index) in enumerate(point_columns, 21):
            values = mapping.get(point, [])
            if point in record.get('촬영포인트', []) and not values:
                put(sheet, index, col, '미연결')
            if image_index < len(values):
                try:
                    target = contained_path(root, values[image_index])
                    if target and target.is_file() and target.is_relative_to(root / 'RT') and target.suffix.lower() in IMAGE_EXTENSIONS:
                        link(sheet, index, col, target, point + (f' ({image_index + 1})' if image_index else ''))
                except ValueError:
                    pass
    for col, options in (('H', 'O,REPAIR'), ('K', ','.join(METHODS))):
        validation = DataValidation(type='list', formula1='"' + options + '"', allow_blank=True)
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(f'{col}2:{col}{len(rows) + 1}')
    pages = DataValidation(type='whole', operator='greaterThanOrEqual', formula1=1, allow_blank=True)
    pages.showErrorMessage = True
    sheet.add_data_validation(pages)
    pages.add(f'L2:L{len(rows) + 1}')
    overview = book.create_sheet('검사내역', 0)
    overview_fields = ('검사일자', '보고서종류', '라인번호', '용접부번호', '용접부특성', '검사결과', '결함종류')
    overview.append((*overview_fields, '보고서번호 (클릭하여 열기)',
                     *(f'RT {p}' + (f' ({i + 1})' if i else '') for p, i in point_columns),
                     'PAUT 데이터 열기', '확인사항', 'RT 영상 연결현황', '검사차수'))
    for row_index, record in enumerate(rows, 2):
        for col, field in enumerate(overview_fields, 1):
            put(overview, row_index, col, record.get(field, ''))
        for target_col, source_col in [(8, 13),
                                       *[(9 + offset, 21 + offset) for offset in range(len(point_columns))],
                                       (9 + len(point_columns), 15)]:
            source_cell = sheet.cell(row_index, source_col)
            if source_col == 13 or source_col >= 21:
                put(overview, row_index, target_col, source_cell.value or '')
            if source_cell.hyperlink:
                cell = put(overview, row_index, target_col, source_cell.value)
                cell.hyperlink = source_cell.hyperlink.target
                cell.style = 'Hyperlink'
        put(overview, row_index, 10 + len(point_columns), sheet.cell(row_index, 16).value)
        expected = record.get('촬영포인트', [])
        mapping = mappings[row_index - 2]
        if record['보고서종류'] == 'RT':
            def valid_image(value):
                try:
                    target = contained_path(root, value)
                    return bool(target and target.is_file() and target.is_relative_to(root / 'RT') and target.suffix.lower() in IMAGE_EXTENSIONS)
                except ValueError:
                    return False
            if expected:
                count = sum(any(valid_image(value) for value in mapping.get(point, [])) for point in expected)
                connection_status = f'{count}/{len(expected)} 포인트 연결'
            else:
                count = sum(any(valid_image(value) for value in values) for values in mapping.values())
                connection_status = f'{count}포인트 연결 / 전체 포인트 미확인'
            put(overview, row_index, 11 + len(point_columns), connection_status)
        put(overview, row_index, 12 + len(point_columns), record.get('검사차수', ''))
    if unmatched:
        diagnostics = book.create_sheet('보고서 대조확인')
        keys = ('검사일자', '보고서종류', '라인번호', '용접부번호', '보고서번호', '대조결과', '출처')
        diagnostics.append(keys)
        for index, row in enumerate(unmatched, 2):
            for col, key in enumerate(keys, 1):
                put(diagnostics, index, col, row.get(key, ''))
    if pending:
        recovery = book.create_sheet('연결 복원보류')
        keys = ('검사일자', '보고서종류', '라인번호', '용접부번호', '용접부특성', '보고서번호', '보고서', '필름영상화', '데이터', 'RT영상', 'id')
        recovery.append(keys)
        for index, row in enumerate(pending, 2):
            for col, key in enumerate(keys, 1):
                value = row.get(key, '')
                put(recovery, index, col, json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value)
    files = book.create_sheet('제출파일목록')
    files.append(('검사종류', '상대경로', '파일 열기'))
    for index, path in enumerate(inventory(root), 2):
        relative = path.relative_to(root)
        put(files, index, 1, relative.parts[0])
        put(files, index, 2, relative.as_posix())
        link(files, index, 3, path, '열기')
    guide = book.create_sheet('작성안내')
    for line in [
        ('항목', '안내'),
        ('폴더 구조', '엑셀을 RT/PAUT/MT/PT 폴더와 같은 최상위 위치에 보관하세요.'),
        ('이동 방법', '메모리카드에 엑셀과 검사 폴더를 함께 복사하고 내부 경로와 파일명을 유지하세요.'),
        ('대조 기준', '검사일자·검사종류·라인번호·용접부번호·용접부특성이 모두 같은 유일한 기록만 병합합니다.'),
        ('기준 목록', '중앙지사 저장 내역만 검사 목록으로 사용합니다. 일치하지 않는 보고서는 보고서 대조확인 시트에 분리합니다.'),
        ('자료 연결', '파일명만으로 확정하지 않습니다. 생성기에서 기록을 선택하고 보고서/영상/데이터를 연결하세요.'),
        ('RT 포인트 영상', 'RT 1-2 / RT 2-3 / RT 3-1 등 포인트 열을 클릭하면 해당 영상 파일이 열립니다. 같은 포인트의 여러 영상도 개별 링크로 표시합니다.'),
        ('RT 영상 수', '연결된 영상 경로 수이며 필수 포인트 수나 검사 완료를 의미하지 않습니다. 보고서의 촬영 포인트와 대조하세요.'),
        ('포인트 연결현황', '중앙지사 구간정보 또는 보고서의 명확한 촬영포인트 목록이 있을 때만 전체 대비 연결 수를 표시합니다. 없으면 전체 포인트 미확인으로 표시합니다.'),
        ('확인사항', '생성 시점의 점검 결과입니다. 엑셀에서 직접 수정해도 자동 갱신되지 않습니다. 생성기에서 재생성하세요.'),
        ('결함종류', '미입력은 무결함을 의미하지 않습니다. 원본 보고서의 코드와 무결함 표기 기준을 확인하세요.'),
        ('원본 대조', '과업내용서 별표1의 누락 항목 및 발주처 지정 항목·파일명은 제출 전에 확인하세요.'),
        ('전용 데이터', 'PAUT 데이터는 해당 프로그램이 설치된 PC에서 열 수 있습니다.'),
        ('생성일시', dt.datetime.now().isoformat(timespec='seconds')),
    ]:
        guide.append(line)
    for ws in book:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='17365D')
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 22
    sheet.column_dimensions['P'].width = 65
    sheet.column_dimensions['Q'].width = 40
    files.column_dimensions['B'].width = 90
    guide.column_dimensions['B'].width = 100
    overview.column_dimensions['C'].width = 38
    overview.column_dimensions[get_column_letter(10 + len(point_columns))].width = 65
    # Atomic creation prevents leaving a half-written deliverable.
    descriptor, temporary_name = tempfile.mkstemp(prefix='.attributes_', suffix='.xlsx', dir=root)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        book.save(temporary)
        check = load_workbook(temporary, read_only=False)
        try:
            if check['속성정보'].max_row != len(rows) + 1:
                raise ValueError('저장 후 행 수 검증 실패')
        finally:
            check.close()
        if destination.exists():
            raise FileExistsError(str(destination))
        temporary.rename(destination)
    finally:
        book.close()
        temporary.unlink(missing_ok=True)
    return destination


def main():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
    from site_apps.central.src.data_sync import get_history_path, get_submission_dir

    window = tk.Tk()
    window.title('중앙지사 속성정보 엑셀 생성')
    window.geometry('1250x760')
    state = {'history': [], 'reports': [], 'rows': [], 'unmatched': []}
    root_var = tk.StringVar(value=get_submission_dir())
    history_var = tk.StringVar(value=get_history_path())
    status = tk.StringVar(value='중앙지사 이력을 불러온 뒤 보고서 앱에서 저장한 JSON 또는 XLSX를 추가하세요.')
    form = ttk.Frame(window, padding=10)
    form.pack(fill='both', expand=True)

    def run(action):
        try:
            action()
        except Exception as error:
            messagebox.showerror('확인 필요', str(error), parent=window)

    def choose_root():
        path = filedialog.askdirectory(title='메모리카드 또는 제출용 최상위 폴더 선택')
        if path:
            persist()
            root_var.set(path)
            rebuild(preserve=False)
            connect_files()

    def choose_history():
        path = filedialog.askopenfilename(title='중앙지사 daily_work_history.json', filetypes=[('JSON', '*.json')])
        if path:
            history_var.set(path)

    for index, (label, variable, command) in enumerate([
        ('제출 폴더', root_var, choose_root), ('중앙지사 이력', history_var, choose_history)
    ]):
        ttk.Label(form, text=label).grid(row=index, column=0, sticky='w')
        ttk.Entry(form, textvariable=variable, state='readonly' if index == 0 else 'normal').grid(row=index, column=1, sticky='ew', padx=6, pady=3)
        ttk.Button(form, text='선택', command=command).grid(row=index, column=2)
    form.columnconfigure(1, weight=1)
    toolbar = ttk.Frame(form)
    toolbar.grid(row=2, column=0, columnspan=3, sticky='ew', pady=8)

    def refresh():
        tree.delete(*tree.get_children())
        for row in state['rows']:
            count = len(issues(row, root_var.get())) if root_var.get() else '폴더 선택'
            tree.insert('', 'end', iid=row['id'], values=tuple(row.get(k, '') for k in visible) + (str(count),))
        pending = pending_saved()
        status.set(f"중앙지사 검사 {len(state['rows'])}건 / 보고서 대조확인 {len(state['unmatched'])}건 / 연결 복원보류 {len(pending)}건. 보류 자료는 자동 연결하지 않고 보존합니다.")

    def pending_saved():
        if not root_var.get():
            return []
        path = Path(root_var.get()) / CACHE_NAME
        if not path.exists():
            return []
        read_connections(path)
        with path.open(encoding='utf-8-sig') as stream:
            return json.load(stream).get('unresolved', [])

    def persist():
        if state['rows'] and root_var.get():
            path = Path(root_var.get()) / CACHE_NAME
            if path.exists():
                read_connections(path)
            write_connections(path, state['rows'], root_var.get())

    def rebuild(preserve=True):
        base = copy.deepcopy(state['history'])
        if root_var.get():
            cache = Path(root_var.get()) / CACHE_NAME
            if cache.exists():
                base = restore_connections(base, read_connections(cache))
        if preserve:
            base = restore_connections(base, state['rows'])
        unmatched = []
        state['rows'] = merge_rows(base, state['reports'], unmatched)
        state['unmatched'] = unmatched
        refresh()

    def load():
        loaded = load_history(history_var.get())
        persist()
        state['history'] = loaded
        rebuild()

    method = tk.StringVar(value='RT')
    def add_report():
        if not state['history']:
            raise ValueError('중앙지사 이력을 먼저 불러오세요. 보고서 데이터는 일치하는 기록만 보완합니다.')
        path = filedialog.askopenfilename(title='보고서 앱 데이터 저장 JSON / 표 형식 XLSX', filetypes=[('보고서 데이터', '*.json *.xlsx')])
        if path:
            loaded = load_report_data(path, method.get())
            existing = {row['id'] for row in state['reports']}
            state['reports'].extend(row for row in loaded if row['id'] not in existing)
            rebuild()
            persist()
            if state['unmatched']:
                messagebox.showinfo('보고서 대조확인', f"{len(state['unmatched'])}건은 중앙지사 기록에 병합하지 않았습니다. 엑셀의 '보고서 대조확인' 시트에서 확인할 수 있습니다.")

    ttk.Button(toolbar, text='중앙지사 이력 불러오기', command=lambda: run(load)).pack(side='left')
    ttk.Combobox(toolbar, textvariable=method, values=METHODS, state='readonly', width=6).pack(side='left', padx=5)
    ttk.Button(toolbar, text='보고서 데이터 추가', command=lambda: run(add_report)).pack(side='left')
    def make_folders():
        if not root_var.get():
            choose_root()
        if not root_var.get():
            return
        count = create_submission_folders(state['rows'], root_var.get())
        messagebox.showinfo('폴더 생성 완료', f'{count}개 검사 기록의 제출 폴더를 만들었습니다.\n보고서·영상·데이터를 넣은 후 자동 연결을 누르세요. 기존 파일은 변경하지 않습니다.')

    def connect_files():
        if not root_var.get() or not state['rows']:
            return
        state['rows'] = auto_connect(state['rows'], root_var.get())
        persist()
        refresh()

    ttk.Button(toolbar, text='제출 폴더 틀 생성', command=lambda: run(make_folders)).pack(side='left', padx=6)
    ttk.Button(toolbar, text='자료 자동 연결', command=lambda: run(connect_files)).pack(side='left')
    ttk.Label(form, text='보고서 프로그램의 「데이터 저장(JSON)」 또는 미리보기 엑셀 내보내기를 사용하세요. 완성 보고서 원본은 행 편집에서 연결합니다.').grid(row=3, column=0, columnspan=3, sticky='w')
    visible = ('검사일자', '보고서종류', '라인번호', '용접부번호', '용접부특성', '검사결과')
    table = ttk.Frame(form)
    table.grid(row=4, column=0, columnspan=3, sticky='nsew', pady=8)
    form.rowconfigure(4, weight=1)
    tree = ttk.Treeview(table, columns=(*visible, '확인 항목 수'), show='headings')
    for field in (*visible, '확인 항목 수'):
        tree.heading(field, text=field)
        tree.column(field, width=145)
    scroll = ttk.Scrollbar(table, orient='vertical', command=tree.yview)
    horizontal = ttk.Scrollbar(table, orient='horizontal', command=tree.xview)
    tree.configure(yscrollcommand=scroll.set, xscrollcommand=horizontal.set)
    tree.grid(row=0, column=0, sticky='nsew')
    scroll.grid(row=0, column=1, sticky='ns')
    horizontal.grid(row=1, column=0, sticky='ew')
    table.rowconfigure(0, weight=1)
    table.columnconfigure(0, weight=1)

    def edit(_event=None):
        selected = tree.selection()
        if not selected:
            return
        if not root_var.get():
            raise ValueError('제출 폴더를 먼저 선택하세요.')
        row = next(row for row in state['rows'] if row['id'] == selected[0])
        dialog = tk.Toplevel(window)
        dialog.title('속성정보 / 원본 자료 연결')
        dialog.transient(window)
        dialog.grab_set()
        variables = {field: tk.StringVar(value=row.get(field, '')) for field in (*FIELDS, *LINKS, '검사결과', '보고서번호')}
        for index, (field, variable) in enumerate(variables.items()):
            ttk.Label(dialog, text=field).grid(row=index, column=0, sticky='w', padx=6)
            original = row.get('원본값', {}).get(field)
            ttk.Entry(dialog, textvariable=variable, width=85, state='readonly' if original else 'normal').grid(row=index, column=1, padx=6, pady=2)
            if field in LINKS:
                def browse(field=field):
                    path = filedialog.askopenfilename(initialdir=root_var.get(), parent=dialog)
                    if path:
                        variables[field].set(contained_path(root_var.get(), path).relative_to(Path(root_var.get()).resolve()).as_posix())
                def folder(field=field):
                    path = filedialog.askdirectory(initialdir=root_var.get(), parent=dialog)
                    if path:
                        variables[field].set(contained_path(root_var.get(), path).relative_to(Path(root_var.get()).resolve()).as_posix())
                ttk.Button(dialog, text='파일', command=lambda f=browse: run(f)).grid(row=index, column=2)
                ttk.Button(dialog, text='폴더', command=lambda f=folder: run(f)).grid(row=index, column=3)
        def save():
            for field in LINKS:
                contained_path(root_var.get(), variables[field].get())
            if variables['필름영상화'].get().strip() != row.get('필름영상화', ''):
                row['자동연결영상'] = False
            row.update({field: var.get().strip() for field, var in variables.items()})
            persist()
            dialog.destroy()
            refresh()
        if row['보고서종류'] == 'RT':
            def edit_images():
                popup = tk.Toplevel(dialog)
                popup.title('RT 포인트별 영상 연결')
                popup.transient(dialog)
                popup.grab_set()
                ttk.Label(popup, text='한 줄에 포인트=상대경로 입력. 같은 포인트의 영상은 줄을 추가하세요.').pack(padx=8, pady=5)
                editor = tk.Text(popup, width=100, height=15)
                editor.pack(padx=8)
                current = point_images({**row, '필름영상화': variables['필름영상화'].get()}, root_var.get())
                editor.insert('1.0', '\n'.join(f'{p}={v}' for p, values in current.items() for v in values))
                point_var = tk.StringVar(value='1-2')
                controls = ttk.Frame(popup)
                controls.pack(fill='x', padx=8, pady=6)
                ttk.Label(controls, text='포인트').pack(side='left')
                ttk.Entry(controls, textvariable=point_var, width=10).pack(side='left')
                def add_images():
                    if not re.fullmatch(r'\d+-\d+', point_var.get().strip()):
                        raise ValueError('포인트를 1-2 형태로 입력하세요.')
                    paths = filedialog.askopenfilenames(parent=popup, initialdir=root_var.get(), title='이 포인트의 영상 선택')
                    additions = [contained_path(root_var.get(), p).relative_to(Path(root_var.get()).resolve()).as_posix() for p in paths]
                    for value in additions:
                        editor.insert('end', '\n' + point_var.get().strip() + '=' + value)
                def apply_images():
                    mapping = {}
                    for line in editor.get('1.0', 'end').splitlines():
                        if not line.strip():
                            continue
                        point, separator, value = line.partition('=')
                        if not separator:
                            raise ValueError('포인트=상대경로 형식으로 입력하세요.')
                        target = contained_path(root_var.get(), value.strip())
                        if not target or not target.is_file() or not target.is_relative_to(Path(root_var.get()).resolve() / 'RT'):
                            raise ValueError('RT 폴더 안의 실제 영상 파일을 선택하세요.')
                        mapping.setdefault(point.strip(), []).append(target.relative_to(Path(root_var.get()).resolve()).as_posix())
                    point_images({**row, 'RT영상': mapping}, root_var.get())
                    row['RT영상'] = mapping
                    row['자동연결영상'] = False
                    persist()
                    popup.destroy()
                    dialog.grab_set()
                ttk.Button(controls, text='영상 파일 추가', command=lambda: run(add_images)).pack(side='left', padx=5)
                ttk.Button(controls, text='포인트 연결 적용', command=lambda: run(apply_images)).pack(side='right')
            ttk.Button(dialog, text='RT 포인트별 영상 연결 (1-2 / 2-3 / 3-1 …)', command=lambda: run(edit_images)).grid(row=17, column=1, pady=4)
        ttk.Label(dialog, text='\n'.join(issues(row, root_var.get())), wraplength=850).grid(row=18, column=0, columnspan=4, sticky='w', padx=6)
        ttk.Button(dialog, text='적용', command=lambda: run(save)).grid(row=19, column=1, pady=8)
    tree.bind('<Double-1>', lambda event: run(lambda: edit(event)))

    def open_selected():
        selected = tree.selection()
        if not selected:
            return
        row = next(row for row in state['rows'] if row['id'] == selected[0])
        target = contained_path(root_var.get(), row.get('보고서'))
        if target is None or not target.exists():
            raise ValueError('보고서를 먼저 연결하세요.')
        os.startfile(target)

    def common_fields():
        if not state['rows']:
            raise ValueError('검사 기록을 먼저 불러오세요.')
        dialog = tk.Toplevel(window)
        dialog.title('공통정보 일괄 입력')
        dialog.transient(window)
        dialog.grab_set()
        variables = {}
        for index, field in enumerate(('공사명', '시공사', '하도사', '용역사')):
            variables[field] = tk.StringVar()
            ttk.Label(dialog, text=field).grid(row=index, column=0, padx=8, pady=4)
            ttk.Entry(dialog, textvariable=variables[field], width=55).grid(row=index, column=1, padx=8)
        ttk.Label(dialog, text='입력한 값은 전체 기록의 빈 칸에만 적용합니다. 기존 값은 유지합니다.').grid(row=4, column=0, columnspan=2, padx=8)
        def apply():
            for row in state['rows']:
                for field, variable in variables.items():
                    if not row.get(field) and variable.get().strip():
                        row[field] = variable.get().strip()
            persist()
            dialog.destroy()
            refresh()
        ttk.Button(dialog, text='빈 칸에 적용', command=apply).grid(row=5, column=1, pady=8)

    def save_connections():
        path = filedialog.asksaveasfilename(title='편집 내용 및 연결정보 저장', defaultextension='.json', initialfile='속성정보_연결정보.json', filetypes=[('JSON', '*.json')])
        if path:
            if not state['rows']:
                raise ValueError('저장할 검사 기록이 없습니다.')
            write_connections(path, state['rows'], root_var.get())
            persist()

    def load_connections():
        if not state['history']:
            raise ValueError('중앙지사 이력을 먼저 불러오세요.')
        path = filedialog.askopenfilename(title='저장한 연결정보 불러오기', filetypes=[('JSON', '*.json')])
        if not path:
            return
        rows = read_connections(path)
        state['rows'] = restore_connections(state['rows'], rows)
        persist()
        refresh()

    def export():
        if not root_var.get():
            raise ValueError('제출 폴더를 선택하세요.')
        count = sum(bool(issues(row, root_var.get())) for row in state['rows'])
        if count and not messagebox.askyesno('검토용 엑셀 생성', f'{count}건에 확인사항이 있습니다. 확인사항을 포함한 검토용 엑셀을 생성할까요?'):
            return
        path = filedialog.asksaveasfilename(initialdir=root_var.get(), initialfile='중앙지사_속성정보.xlsx', defaultextension='.xlsx', filetypes=[('Excel', '*.xlsx')])
        if path:
            persist()
            export_workbook(state['rows'], root_var.get(), path, state['unmatched'], pending_saved())
            messagebox.showinfo('생성 완료', f'{len(state["rows"])}건 저장\n확인사항 있는 기록: {count}건\n{path}')

    bottom = ttk.Frame(form)
    bottom.grid(row=5, column=0, columnspan=3, sticky='ew')
    for label, action in [('공통정보 입력', common_fields), ('선택 기록 편집 / 자료 연결', edit), ('선택 보고서 열기', open_selected),
                          ('연결정보 저장', save_connections), ('연결정보 불러오기', load_connections), ('엑셀 생성', export)]:
        ttk.Button(bottom, text=label, command=lambda a=action: run(a)).pack(side='left', padx=3)
    ttk.Label(form, textvariable=status, wraplength=1150).grid(row=6, column=0, columnspan=3, sticky='w', pady=8)
    def close():
        try:
            persist()
        except Exception as error:
            messagebox.showerror('연결정보 저장 실패', str(error), parent=window)
            return
        window.destroy()
    window.protocol('WM_DELETE_WINDOW', close)
    if Path(history_var.get()).is_file():
        window.after(100, lambda: run(load))
    window.mainloop()


if __name__ == '__main__':
    main()
