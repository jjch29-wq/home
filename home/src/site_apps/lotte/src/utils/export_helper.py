import os
import datetime
import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# [MULTI-EQUIP] 장비명 다중 등록 구분자 (예: "UT장비 (SN1), MT장비 (SN2)")
EQUIP_SEPARATOR = ", "

# K1 PJT 공식 누계기성(제8회 기성내역서 '26.10) 검사방법별 (최종 수량, 금액)
K1_OFFICIAL_CUMULATIVE = {
    'RT': (10064.0, 256632000.0),
    'PAUT': (1986.0, 36741000.0),
    'PT': (6170.0, 25913160.0),
    'PMI': (670.0, 9581000.0),
    'UT': (88.0, 818400.0),
}


def split_equipments(value):
    """장비명 문자열을 개별 장비 리스트로 분리.
    괄호 안의 쉼표(SN 등)는 분리하지 않으며, 빈 값/nan/중복은 제거한다."""
    if value is None:
        return []
    s = str(value).strip()
    if not s or s.lower() in ('nan', 'none', '-'):
        return []
    parts, buf, depth = [], [], 0
    for ch in s:
        if ch in '([':
            depth += 1
        elif ch in ')]':
            depth = max(0, depth - 1)
        if ch == ',' and depth == 0:
            parts.append(''.join(buf))
            buf = []
        else:
            buf.append(ch)
    parts.append(''.join(buf))
    result = []
    for p in parts:
        p = p.strip()
        if p and p.lower() not in ('nan', 'none', '-') and p not in result:
            result.append(p)
    return result


def join_equipments(items):
    """장비 리스트(또는 다중 장비 문자열 포함)를 중복 없이 하나의 문자열로 결합."""
    result = []
    for it in items or []:
        for e in split_equipments(it):
            if e not in result:
                result.append(e)
    return EQUIP_SEPARATOR.join(result)

# Common Styles
TITLE_FONT = Font(name='맑은 고딕', size=16, bold=True)
HEADER_FONT = Font(name='맑은 고딕', size=11, bold=True)
BOLD_FONT = Font(name='맑은 고딕', size=10, bold=True)
NORMAL_FONT = Font(name='맑은 고딕', size=10)
INFO_FONT = Font(name='맑은 고딕', size=11, bold=True)

CENTER_ALIGN = Alignment(horizontal='center', vertical='center')
RIGHT_ALIGN = Alignment(horizontal='right', vertical='center')
LEFT_ALIGN = Alignment(horizontal='left', vertical='center')

THIN_BORDER = Side(border_style="thin", color="000000")
BORDER_ALL = Border(top=THIN_BORDER, left=THIN_BORDER, right=THIN_BORDER, bottom=THIN_BORDER)

HEADER_FILL = PatternFill(start_color='E0E0E0', end_color='E0E0E0', fill_type='solid')
SUB_FILL = PatternFill(start_color='F5F5F5', end_color='F5F5F5', fill_type='solid')
YELLOW_FILL = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
BLUE_FILL = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
GREEN_FILL = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
ORANGE_FILL = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')

def _to_f(val):
    try:
        if val is None or val == '':
            return 0.0
        s = str(val).replace(',', '').replace('₩', '').replace('%', '').strip()
        if '(' in s:
            s = s.split('(')[0].strip()
        return float(s or 0)
    except:
        return 0.0

def _write_table(ws, start_row, title, headers, data_rows, total_label=None, total_val=None, theme_fill=HEADER_FILL):
    """Generic table writer with formatting and optional total row"""
    ws.cell(row=start_row, column=1, value=title).font = HEADER_FONT
    ws.cell(row=start_row, column=1).fill = theme_fill
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=len(headers))
    for col in range(1, len(headers)+1):
        ws.cell(row=start_row, column=col).border = BORDER_ALL
        ws.cell(row=start_row, column=col).alignment = CENTER_ALIGN
        
    r = start_row + 1
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=r, column=col, value=h)
        c.font = BOLD_FONT
        c.fill = SUB_FILL
        c.border = BORDER_ALL
        c.alignment = CENTER_ALIGN
            
    r += 1
    for row_data in data_rows:
        for col, val in enumerate(row_data, 1):
            c = ws.cell(row=r, column=col, value=val)
            c.font = NORMAL_FONT
            c.border = BORDER_ALL
            if isinstance(val, (int, float)):
                c.number_format = '#,##0'
                c.alignment = RIGHT_ALIGN
            elif isinstance(val, str) and val.replace(',', '').replace('.', '').replace('-', '').isdigit():
                c.number_format = '#,##0'
                c.alignment = RIGHT_ALIGN
                try: c.value = float(val.replace(',', ''))
                except: pass
            else:
                c.alignment = CENTER_ALIGN
        r += 1
        
    if total_label is not None:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers)-1)
        c = ws.cell(row=r, column=1, value=total_label)
        c.font = BOLD_FONT
        c.fill = YELLOW_FILL
        c.border = BORDER_ALL
        c.alignment = CENTER_ALIGN
        for col in range(2, len(headers)):
            ws.cell(row=r, column=col).border = BORDER_ALL
            
        c_tot = ws.cell(row=r, column=len(headers), value=total_val)
        c_tot.font = BOLD_FONT
        c_tot.fill = YELLOW_FILL
        c_tot.border = BORDER_ALL
        c_tot.alignment = RIGHT_ALIGN
        c_tot.number_format = '#,##0'
        try: c_tot.value = float(str(total_val).replace(',', '').replace('₩', '').strip())
        except: pass
        r += 1
        
    return r + 1

def _build_budget_sheet(ws, summary_data, labor_data, material_data, expense_data, mode='planned'):
    """Build Planned or Actual summary sheet"""
    if mode == 'actual':
        theme_fill = GREEN_FILL
        cost_header = "사후원가가액"
        title_text = f"[{summary_data.get('site', '')}] 공사실행예산서 (사후원가 실적)"
        summary_title = "1) 사후원가(실적) 요약"
        suffix = " (실적)"
    else:
        theme_fill = BLUE_FILL
        cost_header = "사전원가가액"
        title_text = f"[{summary_data.get('site', '')}] 공사실행예산서 (사정원가)"
        summary_title = "1) 사전예산 요약"
        suffix = ""

    # Title
    ws.merge_cells('A1:G1')
    c = ws['A1']
    c.value = title_text
    c.font = Font(name='맑은 고딕', size=18, bold=True)
    c.alignment = CENTER_ALIGN
    ws.row_dimensions[1].height = 30
    
    # Metadata
    ws.merge_cells('A2:B2')
    ws['A2'] = "출력일시: " + datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    ws['A2'].font = NORMAL_FONT
    ws['A2'].alignment = LEFT_ALIGN

    # Approval Box
    approval_headers = ["소장", "담당", "상무", "본부장", "사장"]
    for i, h in enumerate(approval_headers):
        col_letter = get_column_letter(3 + i)
        h_cell = ws[f"{col_letter}2"]
        h_cell.value = h
        h_cell.font = BOLD_FONT
        h_cell.alignment = CENTER_ALIGN
        h_cell.border = BORDER_ALL
        
        s_cell = ws[f"{col_letter}3"]
        s_cell.border = BORDER_ALL
        s_cell.alignment = CENTER_ALIGN
        
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 60

    # Set Column Widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 12

    # Print Setup
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 9

    # Site Info
    site_info_row = 5
    ws.cell(row=site_info_row, column=1, value="현 장 명 :").font = INFO_FONT
    ws.cell(row=site_info_row, column=2, value=f"{summary_data.get('site', '')} 중 비파괴검사").font = INFO_FONT
    
    ws.cell(row=site_info_row+1, column=1, value="공사기간 :").font = INFO_FONT
    period_str = summary_data.get('period', '')
    ws.cell(row=site_info_row+1, column=2, value=period_str).font = INFO_FONT
    
    days_diff_str = ""
    if period_str and "~" in period_str:
        try:
            start_str, end_str = period_str.split("~")
            start_date = datetime.datetime.strptime(start_str.strip(), "%Y.%m.%d")
            end_date = datetime.datetime.strptime(end_str.strip(), "%Y.%m.%d")
            days_diff_str = f"{(end_date - start_date).days + 1}"
        except: pass

    ws.cell(row=site_info_row+1, column=5, value=days_diff_str).font = INFO_FONT
    ws.cell(row=site_info_row+1, column=5).alignment = RIGHT_ALIGN
    ws.cell(row=site_info_row+1, column=6, value="일").font = INFO_FONT
    ws.cell(row=site_info_row+1, column=6).alignment = LEFT_ALIGN
    
    ws.cell(row=site_info_row+2, column=1, value="계약업체 :").font = INFO_FONT
    ws.cell(row=site_info_row+2, column=2, value="").font = INFO_FONT
    ws.cell(row=site_info_row+3, column=1, value="경인사무소 :").font = INFO_FONT
    ws.cell(row=site_info_row+3, column=2, value="").font = INFO_FONT

    current_row = site_info_row + 5

    # 1. Summary Table
    ws.cell(row=current_row, column=1, value=summary_title).font = HEADER_FONT
    ws.cell(row=current_row, column=1).fill = theme_fill
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
    for col in range(1, 3):
        ws.cell(row=current_row, column=col).border = BORDER_ALL
        ws.cell(row=current_row, column=col).alignment = CENTER_ALIGN
    current_row += 1

    profit_raw = str(summary_data.get('profit', 0)).strip()
    margin_raw = str(summary_data.get('margin', '0%')).strip()
    if '(' in profit_raw and ')' in profit_raw:
        p_part, m_part = profit_raw.split('(')
        profit_val = p_part.strip()
        if not margin_raw or margin_raw in ('0%', '0.0%', '0', 'nan'):
            margin_val = m_part.replace(')', '').strip()
        else: margin_val = margin_raw
    else:
        profit_val = profit_raw
        margin_val = margin_raw

    if '%' not in str(margin_val):
        try: margin_val = f"{float(str(margin_val).replace(',', '').strip() or 0):.1f}%"
        except: margin_val = f"{margin_val}%"

    summary_items = [
        ("계약금액(Revenue)", summary_data.get('revenue', 0)),
        ("매출금액(UnitPrice)", summary_data.get('unitprice', 0)),
        ("실행 노무비(Labor)", summary_data.get('labor', 0)),
        ("실행 재료비(Material)", summary_data.get('material', 0)),
        ("실행 경비(Expense)", summary_data.get('expense', 0)),
        ("실행 외주비(Outsource)", summary_data.get('outsource', 0)),
        ("영업이익(Profit)", profit_val),
        ("이익률(Margin)", margin_val)
    ]
    
    for k, v in summary_items:
        ws.cell(row=current_row, column=1, value=k).font = BOLD_FONT
        ws.cell(row=current_row, column=1).fill = SUB_FILL
        ws.cell(row=current_row, column=1).border = BORDER_ALL
        
        c_val = ws.cell(row=current_row, column=2, value=v)
        c_val.font = NORMAL_FONT
        c_val.border = BORDER_ALL
        if '%' in str(v):
            c_val.alignment = RIGHT_ALIGN
        else:
            c_val.number_format = '#,##0'
            c_val.alignment = RIGHT_ALIGN
            try: c_val.value = float(str(v).replace(',', '').replace('₩', '').strip())
            except: pass
        current_row += 1
        
    current_row += 2

    # 2. Labor Data
    if labor_data:
        reg_rows = []
        reg_total = 0
        regular_keys = ['이사', '부장', '차장', '과장', '대리', '사원', '초급', '중급', '고급', '특급']
        for rank in labor_data:
            if rank not in regular_keys and rank not in ['연장근무', '야간근무', '휴일근무']:
                regular_keys.append(rank)

        for rank in regular_keys:
            if rank in labor_data:
                vals = labor_data[rank]
                p = _to_f(vals.get('personnel'))
                d = _to_f(vals.get('period'))
                u = _to_f(vals.get('unit_price'))
                amt = p * d * u
                if p > 0 or d > 0:
                    d_disp = round(d, 4) if d % 1 else int(d)
                    p_disp = round(p, 2) if p % 1 else int(p)
                    reg_rows.append([rank, p_disp, d_disp, u, amt])
                    reg_total += amt
        
        if reg_rows:
            current_row = _write_table(
                ws, current_row, f"2) 노무비 - 정시근무{suffix}", 
                ["직급", "투입인원(명)", "투입기간(일)", "단가", cost_header], 
                reg_rows, f"정시근무 합계{suffix}", reg_total, theme_fill=HEADER_FILL
            )

        sp_rows = []
        sp_total = 0
        special_keys = ['연장근무', '야간근무', '휴일근무']
        for stype in special_keys:
            if stype in labor_data:
                vals = labor_data[stype]
                p = _to_f(vals.get('personnel'))
                h = _to_f(vals.get('period'))
                u = _to_f(vals.get('unit_price'))
                amt = p * h * u
                if p > 0 or h > 0:
                    h_disp = round(h, 4) if h % 1 else int(h)
                    p_disp = round(p, 2) if p % 1 else int(p)
                    sp_rows.append([stype, p_disp, h_disp, u, amt])
                    sp_total += amt
                
        if sp_rows:
            current_row = _write_table(
                ws, current_row, f"2) 노무비 - 특별근무{suffix}", 
                ["형태", "투입인원(명)", "투입시간(시간)", "단가", cost_header], 
                sp_rows, f"특별근무 합계{suffix}", sp_total, theme_fill=HEADER_FILL
            )

    # 3. Material Data
    if material_data:
        mat_rows = []
        mat_total = 0
        for item in material_data:
            qty = _to_f(item.get('qty'))
            price = _to_f(item.get('price'))
            amt = qty * price
            if qty > 0:
                mat_rows.append([item.get('name', ''), item.get('spec', ''), qty, item.get('unit', ''), price, amt])
                mat_total += amt
        if mat_rows:
            current_row = _write_table(
                ws, current_row, f"3) 재료비{suffix}", 
                ["품목", "사양", "수량", "규격", "단가", cost_header], 
                mat_rows, f"재료비 합계{suffix}", mat_total, theme_fill=HEADER_FILL
            )
            
    # 4. Expense Data
    if expense_data:
        # Site Expense
        exp_rows = []
        exp_total = 0
        for item in expense_data.get('site_expense', []):
            cont = item.get('cont', '')
            ppl = _to_f(item.get('ppl'))
            qty = _to_f(item.get('qty'))
            price = _to_f(item.get('price'))
            amt = _to_f(item.get('amount'))
            if amt == 0 and (qty > 0 or price > 0):
                amt = qty * price
            if qty > 0 or ppl > 0 or amt > 0:
                exp_rows.append([cont, ppl, qty, price, amt])
                exp_total += amt
        if exp_rows:
            current_row = _write_table(
                ws, current_row, f"4) 경비 - 현장경비{suffix}", 
                ["항목(내용)", "투입인원", "수량/일수", "단가", cost_header], 
                exp_rows, "현장경비 합계", exp_total, theme_fill=HEADER_FILL
            )
            
        # Rental
        rent_rows = []
        rent_total = 0
        for item in expense_data.get('rental', []):
            spec = item.get('spec', '')
            qty = _to_f(item.get('qty'))
            period = _to_f(item.get('period'))
            price = _to_f(item.get('price'))
            amt = _to_f(item.get('amount'))
            if amt == 0 and (qty > 0 or period > 0 or price > 0):
                amt = qty * period * price
            if qty > 0 or period > 0 or amt > 0:
                rent_rows.append([spec, qty, period, price, amt])
                rent_total += amt
        if rent_rows:
            current_row = _write_table(
                ws, current_row, f"4) 경비 - 렌탈비{suffix}", 
                ["사양", "수량", "기간(일수)", "단가", cost_header], 
                rent_rows, "렌탈비 합계", rent_total, theme_fill=HEADER_FILL
            )
            
        # Outsource
        out_rows = []
        out_total = 0
        for item in expense_data.get('outsource', []):
            work = item.get('work', '')
            count = _to_f(item.get('count'))
            price = _to_f(item.get('price'))
            amt = _to_f(item.get('amount'))
            if amt == 0 and (count > 0 or price > 0):
                amt = count * price
            if count > 0 or amt > 0:
                out_rows.append([work, count, "", price, amt])
                out_total += amt
        if out_rows:
            current_row = _write_table(
                ws, current_row, f"4) 경비 - 외주비{suffix}", 
                ["작업내용", "인원/횟수", "비고", "단가", cost_header], 
                out_rows, "외주비 합계", out_total, theme_fill=HEADER_FILL
            )
            
        # Depreciation
        dep_rows = []
        dep_total = 0
        for item in expense_data.get('depreciation', []):
            name = item.get('item', '')
            qty = _to_f(item.get('qty'))
            days = _to_f(item.get('days'))
            rate = _to_f(item.get('rate'))
            amt = _to_f(item.get('amount'))
            if amt == 0 and (qty > 0 or days > 0 or rate > 0):
                amt = qty * days * rate
            if qty > 0 or days > 0 or amt > 0:
                dep_rows.append([name, qty, days, rate, amt])
                dep_total += amt
        if dep_rows:
            current_row = _write_table(
                ws, current_row, f"4) 경비 - 감가상각{suffix}", 
                ["장비명", "대수", "사용일수", "상각율", cost_header], 
                dep_rows, "감가상각 합계", dep_total, theme_fill=HEADER_FILL
            )

    return current_row


def _build_proof_inspection_sheet(ws, site, raw_df):
    """Sheet 3: Detailed inspection logs proving revenue/test fee"""
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 9

    # Set Column Widths
    col_widths = {
        'A': 6, 'B': 13, 'C': 11, 'D': 18, 'E': 14, 
        'F': 15, 'G': 25, 'H': 15, 'I': 12, 'J': 11,
        'K': 11, 'L': 15, 'M': 20
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # Title
    ws.merge_cells('A1:M1')
    c = ws['A1']
    c.value = f"[{site}] 사후원가 증빙 1 - 검사 실적 및 매출 증빙 내역 (방법별 요약)"
    c.font = Font(name='맑은 고딕', size=16, bold=True)
    c.alignment = CENTER_ALIGN
    ws.row_dimensions[1].height = 28

    ws.merge_cells('A2:M2')
    ws['A2'] = f"출력일시: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')} | 총 검사 건수: {len(raw_df):,}건"
    ws['A2'].font = NORMAL_FONT
    ws['A2'].alignment = LEFT_ALIGN

    # --- 1. 검사방법별 집계 요약 (Summary by Inspection Method) ---
    method_dict = {}
    for _, row in raw_df.iterrows():
        m = str(row.get('검사방법', '')).strip()
        if not m or m.lower() in ('nan', 'none'):
            m = '기타'
        if m not in method_dict:
            method_dict[m] = {
                'count': 0, 'qty': 0.0, 'fee': 0.0, 
                'travel': 0.0, 'meal': 0.0, 'equips': set()
            }
        q = _to_f(row.get('검사량', 0) or row.get('수량', 0))
        f = _to_f(row.get('검사비', 0))
        tr = _to_f(row.get('출장비', 0))
        ml = _to_f(row.get('일식', 0))
        method_dict[m]['count'] += 1
        method_dict[m]['qty'] += q
        method_dict[m]['fee'] += f
        method_dict[m]['travel'] += tr
        method_dict[m]['meal'] += ml
        # [MULTI-EQUIP] 다중 장비는 개별로 집계
        for eq in split_equipments(row.get('장비명', '')):
            method_dict[m]['equips'].add(eq)

    # K1 PJT는 건수는 실적 그대로 두고, 검사수량·검사금액은 공식 누계기성(최종) 값으로 보정한다.
    if str(site).upper().replace('-', '').replace(' ', '') == 'K1PJT':
        for m_key, (off_qty, off_fee) in K1_OFFICIAL_CUMULATIVE.items():
            if m_key in method_dict:
                method_dict[m_key]['qty'] = off_qty
                method_dict[m_key]['fee'] = off_fee

    total_sum_fee = sum(v['fee'] + v['travel'] + v['meal'] for v in method_dict.values())

    r = 4
    headers_summary = ["No", "검사방법", "건수(건)", "총 검사수량", "검사금액(원)", "출장비(원)", "일식(원)", "총 매출액(원)", "매출 비중(%)", "주요 투입장비 및 비고"]
    ws.cell(row=r, column=1, value="1) 검사방법별 실적 요약 집계 (Summary by Inspection Method)").font = HEADER_FONT
    ws.cell(row=r, column=1).fill = BLUE_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=13)
    for c_i in range(1, 14):
        ws.cell(row=r, column=c_i).border = BORDER_ALL

    r += 1
    # Summary Table Headers
    for c_i, h in enumerate(headers_summary[:9], 1):
        cell = ws.cell(row=r, column=c_i, value=h)
        cell.font = BOLD_FONT
        cell.fill = SUB_FILL
        cell.border = BORDER_ALL
        cell.alignment = CENTER_ALIGN
    # Merge J~M for note/equipment
    ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=13)
    c_note_h = ws.cell(row=r, column=10, value=headers_summary[9])
    c_note_h.font = BOLD_FONT
    c_note_h.fill = SUB_FILL
    c_note_h.alignment = CENTER_ALIGN
    for c_i in range(10, 14):
        ws.cell(row=r, column=c_i).border = BORDER_ALL
        ws.cell(row=r, column=c_i).fill = SUB_FILL

    # Populate Summary Rows (sorted by total revenue descending)
    sorted_methods = sorted(method_dict.items(), key=lambda x: (x[1]['fee'] + x[1]['travel'] + x[1]['meal']), reverse=True)
    r += 1
    sum_tot_cnt = 0
    sum_tot_qty = 0.0
    sum_tot_fee = 0.0
    sum_tot_trv = 0.0
    sum_tot_meal = 0.0
    sum_tot_rev = 0.0

    for s_idx, (m_name, m_data) in enumerate(sorted_methods, 1):
        m_rev = m_data['fee'] + m_data['travel'] + m_data['meal']
        share = (m_rev / total_sum_fee * 100) if total_sum_fee > 0 else 0.0
        eq_str = ', '.join(sorted(list(m_data['equips']))) if m_data['equips'] else "-"

        sum_tot_cnt += m_data['count']
        sum_tot_qty += m_data['qty']
        sum_tot_fee += m_data['fee']
        sum_tot_trv += m_data['travel']
        sum_tot_meal += m_data['meal']
        sum_tot_rev += m_rev

        row_s_vals = [
            s_idx, m_name, m_data['count'], m_data['qty'], 
            m_data['fee'], m_data['travel'], m_data['meal'], 
            m_rev, f"{share:.1f}%"
        ]
        for c_i, val in enumerate(row_s_vals, 1):
            cell = ws.cell(row=r, column=c_i, value=val)
            cell.font = NORMAL_FONT
            cell.border = BORDER_ALL
            if c_i in (3, 4, 5, 6, 7, 8):
                cell.number_format = '#,##0'
                cell.alignment = RIGHT_ALIGN
            elif c_i in (1, 2, 9):
                cell.alignment = CENTER_ALIGN
            else:
                cell.alignment = LEFT_ALIGN

        ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=13)
        c_eq = ws.cell(row=r, column=10, value=eq_str)
        c_eq.font = NORMAL_FONT
        c_eq.alignment = LEFT_ALIGN
        for c_i in range(10, 14):
            ws.cell(row=r, column=c_i).border = BORDER_ALL
        r += 1

    # Summary Total Row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c_tot_s = ws.cell(row=r, column=1, value="검사방법별 합계")
    c_tot_s.font = BOLD_FONT
    c_tot_s.fill = YELLOW_FILL
    c_tot_s.border = BORDER_ALL
    c_tot_s.alignment = CENTER_ALIGN
    ws.cell(row=r, column=2).border = BORDER_ALL
    ws.cell(row=r, column=2).fill = YELLOW_FILL

    s_totals = [
        (3, sum_tot_cnt, '#,##0'),
        (4, sum_tot_qty, '#,##0'),
        (5, sum_tot_fee, '#,##0'),
        (6, sum_tot_trv, '#,##0'),
        (7, sum_tot_meal, '#,##0'),
        (8, sum_tot_rev, '#,##0'),
        (9, "100.0%", '@')
    ]
    for c_i, val, fmt in s_totals:
        cell = ws.cell(row=r, column=c_i, value=val)
        cell.font = BOLD_FONT
        cell.fill = YELLOW_FILL
        cell.border = BORDER_ALL
        cell.number_format = fmt
        cell.alignment = CENTER_ALIGN if c_i == 9 else RIGHT_ALIGN

    ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=13)
    c_eq_tot = ws.cell(row=r, column=10, value=f"총 {len(sorted_methods)}개 검사방법 운용")
    c_eq_tot.font = BOLD_FONT
    c_eq_tot.fill = YELLOW_FILL
    c_eq_tot.alignment = CENTER_ALIGN
    for c_i in range(10, 14):
        ws.cell(row=r, column=c_i).border = BORDER_ALL
        ws.cell(row=r, column=c_i).fill = YELLOW_FILL




# 사후원가 보고서(app.py)와 동일한 OT 분류 함수를 재사용한다.
# (연장 18~22시 / 야간 22시~ / 휴일: 주말, 금요일 자정 이후)
try:
    from site_apps.lotte.src.models.worker_model import (
        _parse_ot_hours_impl, _calculate_split_ot_hours_impl, _calculate_ot_from_worktime_impl,
    )
    from site_apps.lotte.src.utils.helpers import MARKER_PATTERN as _MARKER_PATTERN
except Exception:  # pragma: no cover - 패키지 경로가 다른 환경 대비
    _parse_ot_hours_impl = _calculate_split_ot_hours_impl = _calculate_ot_from_worktime_impl = None
    _MARKER_PATTERN = re.compile(r'\[.*?\]')

DEFAULT_OT_RATES = {'day': 4000.0, 'night': 5000.0, 'holiday': 7500.0}


class _OtCalc:
    """worker_model의 *_impl 함수가 요구하는 self 대체 객체."""
    def _parse_ot_hours(self, value):
        return _parse_ot_hours_impl(self, value)


def _split_ot_hours_like_report(ot_raw, wt_raw, date_val):
    """보고서와 동일하게 (연장, 야간, 휴일) 시간을 반환한다. 계산 불가 시 None."""
    if _calculate_split_ot_hours_impl is None:
        return None
    calc = _OtCalc()
    try:
        dt = pd.to_datetime(date_val)
    except Exception:
        dt = date_val
    ot_compact = ot_raw.replace(',', '')
    if ot_compact.isdigit() and int(ot_compact) > 100:
        ot_hours, _ = _calculate_ot_from_worktime_impl(calc, wt_raw, dt)
        split_source = f"{ot_hours}시간 ({ot_raw}원)"
        clean_wt = _MARKER_PATTERN.sub('', wt_raw).strip()
        if '~' in clean_wt:
            split_source = f"{clean_wt.split('~')[0]}~{ot_hours}시간 ({ot_raw}원)"
    else:
        split_source = ot_raw
    return _calculate_split_ot_hours_impl(calc, split_source, dt)


def _calc_ot_hours_val(time_str):
    """(야간) 22:00~익일06:00 등에서 실제 근무시간(시간)을 부동소수점 숫자로 계산"""
    if not time_str or str(time_str).lower() in ('nan', 'none', '-'):
        return 0.0
    m = re.search(r'(\d{1,2}):(\d{2})\s*~\s*(익일)?\s*(\d{1,2}):(\d{2})', str(time_str))
    if not m:
        return 0.0
    sh, sm, next_day, eh, em = m.groups()
    sh, sm, eh, em = int(sh), int(sm), int(eh), int(em)
    start_min = sh * 60 + sm
    end_min = eh * 60 + em
    if next_day or end_min < start_min:
        end_min += 24 * 60
    return (end_min - start_min) / 60.0


def _format_ot_hours(time_str):
    """(야간) 22:00~익일06:00 형태를 (야간) 8시간 형태로 깔끔하게 변환"""
    if not time_str or str(time_str).lower() in ('nan', 'none', '-'):
        return "-"
    s = str(time_str).strip()
    prefix = ""
    m_pref = re.search(r'(\([^\)]+\))', s)
    if m_pref:
        prefix = m_pref.group(1).strip()
    hours = _calc_ot_hours_val(s)
    if hours > 0:
        h_str = f"{int(hours)}시간" if hours.is_integer() else f"{hours:.1f}시간"
        return f"{prefix} {h_str}".strip() if prefix else h_str
    return s


def _build_proof_labor_sheet(ws, site, raw_df, actual_labor_data):
    """Sheet 4: Detailed labor active days and OT breakdown"""
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 9

    # Column widths (18 columns for Person x Month matrix)
    col_widths = {
        'A': 6, 'B': 12, 'C': 12,
        'D': 18, 'E': 18, 'F': 18, 'G': 13, 'H': 18, 'I': 12,
        'J': 10, 'K': 10, 'L': 10, 'M': 10, 'N': 10, 'O': 10,
        'P': 13, 'Q': 18, 'R': 12
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # Title
    ws.merge_cells('A1:H1')
    c = ws['A1']
    c.value = f"[{site}] 사후원가 증빙 2 - 노무비 투입공수 및 특별근무(OT) 증빙 상세"
    c.font = Font(name='맑은 고딕', size=16, bold=True)
    c.alignment = CENTER_ALIGN
    ws.row_dimensions[1].height = 28

    ws.merge_cells('A2:H2')
    ws['A2'] = f"출력일시: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')} | 실적 작업자 투입 일수 및 OT 상세 집계"
    ws['A2'].font = NORMAL_FONT
    ws['A2'].alignment = LEFT_ALIGN

    ranks = ["이사", "부장", "차장", "과장", "대리", "계장", "주임", "기사"]
    rank_pattern = re.compile(r"\[?(이사|부장|차장|과장|대리|계장|주임|기사)\]?\s*(.*)")
    worker_rank_map = {
        "주진철": "부장", "우명광": "대리", "김진환": "주임", "장승대": "계장",
        "김성렬": "주임", "박광복": "부장", "주영광": "과장", "이경재": "주임",
        "황조현": "주임", "김춘호": "차장", "박원준": "과장", "이봉주": "부장",
        "김성현": "대리"
    }

    # Extract Unique Dates per worker
    worker_dates = {}
    ot_records = []
    processed_ot = set()
    ot_rates = dict(DEFAULT_OT_RATES)
    for _k, _t in (('day', '연장근무'), ('night', '야간근무'), ('holiday', '휴일근무')):
        _p = _to_f((actual_labor_data or {}).get(_t, {}).get('unit_price', 0))
        if _p > 0:
            ot_rates[_k] = _p

    for _, row in raw_df.iterrows():
        d_str = str(row.get('Date', ''))[:10]
        for i in range(1, 11):
            suffix = '' if i == 1 else str(i)
            u_raw = str(row.get(f'User{suffix}', '')).strip()
            if not u_raw or u_raw.lower() in ('nan', 'none', ''):
                continue
            m = rank_pattern.search(u_raw)
            rank = None
            w_name = u_raw
            if m and m.group(1) in ranks:
                rank = m.group(1)
                w_name = m.group(2).strip() or rank
            if not rank:
                clean_name = re.sub(r'\(.*?\)', '', u_raw).strip()
                rank = worker_rank_map.get(clean_name, '기사')
                w_name = clean_name
                
            key = (rank, w_name)
            if key not in worker_dates:
                worker_dates[key] = set()
            worker_dates[key].add(d_str)

            # OT (사후원가 보고서와 동일 기준: 작업자·일자당 최초 1건만, 시간×단가로 산정)
            ot_raw = str(row.get(f'OT{suffix}', '')).strip()
            wt_raw = str(row.get(f'WorkTime{suffix}', '')).strip()
            if ot_raw.casefold() not in ('', '0', '0.0', 'nan', 'none', 'null'):
                ot_key = (d_str, w_name.strip().casefold())
                if ot_key not in processed_ot:
                    processed_ot.add(ot_key)
                    split = _split_ot_hours_like_report(ot_raw, wt_raw, row.get('Date'))
                    if split is None:
                        split = (0.0, 0.0, 0.0)
                    h_day, h_night, h_hol = split
                    ot_amt = (h_day * ot_rates['day'] + h_night * ot_rates['night']
                              + h_hol * ot_rates['holiday'])
                    ot_records.append((d_str, rank, w_name, wt_raw, ot_raw, ot_amt,
                                       h_day, h_night, h_hol))

    # Table 1: Regular Work Days
    r = 4
    headers_reg = ["No", "직급", "작업자명", "투입공수(일)", "일단가(원)", "산출노무비(원)", "최초투입일", "최종투입일"]
    ws.cell(row=r, column=1, value="1) 직급 및 작업자별 정시 투입공수 상세 (Unique Working Days)").font = HEADER_FONT
    ws.cell(row=r, column=1).fill = GREEN_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers_reg))
    for c_i in range(1, len(headers_reg)+1): ws.cell(row=r, column=c_i).border = BORDER_ALL

    r += 1
    for c_i, h in enumerate(headers_reg, 1):
        cell = ws.cell(row=r, column=c_i, value=h)
        cell.font = BOLD_FONT
        cell.fill = SUB_FILL
        cell.border = BORDER_ALL
        cell.alignment = CENTER_ALIGN

    total_days = 0.0
    total_reg_cost = 0.0

    sorted_workers = sorted(
        worker_dates.items(), 
        key=lambda x: (ranks.index(x[0][0]) if x[0][0] in ranks else 99, -len(x[1]))
    )

    r += 1
    for idx, ((rank, w_name), d_set) in enumerate(sorted_workers, 1):
        days = len(d_set)
        dates_sorted = sorted(list(d_set))
        first_d = dates_sorted[0]
        last_d = dates_sorted[-1]
        
        # Unit price from actual_labor_data
        unit_price = 0.0
        if actual_labor_data and rank in actual_labor_data:
            unit_price = _to_f(actual_labor_data[rank].get('unit_price', 0))
        if unit_price == 0:
            unit_price = 230208 if rank in ('이사', '부장') else (198625 if rank == '차장' else 170000)

        cost = days * unit_price
        total_days += days
        total_reg_cost += cost

        row_vals = [idx, rank, w_name, days, unit_price, cost, first_d, last_d]
        for c_i, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r, column=c_i, value=val)
            cell.font = NORMAL_FONT
            cell.border = BORDER_ALL
            if c_i in (4, 5, 6):
                cell.number_format = '#,##0'
                cell.alignment = RIGHT_ALIGN
            else:
                cell.alignment = CENTER_ALIGN
        r += 1

    # Total Row for Table 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c_tot = ws.cell(row=r, column=1, value=f"정시투입 합계 (총 {len(sorted_workers)}명)")
    c_tot.font = BOLD_FONT
    c_tot.fill = YELLOW_FILL
    c_tot.border = BORDER_ALL
    c_tot.alignment = CENTER_ALIGN
    for c_i in range(2, 4):
        ws.cell(row=r, column=c_i).border = BORDER_ALL
        ws.cell(row=r, column=c_i).fill = YELLOW_FILL

    ws.cell(row=r, column=4, value=total_days).number_format = '#,##0'
    ws.cell(row=r, column=4).font = BOLD_FONT
    ws.cell(row=r, column=4).fill = YELLOW_FILL
    ws.cell(row=r, column=4).alignment = RIGHT_ALIGN
    ws.cell(row=r, column=4).border = BORDER_ALL

    ws.cell(row=r, column=5, value="").border = BORDER_ALL
    ws.cell(row=r, column=5).fill = YELLOW_FILL

    ws.cell(row=r, column=6, value=total_reg_cost).number_format = '#,##0'
    ws.cell(row=r, column=6).font = BOLD_FONT
    ws.cell(row=r, column=6).fill = YELLOW_FILL
    ws.cell(row=r, column=6).alignment = RIGHT_ALIGN
    ws.cell(row=r, column=6).border = BORDER_ALL

    ws.cell(row=r, column=7, value="").border = BORDER_ALL
    ws.cell(row=r, column=7).fill = YELLOW_FILL
    ws.cell(row=r, column=8, value="").border = BORDER_ALL
    ws.cell(row=r, column=8).fill = YELLOW_FILL

    r += 3

    # 인월별 매트릭스(구 2번 표)는 요약 성격에 맞지 않아 출력하지 않는다.
    # 합계/비중 계산에 필요한 값만 유지한다.
    total_all_ot_amt = sum(rec[5] for rec in ot_records)
    total_all_ot_hours = sum(rec[6] + rec[7] + rec[8] for rec in ot_records)

    # Table 3: Worker Summary by OT Work Type (Person x Work Type Matrix)
    headers_type_summary = ["No", "직급", "작업자명", "야간근무(시간)", "주간연장(시간)", "휴일근무(시간)", "총 시간(h)", "지급 OT총액(원)", "비중(%)"]
    ws.cell(row=r, column=1, value="2) 작업자별 특별근무(야간/주간/휴일) 시간 및 금액 집계 요약 (Person × OT Type Matrix)").font = HEADER_FONT
    ws.cell(row=r, column=1).fill = ORANGE_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers_type_summary))
    for c_i in range(1, len(headers_type_summary) + 1):
        ws.cell(row=r, column=c_i).border = BORDER_ALL

    r += 1
    for c_i, h in enumerate(headers_type_summary, 1):
        cell = ws.cell(row=r, column=c_i, value=h)
        cell.font = BOLD_FONT
        cell.fill = SUB_FILL
        cell.border = BORDER_ALL
        cell.alignment = CENTER_ALIGN

    # Aggregate by worker and work type
    worker_ot_summary = {}
    for d_str, rank, w_name, wt_str, ot_str, ot_amt, h_day, h_night, h_hol in ot_records:
        key = (rank, w_name)
        if key not in worker_ot_summary:
            worker_ot_summary[key] = {
                'night_hours': 0.0, 'night_amt': 0.0,
                'day_hours': 0.0, 'day_amt': 0.0,
                'holiday_hours': 0.0, 'holiday_amt': 0.0,
                'total_hours': 0.0, 'total_amt': 0.0
            }
        info_ws = worker_ot_summary[key]
        info_ws['night_hours'] += h_night
        info_ws['night_amt'] += h_night * ot_rates['night']
        info_ws['day_hours'] += h_day
        info_ws['day_amt'] += h_day * ot_rates['day']
        info_ws['holiday_hours'] += h_hol
        info_ws['holiday_amt'] += h_hol * ot_rates['holiday']
        info_ws['total_hours'] += h_day + h_night + h_hol
        info_ws['total_amt'] += ot_amt

    sorted_worker_ot = sorted(worker_ot_summary.items(), key=lambda x: x[1]['total_hours'], reverse=True)

    r += 1
    tot_night_h_sum = sum(v['night_hours'] for v in worker_ot_summary.values())
    tot_day_h_sum = sum(v['day_hours'] for v in worker_ot_summary.values())
    tot_hol_h_sum = sum(v['holiday_hours'] for v in worker_ot_summary.values())

    for idx, ((rank, w_name), info) in enumerate(sorted_worker_ot, 1):
        share = (info['total_amt'] / total_all_ot_amt * 100) if total_all_ot_amt > 0 else 0.0
        row_vals = [
            idx, rank, w_name,
            round(info['night_hours'], 1) if info['night_hours'] > 0 else "-",
            round(info['day_hours'], 1) if info['day_hours'] > 0 else "-",
            round(info['holiday_hours'], 1) if info['holiday_hours'] > 0 else "-",
            round(info['total_hours'], 1),
            info['total_amt'],
            f"{share:.1f}%"
        ]

        for c_i, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r, column=c_i, value=val)
            cell.font = NORMAL_FONT
            cell.border = BORDER_ALL
            if c_i in (1, 2):
                cell.alignment = CENTER_ALIGN
            elif c_i == 3:
                cell.alignment = CENTER_ALIGN
            elif c_i in (4, 5, 6, 7):
                if isinstance(val, (int, float)):
                    cell.number_format = '#,##0.0'
                    cell.alignment = RIGHT_ALIGN
                else:
                    cell.alignment = CENTER_ALIGN
            elif c_i == 8:
                cell.number_format = '#,##0'
                cell.alignment = RIGHT_ALIGN
            elif c_i == 9:
                cell.alignment = CENTER_ALIGN
        r += 1

    # Total row for Table 3
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c_tot_type = ws.cell(row=r, column=1, value=f"특별근무 합계 (총 {len(sorted_worker_ot)}명)")
    c_tot_type.font = BOLD_FONT
    c_tot_type.fill = YELLOW_FILL
    c_tot_type.border = BORDER_ALL
    c_tot_type.alignment = CENTER_ALIGN
    for c_i in range(2, 4):
        ws.cell(row=r, column=c_i).border = BORDER_ALL
        ws.cell(row=r, column=c_i).fill = YELLOW_FILL

    # Night hours total
    c_n = ws.cell(row=r, column=4, value=tot_night_h_sum)
    c_n.font = BOLD_FONT; c_n.fill = YELLOW_FILL; c_n.border = BORDER_ALL
    c_n.number_format = '#,##0.0'; c_n.alignment = RIGHT_ALIGN

    # Day hours total
    c_d = ws.cell(row=r, column=5, value=tot_day_h_sum)
    c_d.font = BOLD_FONT; c_d.fill = YELLOW_FILL; c_d.border = BORDER_ALL
    c_d.number_format = '#,##0.0'; c_d.alignment = RIGHT_ALIGN

    # Holiday hours total
    c_h = ws.cell(row=r, column=6, value=tot_hol_h_sum)
    c_h.font = BOLD_FONT; c_h.fill = YELLOW_FILL; c_h.border = BORDER_ALL
    c_h.number_format = '#,##0.0'; c_h.alignment = RIGHT_ALIGN

    # Total hours
    c_th = ws.cell(row=r, column=7, value=total_all_ot_hours)
    c_th.font = BOLD_FONT; c_th.fill = YELLOW_FILL; c_th.border = BORDER_ALL
    c_th.number_format = '#,##0.0'; c_th.alignment = RIGHT_ALIGN

    # Total amount
    c_ta = ws.cell(row=r, column=8, value=total_all_ot_amt)
    c_ta.font = BOLD_FONT; c_ta.fill = YELLOW_FILL; c_ta.border = BORDER_ALL
    c_ta.number_format = '#,##0'; c_ta.alignment = RIGHT_ALIGN

    # 100%
    c_tp = ws.cell(row=r, column=9, value="100.0%")
    c_tp.font = BOLD_FONT; c_tp.fill = YELLOW_FILL; c_tp.border = BORDER_ALL
    c_tp.alignment = CENTER_ALIGN



def _build_proof_materials_and_expenses_sheet(ws, site, raw_df, actual_material_data, actual_expense_data):
    """Sheet 5: Materials consumption, vehicle logs, and special equipment usage"""
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 9

    col_widths = {'A': 6, 'B': 24, 'C': 18, 'D': 18, 'E': 24, 'F': 28, 'G': 18, 'H': 25}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # Title
    ws.merge_cells('A1:G1')
    c = ws['A1']
    c.value = f"[{site}] 사후원가 증빙 3 - 자재 소모·차량 운행·특수장비 투입 증빙 상세"
    c.font = Font(name='맑은 고딕', size=16, bold=True)
    c.alignment = CENTER_ALIGN
    ws.row_dimensions[1].height = 28

    ws.merge_cells('A2:G2')
    ws['A2'] = f"출력일시: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')} | 실적 자재·차량·장비 운용 집계"
    ws['A2'].font = NORMAL_FONT
    ws['A2'].alignment = LEFT_ALIGN

    # --- 1. Materials Consumption Table ---
    r = 4
    headers_mat = ["No", "품목명", "사양/구분", "단위", "실소모수량", "기준단가(원)", "소모금액(원)"]
    ws.cell(row=r, column=1, value="1) 자재 소모 실적 집계 (NDT 약품 및 검사용 필름)").font = HEADER_FONT
    ws.cell(row=r, column=1).fill = BLUE_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers_mat))
    for c_i in range(1, len(headers_mat)+1): ws.cell(row=r, column=c_i).border = BORDER_ALL

    r += 1
    for c_i, h in enumerate(headers_mat, 1):
        cell = ws.cell(row=r, column=c_i, value=h)
        cell.font = BOLD_FONT
        cell.fill = SUB_FILL
        cell.border = BORDER_ALL
        cell.alignment = CENTER_ALIGN

    # Build Mat usage mapping from raw_df
    mat_items = [
        ("세척제", "PT 약품", "CAN", raw_df['NDT_세척제'].sum() if 'NDT_세척제' in raw_df.columns else 0.0, 1500),
        ("침투제", "PT 약품", "CAN", raw_df['NDT_침투제'].sum() if 'NDT_침투제' in raw_df.columns else 0.0, 2300),
        ("현상제", "PT 약품", "CAN", raw_df['NDT_현상제'].sum() if 'NDT_현상제' in raw_df.columns else 0.0, 2000),
        ("흑색자분", "MT 약품", "CAN", raw_df['NDT_흑색자분'].sum() if 'NDT_흑색자분' in raw_df.columns else 0.0, 1800),
        ("백색페인트", "MT 약품", "CAN", raw_df['NDT_백색페인트'].sum() if 'NDT_백색페인트' in raw_df.columns else 0.0, 2350),
        ("형광자분", "MT 약품", "CAN", raw_df['NDT_형광자분'].sum() if 'NDT_형광자분' in raw_df.columns else 0.0, 2500),
        ("방사선 필름", "MX125", "매", raw_df['FilmCount'].sum() if 'FilmCount' in raw_df.columns else 0.0, 990),
    ]

    # If actual_material_data passed, override unit price
    price_map = {}
    if actual_material_data:
        for it in actual_material_data:
            spec = it.get('spec', '') or it.get('name', '')
            price_map[spec] = _to_f(it.get('price', 0))

    r += 1
    total_mat_cost = 0.0
    for idx, (spec, cat, unit, qty, def_price) in enumerate(mat_items, 1):
        price = price_map.get(spec, def_price)
        amt = qty * price
        total_mat_cost += amt

        row_vals = [idx, cat, spec, unit, qty, price, amt]
        for c_i, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r, column=c_i, value=val)
            cell.font = NORMAL_FONT
            cell.border = BORDER_ALL
            if c_i in (5, 6, 7):
                cell.number_format = '#,##0'
                cell.alignment = RIGHT_ALIGN
            else:
                cell.alignment = CENTER_ALIGN
        r += 1

    # Total Row for Materials
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c_tot = ws.cell(row=r, column=1, value="자재 실소모 총 합계")
    c_tot.font = BOLD_FONT
    c_tot.fill = YELLOW_FILL
    c_tot.border = BORDER_ALL
    c_tot.alignment = CENTER_ALIGN
    for c_i in range(2, 7):
        ws.cell(row=r, column=c_i).border = BORDER_ALL
        ws.cell(row=r, column=c_i).fill = YELLOW_FILL

    cell_tot_mat = ws.cell(row=r, column=7, value=total_mat_cost)
    cell_tot_mat.font = BOLD_FONT
    cell_tot_mat.fill = YELLOW_FILL
    cell_tot_mat.border = BORDER_ALL
    cell_tot_mat.number_format = '#,##0'
    cell_tot_mat.alignment = RIGHT_ALIGN

    r += 3

    # --- 2. Vehicles Table ---
    headers_veh = ["No", "차량번호", "운행기록건수(건)", "고유운행일수(일)", "누계 주행거리(km)", "운행기간", "비고"]
    ws.cell(row=r, column=1, value="2) 차량 운행 일지 실적 집계").font = HEADER_FONT
    ws.cell(row=r, column=1).fill = GREEN_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers_veh))
    for c_i in range(1, len(headers_veh)+1): ws.cell(row=r, column=c_i).border = BORDER_ALL

    r += 1
    for c_i, h in enumerate(headers_veh, 1):
        cell = ws.cell(row=r, column=c_i, value=h)
        cell.font = BOLD_FONT
        cell.fill = SUB_FILL
        cell.border = BORDER_ALL
        cell.alignment = CENTER_ALIGN

    # Group Vehicles
    veh_dict = {}
    if '차량번호' in raw_df.columns:
        for _, row in raw_df.iterrows():
            d_str = str(row.get('Date', ''))[:10]
            v_raw = str(row.get('차량번호', '')).strip()
            km_raw = str(row.get('주행거리', '')).strip()
            if not v_raw or v_raw.lower() in ('nan', 'none', ''):
                continue
            nos = [x.strip() for x in v_raw.split('||') if x.strip()]
            kms = [x.strip() for x in km_raw.split('||') if x.strip()]
            for j, no in enumerate(nos):
                if no not in veh_dict:
                    veh_dict[no] = {'dates': set(), 'kms': [], 'count': 0}
                veh_dict[no]['dates'].add(d_str)
                veh_dict[no]['count'] += 1
                if j < len(kms) and kms[j].replace(',', '').isdigit():
                    veh_dict[no]['kms'].append(float(kms[j].replace(',', '')))

    r += 1
    tot_veh_days = 0
    tot_veh_km = 0.0
    for idx, (v_no, info) in enumerate(sorted(veh_dict.items(), key=lambda x: -len(x[1]['dates'])), 1):
        u_days = len(info['dates'])
        tot_veh_days += u_days
        dates_s = sorted(list(info['dates']))
        period_str = f"{dates_s[0]} ~ {dates_s[-1]}" if dates_s else ""
        max_km = sum(info['kms']) if info['kms'] else "-"
        tot_veh_km += sum(info['kms'])

        car_type = "스타렉스" if "스타렉스" in v_no else ("탑차" if "탑차" in v_no else "현장차량")

        row_vals = [idx, v_no, info['count'], u_days, max_km, period_str, car_type]
        for c_i, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r, column=c_i, value=val)
            cell.font = NORMAL_FONT
            cell.border = BORDER_ALL
            if c_i in (3, 4, 5):
                cell.number_format = '#,##0'
                cell.alignment = RIGHT_ALIGN if isinstance(val, (int, float)) else CENTER_ALIGN
            elif c_i in (1, 2, 6, 7):
                cell.alignment = CENTER_ALIGN
            else:
                cell.alignment = LEFT_ALIGN
        r += 1

    # Total Row for Vehicles
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c_totv = ws.cell(row=r, column=1, value=f"차량 운행 총계 (총 {len(veh_dict)}대)")
    c_totv.font = BOLD_FONT
    c_totv.fill = YELLOW_FILL
    c_totv.border = BORDER_ALL
    c_totv.alignment = CENTER_ALIGN
    ws.cell(row=r, column=2).border = BORDER_ALL
    ws.cell(row=r, column=2).fill = YELLOW_FILL

    cell_tot_vcnt = ws.cell(row=r, column=3, value=sum(i['count'] for i in veh_dict.values()))
    cell_tot_vcnt.font = BOLD_FONT
    cell_tot_vcnt.fill = YELLOW_FILL
    cell_tot_vcnt.border = BORDER_ALL
    cell_tot_vcnt.number_format = '#,##0'
    cell_tot_vcnt.alignment = RIGHT_ALIGN

    cell_tot_vdays = ws.cell(row=r, column=4, value=tot_veh_days)
    cell_tot_vdays.font = BOLD_FONT
    cell_tot_vdays.fill = YELLOW_FILL
    cell_tot_vdays.border = BORDER_ALL
    cell_tot_vdays.number_format = '#,##0'
    cell_tot_vdays.alignment = RIGHT_ALIGN

    for c_i in range(5, 8):
        ws.cell(row=r, column=c_i).border = BORDER_ALL
        ws.cell(row=r, column=c_i).fill = YELLOW_FILL
    cell_tot_vkm = ws.cell(row=r, column=5, value=tot_veh_km)
    cell_tot_vkm.font = BOLD_FONT
    cell_tot_vkm.number_format = '#,##0'
    cell_tot_vkm.alignment = RIGHT_ALIGN

    r += 3

    # --- 3. Special Equipment Table ---
    headers_eq = ["No", "장비/검사구분", "투입건수(건)", "고유투입일수(일)", "투입기간", "비고"]
    ws.cell(row=r, column=1, value="3) 특수장비 및 주요 검사방법 투입 실적").font = HEADER_FONT
    ws.cell(row=r, column=1).fill = ORANGE_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers_eq))
    for c_i in range(1, len(headers_eq)+1): ws.cell(row=r, column=c_i).border = BORDER_ALL

    r += 1
    for c_i, h in enumerate(headers_eq, 1):
        cell = ws.cell(row=r, column=c_i, value=h)
        cell.font = BOLD_FONT
        cell.fill = SUB_FILL
        cell.border = BORDER_ALL
        cell.alignment = CENTER_ALIGN

    eq_dict = {}
    for _, row in raw_df.iterrows():
        d_str = str(row.get('Date', ''))[:10]
        method = str(row.get('검사방법', '')).strip().upper()
        equip_list = [e.upper() for e in split_equipments(row.get('장비명', ''))]

        def _classify(equip):
            if 'PAUT' in method or 'PAUT' in equip: return "PAUT 장비"
            if 'PMI' in method or 'PMI' in equip: return "PMI 장비"
            if 'MT' in method or 'YOKE' in equip: return "MT (YOKE) 장비"
            if 'UT' in method: return "UT 초음파탐상장비"
            if 'RT' in method: return "RT 방사선투과검사장비"
            return None

        targets = set()
        if len(equip_list) > 1:
            # [MULTI-EQUIP] 장비별로 분류 (검사방법이 아닌 장비명 기준)
            for equip in equip_list:
                if 'PAUT' in equip: targets.add("PAUT 장비")
                elif 'PMI' in equip: targets.add("PMI 장비")
                elif 'YOKE' in equip or 'MT' in equip: targets.add("MT (YOKE) 장비")
                elif 'UT' in equip: targets.add("UT 초음파탐상장비")
                elif 'RT' in equip: targets.add("RT 방사선투과검사장비")
            if not targets:
                t = _classify(' '.join(equip_list))
                if t: targets.add(t)
        else:
            t = _classify(equip_list[0] if equip_list else '')
            if t: targets.add(t)

        for target in targets:
            if target not in eq_dict: eq_dict[target] = {'dates': set(), 'count': 0}
            eq_dict[target]['dates'].add(d_str)
            eq_dict[target]['count'] += 1

    r += 1
    for idx, (eq_name, info) in enumerate(sorted(eq_dict.items(), key=lambda x: -len(x[1]['dates'])), 1):
        u_days = len(info['dates'])
        dates_s = sorted(list(info['dates']))
        p_str = f"{dates_s[0]} ~ {dates_s[-1]}" if dates_s else ""
        row_vals = [idx, eq_name, info['count'], u_days, p_str, "현장 실적"]
        for c_i, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r, column=c_i, value=val)
            cell.font = NORMAL_FONT
            cell.border = BORDER_ALL
            if c_i in (3, 4):
                cell.number_format = '#,##0'
                cell.alignment = RIGHT_ALIGN
            elif c_i in (1, 2, 5, 6):
                cell.alignment = CENTER_ALIGN
        r += 1


def export_budget_estimation_to_excel(filepath, summary_data, labor_data, material_data, expense_data,
                                     actual_summary_data=None, actual_labor_data=None,
                                     actual_material_data=None, actual_expense_data=None,
                                     raw_df=None):
    """
    Export the detailed budget estimation to a multi-sheet Excel file:
      - Sheet 1: 사전예산서 (Planned)
      - Sheet 2: 사후원가(실적) 보고서 (Actuals)
      - Sheet 3: 증빙_검사실적 (Inspection Revenue Backup)
      - Sheet 4: 증빙_노무투입 (Labor Cost & OT Backup)
      - Sheet 5: 증빙_자재·경비·차량 (Materials, Expenses & Vehicle Backup)
    """
    wb = Workbook()
    site = (summary_data or {}).get('site', '') or (actual_summary_data or {}).get('site', '공사현장')
    
    # 1. Sheet 1: 사전예산서
    ws_planned = wb.active
    ws_planned.title = "사전예산서"
    _build_budget_sheet(ws_planned, summary_data or {}, labor_data or {}, material_data or [], expense_data or {}, mode='planned')
    
    # 2. Sheet 2: 사후원가(실적) 보고서
    has_actual_data = (
        bool(actual_summary_data) or 
        bool(actual_labor_data) or 
        bool(actual_material_data) or 
        bool(actual_expense_data)
    )
    if has_actual_data:
        ws_actual = wb.create_sheet(title="사후원가(실적) 보고서")
        _build_budget_sheet(ws_actual, actual_summary_data or {}, actual_labor_data or {}, 
                            actual_material_data or [], actual_expense_data or {}, mode='actual')

    # 3. Supporting Evidence Sheets (Sheets 3, 4, 5)
    if raw_df is not None and not raw_df.empty:
        ws_proof_insp = wb.create_sheet(title="증빙_검사실적")
        _build_proof_inspection_sheet(ws_proof_insp, site, raw_df)

        ws_proof_labor = wb.create_sheet(title="증빙_노무투입")
        _build_proof_labor_sheet(ws_proof_labor, site, raw_df, actual_labor_data or {})

        ws_proof_mat = wb.create_sheet(title="증빙_자재·경비·차량")
        _build_proof_materials_and_expenses_sheet(ws_proof_mat, site, raw_df, actual_material_data or [], actual_expense_data or {})

    wb.save(filepath)
    return True
