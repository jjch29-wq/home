import csv
import io
import json
from pathlib import Path
import sys
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from paut_fields import TCG_MAX, normalize_tcg
from paut_report_layout import prepare_indication_row
import openpyxl


class TcgFieldsTests(unittest.TestCase):
    def test_daily_section_tcg_is_imported_and_restored(self):
        from ndt_history_import import paut_flaw_findings
        text = 'X:203.5-209.0,L:5.8,D:1.175~1.259,H:0.83,TCG30%,Surface'
        result = paut_flaw_findings(text, '4.9')
        self.assertEqual(result[TCG_MAX], '30')
        self.assertEqual(result['Remarks'], 'Surface')
        restored = normalize_tcg({'_source_paut_section': text, TCG_MAX: ''})
        self.assertEqual(restored[TCG_MAX], '30')
        self.assertEqual(normalize_tcg({'Remarks': '30% unrelated'})[TCG_MAX], '')

    def test_existing_names_and_zero_are_restored(self):
        for name in ('TCG Max', 'TCG MAX (%)', 'TCG\nMax\n(%)', 'tcg_max_percent'):
            with self.subTest(name=name):
                item = {name: 0}
                normalize_tcg(item)
                self.assertEqual(item[TCG_MAX], 0)

    def test_primary_value_is_preserved(self):
        item = {TCG_MAX: '65%', 'TCG Max': '80%'}
        normalize_tcg(item)
        self.assertEqual(item[TCG_MAX], '65%')

    def test_missing_values_stay_empty(self):
        for value in (None, '', float('nan')):
            self.assertEqual(normalize_tcg({'TCG Max': value})[TCG_MAX], '')

    def test_json_and_csv_roundtrip(self):
        item = normalize_tcg(json.loads('{"TCG Max": 72.5}'))
        stream = io.StringIO()
        writer = csv.DictWriter(stream, fieldnames=list(item))
        writer.writeheader()
        writer.writerow(item)
        stream.seek(0)
        restored = normalize_tcg(next(csv.DictReader(stream)))
        self.assertEqual(restored[TCG_MAX], '72.5')

    def test_accepted_row_retains_tcg_in_report_after_reload(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.merge_cells('N9:X9')
        prepare_indication_row(ws, 9, False, has_tcg=True)
        ws['N9'] = 'NO RECORDABLE INDICATION'
        ws['W9'] = 72.5
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        ws = openpyxl.load_workbook(buf).active
        self.assertEqual(ws['W9'].value, 72.5)
        self.assertEqual(ws['N9'].value, 'NO RECORDABLE INDICATION')
