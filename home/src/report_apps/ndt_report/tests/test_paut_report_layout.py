import sys
from pathlib import Path
import unittest
import openpyxl
import io

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from paut_report_layout import (prepare_indication_row, compact_report_height,
                                align_report_heights, printed_height)
from paut_report_layout import solid_right_border, fit_height_heading
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from types import SimpleNamespace


class ReportHeightTests(unittest.TestCase):
    def test_height_heading_has_two_lines_and_fits_narrow_column(self):
        ws = openpyxl.Workbook().active
        ws.merge_cells('V7:V8')
        ws['V7'] = '높이\n(mm)'
        ws['T7'] = '깊이\n(mm)'
        original_width = ws.column_dimensions['V'].width
        fit_height_heading(ws)
        self.assertEqual(ws['V7'].value, '높이\n(mm)')
        self.assertEqual(ws['V7'].font.sz, 9)
        self.assertEqual(ws['T7'].font.sz, 11)
        self.assertEqual(ws.column_dimensions['V'].width, original_width)

    def test_solid_merged_outer_border_survives_save_reload(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        from openpyxl.styles import Border, Side
        ws.merge_cells('AH4:AJ6')
        ws['AH4'].border = Border(right=Side(style='dotted'),
                                 bottom=Side(style='dotted'))
        solid_right_border(ws)
        for _ in range(2):
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            wb = openpyxl.load_workbook(buf)
        ws = wb.active
        self.assertEqual(ws['AH4'].border.right.style, 'thin')
        self.assertEqual(ws['AH4'].border.bottom.style, 'dotted')
        self.assertTrue(all(ws.cell(r, 36).border.right.style == 'thin'
                            for r in range(4, 46)))

    def test_excel_width_scaling_produces_equal_printed_heights(self):
        wb = openpyxl.Workbook()
        a = wb.active
        b = wb.create_sheet()
        for ws in (a, b):
            ws.print_area = 'A1:D10'
            ws.page_margins.left = ws.page_margins.right = .4
            for row in range(1, 11):
                ws.row_dimensions[row].height = 30
        align_report_heights(wb.worksheets, measured_widths={a.title: 540, b.title: 590})
        self.assertAlmostEqual(printed_height(a) * a.page_setup.scale,
                               printed_height(b) * b.page_setup.scale)
        self.assertNotEqual(a.page_setup.scale, b.page_setup.scale)
        self.assertFalse(a.sheet_properties.pageSetUpPr.fitToPage)
        available = 210 / 25.4 * 72 - .8 * 72
        for ws, width in ((a, 540), (b, 590)):
            self.assertLessEqual(width * ws.page_setup.scale / 100,
                                 available - width * .02)

    def test_different_source_heights_finish_at_same_height(self):
        wb = openpyxl.Workbook()
        cover = wb.active
        data = wb.create_sheet()
        for ws, height in ((cover, 20), (data, 25)):
            ws.print_area = 'A1:D10'
            for row in range(1, 11):
                ws.row_dimensions[row].height = height
        align_report_heights(wb.worksheets)
        self.assertAlmostEqual(printed_height(cover), 189)
        self.assertAlmostEqual(printed_height(data), 189)

    def test_printed_rows_and_drawing_shrink_only_vertically(self):
        ws = openpyxl.Workbook().active
        ws.print_area = 'A1:D3'
        ws.row_dimensions[1].height = 20
        ws.row_dimensions[4].height = 30
        ws.column_dimensions['A'].width = 12
        anchor = OneCellAnchor(_from=AnchorMarker(row=1, rowOff=1000),
                               ext=XDRPositiveSize2D(cx=2000, cy=4000))
        ws._images.append(SimpleNamespace(anchor=anchor))
        compact_report_height(ws)
        self.assertAlmostEqual(ws.row_dimensions[1].height, 18.9)
        self.assertAlmostEqual(ws.row_dimensions[2].height, 15 * .945)
        self.assertEqual(ws.row_dimensions[4].height, 30)
        self.assertEqual(ws.column_dimensions['A'].width, 12)
        self.assertEqual(anchor.ext.cx, 2000)
        self.assertEqual(anchor.ext.cy, 3780)
        self.assertEqual(anchor._from.rowOff, 945)


class IndicationLayoutTests(unittest.TestCase):
    def test_merged_template_accepts_all_indication_values(self):
        ws = openpyxl.Workbook().active
        ws.merge_cells('N9:X9')
        prepare_indication_row(ws, 9, True)
        values = {'N9': '203.5', 'R9': '5.8', 'T9': '1.175', 'V9': '0.83', 'AE9': 'Surface'}
        for address, value in values.items():
            ws[address] = value
        self.assertEqual({address: ws[address].value for address in values}, values)
        self.assertNotIn('N9:X9', {str(r) for r in ws.merged_cells.ranges})
        self.assertIn('W9:X9', {str(r) for r in ws.merged_cells.ranges})
        self.assertNotIn('V9:W9', {str(r) for r in ws.merged_cells.ranges})
        self.assertEqual(ws['V9'].border.right.style, 'dotted')
        self.assertEqual(ws['W9'].border.left.style, 'dotted')
        prepare_indication_row(ws, 9, False)
        ws['N9'] = 'NO RECORDABLE INDICATION'
        self.assertIn('N9:X9', {str(r) for r in ws.merged_cells.ranges})
