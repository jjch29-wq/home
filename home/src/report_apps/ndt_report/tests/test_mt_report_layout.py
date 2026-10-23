import sys
from pathlib import Path
import unittest
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from mt_report_layout import mt_data_positions, mt_length_totals, hot_tapping_points


class MTPageTests(unittest.TestCase):
    def test_hot_tapping_sums_dates_for_same_point_and_keeps_lines_separate(self):
        rows = [{'Date': '2026-10-12', 'Dwg': 'L1', 'Joint': 'HT-R01', '_source_tested_length': '0.4', 'Result': '합격'},
                {'Date': '2026-10-20', 'Dwg': 'L1', 'Joint': 'R01', '_source_tested_length': '0.6', 'Result': '합격'},
                {'Date': '2026-10-20', 'Dwg': 'L2', 'Joint': 'R01', '_source_tested_length': '0.5', 'Result': '합격'}]
        points = hot_tapping_points(rows)
        self.assertEqual(len(points), 2)
        self.assertEqual((points[0]['Joint'], points[0]['Tested Length']), ('HT-R01', '1.0'))
        self.assertIn('Length: 1.0000 M', points[0]['Remarks'])
        self.assertEqual(rows[0]['Joint'], 'HT-R01')

    def test_lengths_use_source_reinspection_not_rejected_result(self):
        rows = [{'_source_id': 'a', '_source_tested_length': '1.25', 'Result': 'Rej'},
                {'_source_id': 'b', '_source_tested_length': '0.5', '_source_work_type': '재검'}]
        original, repair, missing = mt_length_totals(rows + [rows[0]])
        self.assertEqual(str(original), '1.25')
        self.assertEqual(str(repair), '0.5')
        self.assertFalse(missing)
        self.assertTrue(mt_length_totals([{'_source_id': 'missing'}])[2])

    def test_cover_and_continuation_boundaries(self):
        for count, sheets in [(8, 1), (9, 2), (37, 2), (38, 3)]:
            with self.subTest(count=count):
                book = openpyxl.Workbook()
                cover = book.active
                continuation = book.create_sheet('MT(을)')
                for row in range(10, 39):
                    continuation.merge_cells(f'A{row}:E{row}')
                positions = mt_data_positions(book, count, lambda page: book.copy_worksheet(continuation))
                self.assertEqual(len(book.worksheets), sheets)
                self.assertEqual(len(positions), count)
                self.assertEqual(positions[0], (cover, 29))
                self.assertEqual(positions[7], (cover, 36))
                if count > 8:
                    self.assertEqual(positions[8], (continuation, 10))
                if count >= 37:
                    self.assertEqual(positions[36], (continuation, 38))
                if count == 38:
                    self.assertEqual(positions[37], (book.worksheets[2], 10))
                for index, (sheet, row) in enumerate(positions):
                    sheet[f'B{row}'] = f'R{index + 1:02}'
                    self.assertEqual(sheet[f'B{row}'].value, f'R{index + 1:02}')
