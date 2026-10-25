import sys
from pathlib import Path
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from rt_report_layout import report_date_from_exam_period, header_stamp_anchor
import openpyxl


class ReportDateTests(unittest.TestCase):
    def test_stamp_uses_each_sheets_title_boundary(self):
        for title_column in (6, 8):
            ws = openpyxl.Workbook().active
            for column in range(1, 9):
                ws.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 10
            ws.cell(1, title_column, 'RADIOGRAPHIC EXAMINATION REPORT')
            # Match dimensions as stored in the template files.
            for column, dim in enumerate(ws.column_dimensions.values(), 1):
                dim.min = dim.max = column
            anchor = header_stamp_anchor(ws)
            right = (anchor._from.col * 75 +
                     (anchor._from.colOff + anchor.ext.cx) / 9525)
            self.assertAlmostEqual(right, (title_column - 1) * 75)

    def test_full_and_abbreviated_periods(self):
        for period in ('2026-10-21~10-23', '2026-10-21 ~ 2026-10-23',
                       '2026-10-21~23'):
            self.assertEqual(report_date_from_exam_period(period),
                             '2026 년   10 월   23 일')
        self.assertEqual(report_date_from_exam_period('2026-12-31~2027-01-02'),
                         '2027 년   01 월   02 일')

    def test_single_or_invalid_date_is_preserved(self):
        for value in ('2026-10-21', '', '2026-10-21~10-32'):
            self.assertEqual(report_date_from_exam_period(value), value)
