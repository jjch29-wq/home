import io
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
import openpyxl
from openpyxl.drawing.image import Image
from PIL import Image as PILImage

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from report_signatures import append_rt_signatures


class SignatureTests(unittest.TestCase):
    def test_adds_both_signatures_without_removing_original_logo(self):
        def png():
            data = io.BytesIO()
            PILImage.new('RGBA', (40, 20), (0, 0, 0, 255)).save(data, format='PNG')
            data.seek(0)
            return data
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'report.xlsx'
            book = openpyxl.Workbook()
            book.active.add_image(Image(png()), 'A1')
            book.active['D32'] = '우명광'
            book.active['D35'] = '주진철'
            book.create_sheet('을지')
            book.create_sheet('추가 을지')
            book.save(path)
            with zipfile.ZipFile(path) as archive:
                original_logo = archive.read('xl/media/image1.png')
            append_rt_signatures(path, [(png(), 'E32', 65, 23), (png(), 'E35', 65, 23)])
            for sheet_index in (1, 2):
                append_rt_signatures(path, [(png(), 'G1', 72, 72)], sheet_index=sheet_index)
            with zipfile.ZipFile(path) as archive:
                self.assertEqual(archive.read('xl/media/image1.png'), original_logo)
                self.assertIsNone(archive.testzip())
            result = openpyxl.load_workbook(path)
            self.assertEqual(len(result.active._images), 3)
            for sheet in result.worksheets[1:]:
                self.assertEqual(len(sheet._images), 1)
                self.assertEqual(sheet._images[0].anchor._from.col, 6)
                self.assertEqual(sheet._images[0].anchor._from.row, 0)
            self.assertEqual(result.active['D32'].value, '우명광')
            result.close()
