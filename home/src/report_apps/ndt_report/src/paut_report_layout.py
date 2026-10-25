"""Restore the district-heating template's indication columns per data row."""
from copy import copy
from openpyxl.styles import Side
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.print_settings import PrintArea


def printed_row_numbers(ws):
    areas = list(PrintArea.from_string(str(ws.print_area)).ranges) if ws.print_area else []
    bounds = [range_boundaries(str(area)) for area in areas]
    return {row for _, first, _, last in bounds for row in range(first, last + 1)}


def solid_right_border(ws, column=36, first_row=4, last_row=45):
    """Update merged masters too, so reloading cannot restore dotted edges."""
    solid = Side(style='thin', color='000000')
    for merged in ws.merged_cells.ranges:
        if merged.max_col == column and merged.min_row <= last_row and merged.max_row >= first_row:
            master = ws.cell(merged.min_row, merged.min_col)
            border = copy(master.border)
            border.right = solid
            master.border = border
    for row in range(first_row, last_row + 1):
        cell = ws.cell(row, column)
        border = copy(cell.border)
        border.right = solid
        cell.border = border


def fit_height_heading(ws):
    """The single-column height heading needs a smaller font than its peers."""
    for row in ws.iter_rows(min_row=7, max_row=8, min_col=1, max_col=34):
        for cell in row:
            text = str(cell.value or '').replace(' ', '').replace('\n', '')
            if text != '높이(mm)':
                continue
            cell.value = '높이\n(mm)'
            font = copy(cell.font)
            font.sz = 9
            cell.font = font
            alignment = copy(cell.alignment)
            alignment.wrapText = True
            alignment.shrinkToFit = False
            alignment.horizontal = 'center'
            alignment.vertical = 'center'
            cell.alignment = alignment


def printed_height(ws):
    default = ws.sheet_format.defaultRowHeight or 15
    return sum((ws.row_dimensions[row].height
                if ws.row_dimensions[row].height is not None else default)
               for row in printed_row_numbers(ws)
               if not ws.row_dimensions[row].hidden)


def compact_report_height(ws, factor=0.945):
    """Compress the printed rows and drawing heights without changing width."""
    if not 0 < factor <= 1:
        raise ValueError("factor must be between 0 and 1")
    rows = printed_row_numbers(ws)
    if not rows:
        return
    default_height = ws.sheet_format.defaultRowHeight or 15
    for row in rows:
        dim = ws.row_dimensions[row]
        dim.height = (dim.height if dim.height is not None else default_height) * factor
    for drawing in [*ws._images, *ws._charts]:
        anchor = drawing.anchor
        if isinstance(anchor, str):
            from openpyxl.utils.cell import coordinate_to_tuple
            if coordinate_to_tuple(anchor)[0] in rows:
                drawing.height *= factor
            continue
        start = getattr(anchor, '_from', None)
        if start is None or start.row + 1 not in rows:
            continue
        start.rowOff = round(start.rowOff * factor)
        end = getattr(anchor, 'to', None)
        if end is not None:
            end.rowOff = round(end.rowOff * factor)
        extent = getattr(anchor, 'ext', None)
        if extent is not None:
            extent.cy = round(extent.cy * factor)


def align_report_heights(sheets, factor=0.945, measured_widths=None, measured_heights=None):
    """Use one height for all report pages, without enlarging any page."""
    heights = [(ws, printed_height(ws)) for ws in sheets]
    heights = [(ws, height) for ws, height in heights if height > 0]
    if measured_heights is not None:
        heights = [(ws, measured_heights[ws.title]) for ws, _ in heights]
    if not heights:
        return
    scales = {}
    if measured_widths is not None:
        import math
        for ws, _ in heights:
            width = measured_widths[ws.title]
            available = 210 / 25.4 * 72 - (
                ws.page_margins.left + ws.page_margins.right
            ) * 72
            # Excel/printer rounding can push the last column across a page
            # boundary even when the nominal width fits. Reserve two percent.
            scales[ws.title] = max(10, min(400, math.floor(available / width * 100) - 2))
    else:
        scales = {ws.title: 100 for ws, _ in heights}
    target = min(height * scales[ws.title] / 100 for ws, height in heights) * factor
    for ws, height in heights:
        scale = scales[ws.title]
        compact_report_height(ws, target / (height * scale / 100))
        if measured_widths is not None:
            ws.sheet_properties.pageSetUpPr.fitToPage = False
            ws.page_setup.fitToWidth = None
            ws.page_setup.fitToHeight = None
            ws.page_setup.scale = scale


def measure_excel_print_widths(path, sheet_names, include_heights=False):
    """Read Excel's real font-dependent widths using a separate hidden instance."""
    import pythoncom
    import win32com.client
    from pathlib import Path
    pythoncom.CoInitialize()
    excel = book = None
    try:
        excel = win32com.client.DispatchEx('Excel.Application')
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AutomationSecurity = 3
        book = excel.Workbooks.Open(str(Path(path).resolve()), UpdateLinks=0, ReadOnly=True)
        areas = {name: book.Worksheets(name).Range(
            book.Worksheets(name).PageSetup.PrintArea) for name in sheet_names}
        widths = {name: float(area.Width) for name, area in areas.items()}
        if include_heights:
            return widths, {name: float(area.Height) for name, area in areas.items()}
        return widths
    finally:
        try:
            if book is not None:
                book.Close(False)
        finally:
            if excel is not None:
                excel.Quit()
            pythoncom.CoUninitialize()


def prepare_indication_row(ws, row, has_indication, has_tcg=False):
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row == merged.max_row == row and merged.min_col >= 14 and merged.max_col <= 24:
            ws.unmerge_cells(str(merged))
    if has_indication:
        for first, last in ((14, 15), (16, 17), (18, 19), (20, 21), (23, 24)):
            ws.merge_cells(start_row=row, end_row=row, start_column=first, end_column=last)
    elif has_tcg:
        ws.merge_cells(start_row=row, end_row=row, start_column=14, end_column=22)
        ws.merge_cells(start_row=row, end_row=row, start_column=23, end_column=24)
    else:
        ws.merge_cells(start_row=row, end_row=row, start_column=14, end_column=24)
    if not has_indication:
        return
    groups = [(14, 15), (16, 17), (18, 19), (20, 21), (22, 22), (23, 24)]
    dotted = Side(style='dotted', color='000000')
    for first, last in groups:
        for column in range(first, last + 1):
            cell = ws.cell(row, column)
            border = copy(cell.border)
            if column != 14:
                border.left = dotted if column == first else Side()
            if column != 24:
                border.right = dotted if column == last else Side()
            cell.border = border
