"""Allocate MT records to the cover and continuation template rows."""
from decimal import Decimal, InvalidOperation
from openpyxl.styles import Alignment
import re


def hot_tapping_points(items):
    """Combine daily work lengths for the same line and tapping point."""
    groups = {}
    for item in items:
        point = re.sub(r'^HT\s*-\s*', '', str(item.get('Joint', '')).strip(), flags=re.I).upper()
        groups.setdefault((str(item.get('Dwg', '')).strip(), point), []).append(item)
    result = []
    for (_, point), rows in groups.items():
        combined = dict(rows[0], Joint=f'HT-{point}' if point else '')
        original, repair, missing = mt_length_totals(rows)
        combined['Tested Length'] = '' if missing else str(original + repair)
        for key in ('Welder', 'Remarks', 'Interpretation'):
            combined[key] = '\n'.join(dict.fromkeys(str(row.get(key, '')).strip() for row in rows if str(row.get(key, '')).strip()))
        length_note = 'Length: 확인 필요' if missing else f'Length: {original + repair:.4f} M'
        combined['Remarks'] = '\n'.join(value for value in (combined['Remarks'], length_note) if value)
        outcomes = [str(row.get('Result', '')).strip().upper() for row in rows]
        if any(value in ('불합격', 'REJ', 'REJECT', 'REJECTED') for value in outcomes):
            combined['Result'] = 'Rej'
        elif all(value in ('합격', 'ACC', 'ACCEPT', 'ACCEPTED') for value in outcomes):
            combined['Result'] = 'Acc'
        else:
            combined['Result'] = ''
        result.append(combined)
    return result


def fit_equipment_labels(sheet):
    """Give bilingual equipment labels their full reserved label area."""
    for merged in list(sheet.merged_cells.ranges):
        if merged.min_row == merged.max_row == 12 and merged.min_col >= 2 and merged.max_col <= 9:
            sheet.unmerge_cells(str(merged))
    sheet.merge_cells('B12:I12')
    sheet['B12'] = 'Equipment 탐상기'
    for row, text in ((13, 'Maker 제조자 :'), (14, 'Model 형식 :'), (15, 'Sr. No. 일련번호 :')):
        sheet.merge_cells(start_row=row, end_row=row, start_column=2, end_column=3)
        sheet.cell(row, 2, text)
    for address in ('B12', 'B13', 'B14', 'B15'):
        sheet[address].alignment = Alignment(horizontal='left', vertical='center', shrink_to_fit=True)


def mt_length_totals(items):
    original = repair = Decimal('0')
    seen = set()
    missing = False
    for item in items:
        key = item.get('_source_id') or (item.get('Date'), item.get('Dwg'), item.get('Joint'))
        if key in seen:
            continue
        seen.add(key)
        value = str(item.get('Tested Length') or item.get('_source_tested_length') or '').strip()
        try:
            length = Decimal(value)
            if not length.is_finite() or length < 0:
                raise InvalidOperation
        except InvalidOperation:
            missing = True
            continue
        if item.get('_source_work_type') == '재검':
            repair += length
        else:
            original += length
    return original, repair, missing


def mt_data_positions(workbook, count, copy_page):
    cover = workbook.worksheets[0]
    continuation = workbook.worksheets[1] if len(workbook.worksheets) > 1 else None
    positions = [(cover, row) for row in range(29, 37)][:count]
    remaining = max(0, count - 8)
    if not remaining:
        if continuation is not None:
            workbook.remove(continuation)
        return positions
    if continuation is None:
        raise ValueError('추가 데이터용 MT(을) 시트가 필요합니다.')
    pages = [continuation]
    for number in range(1, (remaining + 28) // 29):
        pages.append(copy_page(number + 2))
    for sheet in pages:
        for row in range(10, 39):
            for merged in list(sheet.merged_cells.ranges):
                if merged.min_row == merged.max_row == row and merged.min_col == 1 and merged.max_col == 5:
                    sheet.unmerge_cells(str(merged))
                    sheet.merge_cells(start_row=row, end_row=row, start_column=2, end_column=5)
        rows = min(remaining, 29)
        positions.extend((sheet, row) for row in range(10, 10 + rows))
        remaining -= rows
    return positions
