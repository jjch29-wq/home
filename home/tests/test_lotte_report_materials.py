import datetime
from pathlib import Path
import sys
import tempfile
import unittest

import openpyxl

SRC = Path(__file__).resolve().parents[1] / 'src'
sys.path.insert(0, str(SRC))
from site_apps.lotte.src.daily_work_report_manager import DailyWorkReportManager
from site_apps.lotte.src.worker_assignment import include_shared_worker_methods


class ReportMaterialsTests(unittest.TestCase):
    def test_report_includes_mt_with_shared_ut_worker(self):
        attendance = [dict(names='주진철', company='함감속기', method='UT', ot_hours='09:00~18:00', ot_amount='0')]
        records = [dict(Date='2026-10-14', Site='함감속기', 업체명='함감속기', 검사방법='UT', User='주진철'),
                   dict(Date='2026-10-14', Site='함감속기', 업체명='함감속기', 검사방법='MT', User='')]
        attendance = include_shared_worker_methods(attendance, records)
        manager = DailyWorkReportManager(str(SRC / 'site_apps/lotte/resources/Template_DailyWorkReport.xlsx'))
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / 'report.xlsx'
            manager.generate_report({'date': datetime.date(2026, 10, 14), 'ot_status': attendance}, str(output))
            wb = openpyxl.load_workbook(output)
            values = [cell.value for row in wb.active for cell in row]
            self.assertIn('UT, MT', values)
            self.assertEqual(values.count('주진철'), 1)
            wb.close()

    def test_fluorescent_powder_has_separate_row_and_pt_is_preserved(self):
        manager = DailyWorkReportManager(str(SRC / 'site_apps/lotte/resources/Template_DailyWorkReport.xlsx'))
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / 'report.xlsx'
            manager.generate_report({'date': datetime.date(2026, 10, 14),
                'materials': {'MT WHITE': {'used': 1}, 'MT FLUORESCENT': {'used': 2.5},
                              'PT Penetrant': {'used': 3}}}, str(output))
            wb = openpyxl.load_workbook(output)
            ws = wb.active
            rows = {cell.value: cell.row for cells in ws for cell in cells
                    if cell.value in ('백색페인트', '형광자분', '침투액')}
            self.assertEqual(ws.cell(rows['형광자분'], 13).value, 2.5)
            self.assertEqual(ws.cell(rows['백색페인트'], 13).value, 1)
            self.assertEqual(ws.cell(rows['침투액'], 13).value, 3)
            self.assertLess(rows['형광자분'], rows['침투액'])
            bottom = next(cell.row for cells in ws for cell in cells if cell.value == '현상액')
            for row in range(42, bottom + 1):
                self.assertEqual(ws.cell(row, 2).border.left.style, 'thin')
                self.assertEqual(ws.cell(row, 19).border.right.style, 'thin')
            for col in range(2, 20):
                self.assertEqual(ws.cell(42, col).border.top.style, 'thin')
                self.assertEqual(ws.cell(bottom, col).border.bottom.style, 'thin')
            wb.close()


if __name__ == '__main__':
    unittest.main()
