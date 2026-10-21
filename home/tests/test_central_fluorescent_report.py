"""Verify the native report pipeline keeps fluorescent powder separate."""
import contextlib
import datetime
import io
from pathlib import Path
import sys
import tempfile
import unittest

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'home/src'))
from site_apps.central.src.services.report_materials import collect_chemical_usage
from site_apps.central.src.daily_work_report_manager import DailyWorkReportManager
from site_apps.central.src.services.report_inspection import collect_inspection_details


class Entry:
    def __init__(self, value):
        self.value = value

    def get(self):
        return self.value


class FluorescentReportTests(unittest.TestCase):
    def test_rt_detail_uses_ori_points_and_reject_films_excludes_rep(self):
        records = pd.DataFrame([
            {'검사방법': 'PAUT', '관경(Inch)': '400A', '검사구분': 'ORI', '조인트수': 7, 'Usage': 8.9369},
            {'검사방법': 'RT', '관경(Inch)': '100A', '검사구분': 'ORI', '조인트수': 26, 'Usage': 80, '불량수': 2},
            {'검사방법': 'RT', '관경(Inch)': '100A', '검사구분': 'REP', '조인트수': 2, 'Usage': 6},
        ])
        details = collect_inspection_details(records, {})
        rt = next(item for item in details if item['method'] == 'RT')
        self.assertEqual(rt['points'], 26)
        self.assertEqual(rt['reject_films'], 2)
        self.assertEqual(rt['quantity'], 80)
        self.assertEqual(records.iloc[2]['조인트수'], 2)
        paut = next(item for item in details if item['method'] == 'PAUT')
        self.assertEqual(paut['points'], 7)
        self.assertEqual(paut['quantity'], 8.9369)

    def test_paut_detail_keeps_legacy_point_recovery(self):
        details = collect_inspection_details(pd.DataFrame([
            {'검사방법': 'PAUT', '관경(Inch)': '250A', 'Usage': 1.6802}]), {'250A': .8401})
        self.assertEqual(details[0]['points'], 2)

    def test_saved_materials_remain_separate_and_aliases_do_not_double(self):
        records = pd.DataFrame([
            {'NDT_형광자분': 2.5, 'DT_형광자분': 2.5, 'NDT_백색페인트': 3, 'NDT_흑색자분': 4},
            {'NDT_형광자분': 1, 'NDT_백색페인트': 2, 'NDT_흑색자분': 3},
        ])
        result = collect_chemical_usage(records)
        self.assertEqual(result['MT FLUORESCENT']['used'], 3.5)
        self.assertEqual(result['MT WHITE']['used'], 5)
        self.assertEqual(result['MT 7C-BLACK']['used'], 7)

    def test_ui_fallback_uses_names_instead_of_widget_positions(self):
        entries = [{'흑색자분': Entry('4'), '형광자분': Entry('2'), '백색페인트': Entry('3')},
                   {'백색페인트': Entry('1'), '형광자분': Entry('0.5')}]
        result = collect_chemical_usage(pd.DataFrame(), entries)
        self.assertEqual(result['MT FLUORESCENT']['used'], 2.5)
        self.assertEqual(result['MT WHITE']['used'], 4)
        self.assertEqual(result['MT 7C-BLACK']['used'], 4)
        self.assertEqual(collect_chemical_usage(pd.DataFrame([{'NDT_형광자분': 0}]), entries), {})

    def test_historical_named_material_is_not_white_paint(self):
        result = collect_chemical_usage(pd.DataFrame([{'품목명': 'MT 형광자분', 'Usage': 2}]))
        self.assertEqual(result, {'MT FLUORESCENT': {'used': 2}})

    def test_generated_report_includes_mt_row_and_preserves_pt_and_print_area(self):
        template = ROOT / 'home/src/site_apps/central/resources/Template_DailyWorkReport.xlsx'
        for rt_count, ot_count, method_count in ((1, 0, 2), (4, 3, 5)):
            with self.subTest(rt=rt_count, ot=ot_count, methods=method_count):
                materials = {f'RT_ROW_{i}': {'name': f'T200-{i}', 'spec': '3⅓×12', 'unit': '매', 'used': 80}
                             for i in range(rt_count)}
                materials.update({key: {'used': used} for key, used in (
                    ('MT WHITE', 3), ('MT 7C-BLACK', 5), ('MT FLUORESCENT', 7.5),
                    ('PT Penetrant', 11), ('PT Cleaner', 13), ('PT Developer', 17))})
                data = {'date': datetime.date(2026, 10, 19), 'materials': materials,
                        'methods': {key: {'qty': 1} for key in ('RT', 'PAUT', 'MT', 'PT', 'UT')[:method_count]},
                        'ot_status': [{'names': f'작업자{i}', 'ot_hours': '3', 'ot_amount': 12000} for i in range(ot_count)]}
                data['inspection_details'] = collect_inspection_details(pd.DataFrame([
                    {'검사방법': 'RT', '관경(Inch)': '100A', '검사구분': 'ORI', '조인트수': 26, 'Usage': 80, '불량수': 2},
                    {'검사방법': 'RT', '관경(Inch)': '100A', '검사구분': 'REP', '조인트수': 2, 'Usage': 0},
                ]), {})
                with tempfile.TemporaryDirectory(dir=ROOT / 'tmp') as directory:
                    output = Path(directory) / 'daily.xlsx'
                    with contextlib.redirect_stdout(io.StringIO()):
                        DailyWorkReportManager(str(template)).generate_report(data, str(output))
                    workbook = openpyxl.load_workbook(output)
                    sheet = workbook.active
                    try:
                        self.assertAlmostEqual(sheet.page_margins.left, 10 / 25.4)
                        self.assertAlmostEqual(sheet.page_margins.right, 10 / 25.4)
                        self.assertAlmostEqual(sheet.page_margins.top, 10 / 25.4)
                        widths = [sheet.column_dimensions[openpyxl.utils.get_column_letter(c)].width for c in range(1, 20)]
                        self.assertAlmostEqual(sum(w * 7 + 5 for w in widths), 190 / 25.4 * 96)
                        height = sum(sheet.row_dimensions[r].height or 15 for r in range(1, sheet.max_row + 1)
                                     if not sheet.row_dimensions[r].hidden)
                        self.assertLessEqual(height, 277 / 25.4 * 72 + 0.01)
                        self.assertTrue(sheet.print_options.horizontalCentered)
                        self.assertFalse(sheet.print_options.verticalCentered)
                        expected = {'백색페인트': 3, '흑색자분': 5, '형광자분': 7.5, '침투액': 11, '세척액': 13, '현상액': 17}
                        rows = {sheet.cell(row, 4).value: row for row in range(40, sheet.max_row + 1)
                                if sheet.cell(row, 4).value in expected}
                        self.assertEqual(set(rows), set(expected))
                        for name, used in expected.items():
                            row = rows[name]
                            self.assertEqual(sheet.cell(row, 13).value, used)
                            self.assertEqual(sheet.cell(row, 8).value, 'CAN')
                            self.assertFalse(sheet.row_dimensions[row].hidden)
                        self.assertEqual(rows['형광자분'], rows['흑색자분'] + 1)
                        self.assertEqual(rows['침투액'], rows['형광자분'] + 1)
                        self.assertIn(f'$S${rows["현상액"] + 1}', str(sheet.print_area))
                        rt_names = [sheet.cell(row, 4).value for row in range(40, sheet.max_row + 1)
                                    if str(sheet.cell(row, 4).value).startswith('T200-')]
                        self.assertEqual(rt_names, [f'T200-{i}' for i in range(rt_count)])
                        content = '\n'.join(str(cell.value or '') for row in sheet for cell in row)
                        self.assertIn('100A: 26P / REJ 2매 / 검사매수 80매', content)
                        self.assertNotIn('28 POINT', content)
                    finally:
                        workbook.close()


if __name__ == '__main__':
    unittest.main()
