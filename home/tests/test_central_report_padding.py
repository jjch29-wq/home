"""Padding belongs to the preceding report page, never a blank page."""
import ast
import sys
import unittest
import os
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.worksheet.pagebreak import Break

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from site_apps.central.src.services.monthly_report_manager import MonthlyReportManager


class ReportPaddingTest(unittest.TestCase):
    def test_whole_detail_table_moves_with_parent_without_orphan_title(self):
        sheet = openpyxl.Workbook().active
        sheet.cell(1, 6, '월간 용역 진도 보고서')
        sheet.cell(5, 6, '3.0 비파괴검사현황')
        sheet.cell(19, 2, '2.0 비파괴검사 세부현황')
        sheet.cell(20, 3, '2.1 비파괴검사 세부현황(PAUT)')
        sheet.cell(21, 2, '업체')
        sheet.cell(23, 2, 'DATA')
        manager = MonthlyReportManager('unused.xlsx')
        inserted = manager._move_table_to_continuation(sheet, 20, 1, 21)
        self.assertEqual(inserted, 9)
        self.assertIsNone(sheet.cell(19, 2).value)
        self.assertEqual(sheet.cell(28, 2).value, '2.0 비파괴검사 세부현황')
        self.assertEqual(sheet.cell(29, 3).value, '2.1 비파괴검사 세부현황(PAUT)')
        self.assertEqual(sheet.cell(30, 2).value, '업체')
        self.assertEqual(sheet.cell(32, 2).value, 'DATA')

    def test_real_toc_is_kept_and_general_and_photo_headers_are_corrected(self):
        sheet = openpyxl.Workbook().active
        for start in (1, 30, 60, 90):
            sheet.merge_cells(start_row=start, start_column=6,
                              end_row=start + 3, end_column=15)
        sheet.cell(5, 6, '목 차')
        sheet.cell(9, 2, 'Section Title')
        sheet.cell(34, 6, '목 차')
        sheet.cell(38, 2, '1. 용역명')
        sheet.cell(64, 6, '6.0 사진대지')
        sheet.cell(94, 6, '6.0 사진대지')
        sheet.print_area = 'B1:W89'
        manager = MonthlyReportManager('unused.xlsx')
        manager._fix_report_section_header_labels(sheet)
        self.assertEqual(sheet.cell(5, 6).value, '목 차')
        self.assertEqual(sheet.cell(34, 6).value, '1.0 용역업무일반')
        self.assertEqual(sheet.cell(64, 6).value, '7.0 사진대지')
        self.assertIn('1     of     1', sheet.cell(64, 16).value)
        self.assertIsNone(sheet.cell(94, 16).value)

    def test_continuation_parent_uses_first_page_header_gap(self):
        sheet = openpyxl.Workbook().active
        for start in (1, 50):
            sheet.cell(start, 6, '월간 용역 진도 보고서')
        sheet.cell(9, 2, '1.0 비파괴검사사용량 집계표')
        sheet.row_dimensions[8].height = 5
        sheet.cell(58, 2, '1.2 SECTION별 검사현황')
        sheet.row_dimensions[57].height = 20
        sheet.cell(59, 3, '1.2.3 방사선투과검사 (계속)')
        manager = MonthlyReportManager('unused.xlsx')
        manager._align_ndt_parent_top_gaps(sheet)
        reference_gap = sum(
            sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight
            for row in (7, 8)
        )
        continuation_gap = sum(
            sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight
            for row in (56, 57)
        )
        self.assertAlmostEqual(continuation_gap, reference_gap)
        self.assertTrue(all(sheet.row_dimensions[row].height > 0
                            for row in (56, 57)))
        self.assertEqual(sheet.cell(59, 3).value, '1.2.3 방사선투과검사 (계속)')

    def test_continuation_gap_includes_last_blank_header_row(self):
        sheet = openpyxl.Workbook().active
        for start in (1, 50):
            sheet.cell(start, 6, '월간 용역 진도 보고서')
        sheet.cell(8, 2, '1.0 비파괴검사사용량 집계표')
        sheet.row_dimensions[7].height = 22.5
        sheet.cell(59, 2, '1.2 SECTION별 검사현황')
        sheet.row_dimensions[56].height = 28.5
        sheet.row_dimensions[57].height = 13.5
        sheet.row_dimensions[58].height = 13.5
        manager = MonthlyReportManager('unused.xlsx')
        manager._align_ndt_parent_top_gaps(sheet)
        heights = [sheet.row_dimensions[row].height for row in (56, 57, 58)]
        self.assertAlmostEqual(sum(heights), 22.5)
        self.assertTrue(all(value > 0 for value in heights))
        self.assertFalse(any(sheet.row_dimensions[row].hidden
                             for row in (56, 57, 58)))

    def test_section_parent_matches_summary_title_font_and_left_edge(self):
        from copy import copy
        from openpyxl.styles import Font, Alignment
        sheet = openpyxl.Workbook().active
        reference = sheet.cell(10, 2, '1.0 비파괴검사사용량 집계표')
        reference.font = Font(name='Arial', size=10, bold=True)
        reference.alignment = Alignment(horizontal='left', vertical='top')
        sheet.cell(30, 3, '1.2 SECTION별 검사현황')
        sheet.merge_cells('C30:W30')
        sheet.cell(31, 3, '1.2.1 PAUT')
        manager = MonthlyReportManager('unused.xlsx')
        manager._normalize_ndt_parent_caption_style(sheet)
        title = sheet.cell(30, 2)
        self.assertEqual(title.value, '1.2 SECTION별 검사현황')
        self.assertEqual(copy(title.font), copy(reference.font))
        self.assertEqual(copy(title.alignment), copy(reference.alignment))
        self.assertIsNone(sheet.cell(30, 3).value)
        self.assertIn('B30:W30', {str(m) for m in sheet.merged_cells.ranges})
        self.assertEqual(sheet.cell(31, 3).value, '1.2.1 PAUT')

    def test_last_rt_page_repeats_detail_parent_and_keeps_pt(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.cell(1, 6, '월간 용역 진도 보고서')
        sheet.cell(5, 6, '3.0 비파괴검사현황')
        sheet.cell(10, 3, '2.3 비파괴검사 세부현황(RT) (계속)')
        sheet.cell(11, 2, '업체')
        sheet.cell(20, 2, 'RT 68')
        sheet.cell(21, 2, 'TOTAL')
        sheet.cell(22, 3, '2.4 비파괴검사 세부현황(PT)')
        sheet.cell(40, 2, '2.0 비파괴검사 세부현황')
        inserted = manager._insert_table_continuation(sheet, 20, 1, 11)
        self.assertEqual(inserted, 12)
        self.assertEqual(sheet.cell(28, 2).value, '2.0 비파괴검사 세부현황')
        self.assertEqual(sheet.cell(29, 3).value,
                         '2.3 비파괴검사 세부현황(RT) (계속)')
        self.assertEqual(sheet.cell(32, 2).value, 'RT 68')
        self.assertEqual(sheet.cell(33, 2).value, 'TOTAL')
        self.assertEqual(sheet.cell(34, 3).value, '2.4 비파괴검사 세부현황(PT)')

    def test_detail_parent_repeats_above_paut_not_between_paut_and_mt(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.cell(1, 6, '월간 용역 진도 보고서')
        sheet.cell(5, 6, '3.0 비파괴검사현황')
        sheet.cell(10, 2, '2.0 비파괴검사 세부현황')
        sheet.cell(11, 3, '2.1 비파괴검사 세부현황(PAUT)')
        sheet.cell(12, 2, '업체')
        sheet.cell(30, 2, 'PAUT DATA')
        sheet.cell(31, 2, 'TOTAL')
        sheet.cell(32, 2, '2.0 비파괴검사 세부현황')
        sheet.cell(33, 3, '2.2 비파괴검사 세부현황(MT)')
        manager._remove_duplicate_ndt_detail_captions(sheet, 1, 40)
        inserted = manager._insert_table_continuation(sheet, 30, 1, 12)
        self.assertEqual(inserted, 12)
        self.assertEqual(sheet.cell(38, 2).value, '2.0 비파괴검사 세부현황')
        self.assertEqual(sheet.cell(39, 3).value,
                         '2.1 비파괴검사 세부현황(PAUT) (계속)')
        self.assertEqual(sheet.cell(42, 2).value, 'PAUT DATA')
        self.assertIsNone(sheet.cell(44, 2).value)
        self.assertEqual(sheet.cell(45, 3).value, '2.2 비파괴검사 세부현황(MT)')

    def test_section_parent_is_repeated_above_rt_continuation(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.cell(1, 6, '월간 용역 진도 보고서')
        sheet.cell(5, 6, '3.0 비파괴검사현황')
        sheet.cell(10, 3, '1.2 SECTION별 검사현황')
        sheet.row_dimensions[10].height = 18
        sheet.merge_cells('C10:H10')
        sheet.cell(20, 3, '1.2.3 방사선투과검사')
        sheet.cell(21, 2, '업체')
        sheet.cell(30, 2, 'TOTAL')

        inserted = manager._insert_table_continuation(sheet, 30, 1, 21)

        self.assertEqual(inserted, 12)
        self.assertEqual(sheet.cell(38, 3).value, '1.2 SECTION별 검사현황')
        self.assertEqual(sheet.row_dimensions[38].height, 18)
        self.assertIn('C38:H38', {str(m) for m in sheet.merged_cells.ranges})
        self.assertEqual(sheet.cell(39, 3).value, '1.2.3 방사선투과검사 (계속)')
        self.assertEqual(sheet.cell(42, 2).value, 'TOTAL')

    def test_summary_caption_moves_up_without_moving_quantity_table(self):
        sheet = openpyxl.Workbook().active
        sheet.row_dimensions[10].height = 20
        sheet.row_dimensions[11].height = 14
        sheet.cell(11, 2, '1.0 비파괴검사사용량 집계표')
        sheet.cell(12, 3, '1.1 비파괴검사 물량표')
        sheet.cell(13, 2, '방법')
        sheet.row_dimensions[20].height = 10
        sheet.row_dimensions[21].height = 14
        sheet.cell(21, 3, '1.2 SECTION별 검사현황')
        sheet.cell(22, 3, '1.2.1 위상배열초음파탐상검사')
        sheet.row_breaks.append(Break(id=60))
        manager = MonthlyReportManager('unused.xlsx')

        self.assertTrue(manager._space_ndt_summary_caption(sheet))

        self.assertEqual(sheet.row_dimensions[10].height, 14)
        self.assertEqual(sheet.row_dimensions[11].height, 20)
        self.assertEqual(sheet.cell(11, 2).alignment.vertical, 'top')
        self.assertEqual(sheet.cell(12, 3).value, '1.1 비파괴검사 물량표')
        self.assertEqual(sheet.cell(13, 2).value, '방법')
        self.assertEqual(sheet.row_dimensions[20].height, 10)
        self.assertEqual(sheet.row_dimensions[21].height, 14)
        manager._space_ndt_section_caption(sheet)
        self.assertEqual(sheet.row_dimensions[20].height, 10)
        self.assertEqual(sheet.row_dimensions[21].height, 18)
        self.assertEqual(sheet.cell(21, 3).alignment.vertical, 'top')
        self.assertEqual(sheet.cell(22, 3).value, '1.2.1 위상배열초음파탐상검사')
        self.assertEqual([b.id for b in sheet.row_breaks.brk], [60])

    def test_only_ndt_header_spacer_is_increased(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.sheet_format.defaultRowHeight = 15

        sheet.cell(10, 6, '월간용역진도보고서')
        sheet.cell(14, 6, '3.0 비파괴검사현황')
        sheet.cell(30, 6, '월간용역진도보고서')
        sheet.cell(34, 6, '4.0 용접결함, 불량률 현황')
        sheet.row_dimensions[16].height = 13.5
        sheet.row_dimensions[36].height = 13.5

        spacer_row = manager._lower_ndt_page_titles(sheet)

        self.assertEqual(spacer_row, 16)
        self.assertEqual(sheet.row_dimensions[16].height, 18.5)
        self.assertEqual(sheet.row_dimensions[36].height, 13.5)

    def test_detail_rows_use_one_horizontal_border_per_boundary(self):
        from openpyxl.styles import Border, Side

        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        thin = Side(style='thin')
        upper = sheet.cell(10, 2)
        lower = sheet.cell(11, 2)
        upper.border = Border(top=thin, bottom=thin)
        lower.border = Border(top=thin, bottom=thin)

        manager._keep_single_lower_horizontal_border(lower)

        self.assertEqual(upper.border.bottom.style, 'thin')
        self.assertIsNone(lower.border.top.style)
        self.assertEqual(lower.border.bottom.style, 'thin')

    def test_tiny_blank_rows_between_detail_records_bridge_vertical_grid(self):
        from openpyxl.styles import Border, Side

        sheet = openpyxl.Workbook().active
        sheet.cell(10, 3, 47)
        sheet.cell(13, 3, 48)
        sheet.row_dimensions[11].height = 1.95
        sheet.row_dimensions[12].height = 1.95
        thin = Side(style='thin')
        medium = Side(style='medium')
        sheet.cell(10, 2).border = Border(left=medium, right=thin, bottom=thin)
        sheet.cell(10, 3).border = Border(left=thin, right=thin, bottom=thin)

        bridged = MonthlyReportManager._bridge_blank_detail_row_borders(sheet)

        self.assertEqual(bridged, [11, 12])
        for row in (11, 12):
            self.assertEqual(sheet.cell(row, 2).border.left.style, 'medium')
            self.assertEqual(sheet.cell(row, 2).border.right.style, 'thin')
            self.assertEqual(sheet.cell(row, 3).border.left.style, 'thin')
            self.assertEqual(sheet.cell(row, 3).border.right.style, 'thin')
            self.assertIsNone(sheet.cell(row, 2).border.top.style)
        self.assertIsNone(sheet.cell(11, 2).border.bottom.style)
        self.assertEqual(sheet.cell(12, 2).border.bottom.style, 'thin')

    def test_merged_general_pages_place_equipment_table_below_heading(self):
        from types import SimpleNamespace

        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        for break_id in (235, 278, 325):
            sheet.row_breaks.append(Break(id=break_id))
        sheet.cell(306, 2, '거. PAUT 검사장비 최소사양')
        anchor = SimpleNamespace(
            _from=SimpleNamespace(row=306, col=2),
            to=SimpleNamespace(row=318, col=20, rowOff=1),
        )
        image = SimpleNamespace(anchor=anchor, width=686.0, height=294.0)
        sheet._images = [image]

        manager._merge_pages7_and8(sheet)

        # Deleting the repeated header moves row 306 to 299. The table then
        # starts on the following blank row 300 instead of covering paragraphs.
        self.assertEqual(image.anchor._from.row + 1, 300)
        self.assertEqual(image.anchor.to.row + 1, 307)
        self.assertEqual(image.height, 180)

    def test_rt_summary_total_uses_free_row_on_previous_page(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.row_breaks.append(Break(id=20))
        sheet.row_breaks.append(Break(id=50))
        sheet.cell(25, 2, '1.2 SECTION별 검사현황')
        sheet.cell(26, 3, '1.2.3 방사선투과검사 (계속)')
        sheet.cell(27, 6, 'Line No.')
        sheet.cell(28, 17, "ORI'")
        sheet.cell(29, 2, 'TOTAL')
        sheet.cell(29, 17, 281)
        sheet.merge_cells('B29:P29')
        sheet.cell(32, 3, '1.2.4 침투탐상검사')

        changed = manager._pull_rt_summary_total_to_previous_page(sheet)

        self.assertTrue(changed)
        self.assertEqual(sheet.cell(21, 2).value, 'TOTAL')
        self.assertEqual(sheet.cell(21, 17).value, 281)
        self.assertIn('B21:P21', {str(merged) for merged in sheet.merged_cells.ranges})
        self.assertEqual([brk.id for brk in sheet.row_breaks.brk], [21, 47])
        self.assertEqual(sheet.cell(29, 3).value, '1.2.4 침투탐상검사')
        self.assertFalse(any(
            '(계속' in ''.join(
                str(sheet.cell(row=row, column=col).value or '')
                for col in range(1, 24)
            )
            for row in range(1, 40)
        ))

    def test_equipment_table_anchor_is_fixed_without_korean_text_matching(self):
        from types import SimpleNamespace

        sheet = openpyxl.Workbook().active
        sheet.cell(299, 2, '7. PAUT minimum specification')
        anchor = SimpleNamespace(
            _from=SimpleNamespace(row=292, col=2),
            to=SimpleNamespace(row=304, col=20),
        )
        image = SimpleNamespace(anchor=anchor, width=686.0, height=180.0)
        sheet._images = [image]

        target = MonthlyReportManager._place_general_equipment_table_image(sheet)

        self.assertEqual(target, 300)
        self.assertEqual(image.anchor._from.row + 1, 300)
        self.assertEqual(image.anchor.to.row + 1, 312)

    def test_ndt_summary_compacts_titles_without_shrinking_records(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.cell(10, 2, '1.0 Summary')
        sheet.cell(11, 3, '1.2.1 PAUT')
        sheet.cell(12, 2, 'Contractor')
        for row in range(10, 13):
            sheet.row_dimensions[row].height = 16.5

        manager._compact_ndt_summary_title_rows(sheet, 10, 13)

        self.assertEqual(sheet.row_dimensions[10].height, 13.75)
        self.assertEqual(sheet.row_dimensions[11].height, 13.75)
        self.assertEqual(sheet.row_dimensions[12].height, 16.5)

    def test_unmarked_duplicate_ndt_heading_is_rebuilt_as_continuation(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active

        # The parent section caption is two rows above the table heading and
        # must not be mistaken for the subsection caption.
        sheet.cell(4, 3, '2.0 Details (계속)')
        sheet.cell(5, 3, '2.2 MT (계속)')
        sheet.cell(6, 6, 'Line No.')
        sheet.cell(10, 3, '2.3 RT')
        sheet.cell(11, 6, 'Line No.')
        sheet.cell(20, 3, '2.3 RT')
        sheet.cell(21, 6, 'Line No.')
        sheet.cell(30, 3, '2.3 RT (계속)')
        sheet.cell(31, 6, 'Line No.')
        sheet.cell(40, 3, '2.4 PT')
        sheet.cell(41, 6, 'Line No.')

        stale, first_continuations = (
            manager._stale_ndt_continuation_blocks(sheet, 1, 50)
        )

        self.assertEqual(stale, [(20, 3), (30, 3)])
        self.assertEqual(first_continuations, [(5, 3, '2.2 MT')])

    def test_report_height_scale_matches_excel_pdf_calibration(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active

        self.assertAlmostEqual(
            manager._report_height_scale(sheet),
            manager._paut_print_scale(sheet) * 0.91,
        )
        self.assertAlmostEqual(
            manager._ndt_reflow_height_scale(sheet, page_number=1),
            manager._paut_print_scale(sheet) * 0.905,
        )
        self.assertAlmostEqual(
            manager._ndt_reflow_height_scale(sheet, page_number=2),
            manager._paut_print_scale(sheet) * 0.905,
        )
        self.assertAlmostEqual(
            manager._ndt_reflow_height_scale(sheet, page_number=3),
            manager._report_height_scale(sheet),
        )

    def test_first_two_ndt_pages_use_same_bottom_boundary(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active

        self.assertEqual(
            manager._ndt_reflow_height_scale(sheet, page_number=1),
            manager._ndt_reflow_height_scale(sheet, page_number=2),
        )

    def test_header_only_ndt_page_is_removed(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.cell(11, 6, '3.0 鍮꾪뙆愿닿??ы쁽??')
        sheet.cell(15, 2, '1.2 SECTION')
        sheet.cell(21, 6, '3.0 鍮꾪뙆愿닿??ы쁽??')
        sheet.cell(25, 6, 'Line No.')
        sheet.cell(27, 3, 1)
        sheet.row_breaks.append(Break(id=10))
        sheet.row_breaks.append(Break(id=20))
        sheet.row_breaks.append(Break(id=30))

        removed = manager._remove_header_only_ndt_pages(sheet)

        self.assertEqual(removed, 1)
        self.assertEqual(sorted({b.id for b in sheet.row_breaks.brk}), [10, 20])
        self.assertIn('3.0', str(sheet.cell(11, 6).value))
        self.assertEqual(sheet.cell(17, 3).value, 1)

    def test_ndt_reflow_uses_bottom_margin_without_double_footer_gap(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.page_margins.top = sheet.page_margins.bottom = 0.35
        sheet.page_margins.footer = 0.3
        for row in range(1, 110):
            sheet.row_dimensions[row].height = 9
        previous = manager._first_print_overflow_row(sheet, 1, 109, 90.5)
        expanded = manager._first_print_overflow_row(
            sheet, 1, 109, 90.5, reserve_footer=False)
        self.assertGreater(expanded, previous)
        capacity = (11.69 - 0.35 - 0.35) * 72 / 0.905
        self.assertLessEqual((expanded - 1) * 9, capacity)
        self.assertGreater(expanded * 9, capacity)

    def test_report_header_contract_title_uses_thin_top_border(self):
        from openpyxl.styles import Border, Side

        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        medium = Side(style='medium')
        sheet.merge_cells('F301:O304')
        sheet.merge_cells('P301:W302')
        for column in range(6, 16):
            sheet.cell(301, column).border = Border(top=medium)

        # Similar body geometry without the document-number block must remain
        # untouched.
        sheet.merge_cells('F320:O323')
        sheet.cell(320, 6).border = Border(top=medium)

        manager._normalize_report_header_top_borders(sheet)

        self.assertTrue(all(
            sheet.cell(301, column).border.top.style == 'thin'
            for column in range(6, 16)
        ))
        self.assertEqual(sheet.cell(320, 6).border.top.style, 'medium')

    def test_monthly_paut_subtotal_includes_250a_night_shift(self):
        source = Path(
            'home/src/site_apps/central/src/services/monthly_report_manager.py'
        )
        tree = ast.parse(source.read_text(encoding='utf-8'))
        paut_keys = next(
            ast.literal_eval(node.value)
            for node in ast.walk(tree)
            if isinstance(node, ast.Assign)
            and any(
                isinstance(target, ast.Name) and target.id == 'paut_keys'
                for target in node.targets
            )
        )

        self.assertIn('PAUT_250A-야간', paut_keys)

    def test_missing_cl_column_is_added_and_counted_without_losing_quantities(self):
        import io
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.cell(1, 2, '1.1 방사선투과검사 불량률')
        sheet.cell(2, 4, '발생 결함 내역')
        sheet.merge_cells('D2:O2')
        for col, label in ((16, '불합격부 재검사 물량(A)'),
                           (17, '당초 시공물량(B)'),
                           (18, '용접 불량률(A/B%)')):
            sheet.cell(2, col, label)
            sheet.merge_cells(start_row=2, end_row=3, start_column=col, end_column=col)
        for col, label in enumerate(('ID', '성명', 'P', 'SI', 'LF', 'IP', 'RC',
                                      'EC', 'EP', 'TI', 'IUC', 'EUC', 'G', 'M'), 2):
            sheet.cell(3, col, label)
        sheet.cell(6, 2, '총계')
        sheet.cell(9, 16, '다른 표 보존')
        manager._write_welder_defect_summary(
            sheet, {'W-1': {'P': 1, 'CL': 1}}, {'W-1': 2}, {'W-1': 117},
            '1.1', ('방사선',))
        self.assertEqual(sheet.cell(3, 16).value, 'CL')
        self.assertEqual(sheet.cell(4, 16).value, 1)
        self.assertEqual(sheet.cell(6, 16).value, 1)
        self.assertEqual(sheet.cell(4, 4).value, 1)
        self.assertEqual(sheet.cell(9, 16).value, '다른 표 보존')
        manager._widen_weld_defect_table(sheet)
        buffer = io.BytesIO()
        sheet.parent.save(buffer)
        buffer.seek(0)
        result = openpyxl.load_workbook(buffer).active
        self.assertEqual(result.cell(4, 16).value, 1)
        self.assertEqual(result.cell(4, 17).value, 2)
        self.assertEqual(result.cell(4, 19).value, 117)
        self.assertAlmostEqual(result.cell(4, 22).value, 2 / 117)
        self.assertIn('D2:P2', {str(m) for m in result.merged_cells.ranges})

    def test_photo_frame_tail_fits_without_resizing_photos_or_caption(self):
        from types import SimpleNamespace
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.page_setup.scale = 100
        sheet.page_margins.top = sheet.page_margins.bottom = .75
        sheet.page_margins.footer = .3
        sheet.print_area = 'B1:W45'
        sheet.cell(5, 6, '6.0 사진대지')
        sheet.cell(41, 2, '사진 설명')
        for r in range(1, 46):
            sheet.row_dimensions[r].height = 18
        anchor = SimpleNamespace(_from=SimpleNamespace(row=10), to=SimpleNamespace(row=39))
        sheet._images.append(SimpleNamespace(anchor=anchor))
        sheet.row_breaks.append(Break(id=45))
        scale = manager._report_height_scale(sheet)
        self.assertIsNotNone(manager._first_print_overflow_row(sheet, 1, 45, scale))
        manager._fit_photo_page_trailing_rows(sheet)
        self.assertIsNone(manager._first_print_overflow_row(sheet, 1, 45, scale))
        self.assertTrue(all(sheet.row_dimensions[r].height == 18 for r in range(1, 42)))
        self.assertEqual(sheet.cell(41, 2).value, '사진 설명')
        self.assertEqual((anchor._from.row, anchor.to.row), (10, 39))
        self.assertEqual([b.id for b in sheet.row_breaks.brk], [45])

    def test_rt_gap_compaction_keeps_last_record_and_total_on_page(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.page_setup.scale = 100
        for r in range(1, 39):
            sheet.row_dimensions[r].height = 20
        sheet.cell(15, 2, 'MT TOTAL')
        sheet.cell(20, 2, '2.3 비파괴검사 세부현황(RT)')
        sheet.cell(21, 6, 'Line No.')
        for r in range(23, 38):
            sheet.cell(r, 3, r - 22)
        sheet.cell(38, 2, 'TOTAL')
        self.assertIsNotNone(manager._detail_page_overflow_row(sheet, 23, 15))
        manager._compact_rt_leading_gap(sheet, 21)
        self.assertIsNone(manager._detail_page_overflow_row(sheet, 23, 15))
        self.assertEqual([sheet.row_dimensions[r].height for r in range(16, 20)], [2] * 4)
        for r in (15, 20, 21, 22, 23, 37, 38):
            self.assertEqual(sheet.row_dimensions[r].height, 20)
        self.assertEqual(sheet.cell(37, 3).value, 15)
        self.assertEqual(sheet.cell(38, 2).value, 'TOTAL')
        self.assertFalse(sheet.row_breaks.brk)

    def test_defect_tables_stay_together_by_reclaiming_empty_gap(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.page_setup.scale = 100
        sheet.cell(1, 6, '월간 용역 진도 보고서')
        sheet.cell(5, 6, '4.0 용접결함, 불량률 현황')
        sheet.cell(14, 2, '총계')
        sheet.cell(35, 2, '1.2 PAUT 용접 불량률')
        sheet.cell(38, 2, 'ID')
        sheet.cell(38, 3, '성명')
        sheet.cell(47, 2, '총계')
        for r in range(1, 48):
            sheet.row_dimensions[r].height = 20
        sheet.row_breaks.append(Break(id=47))
        originals = {f'W-{i}': 100 for i in range(8)}
        manager._write_welder_defect_summary(sheet, {}, {}, originals, '1.2', ('PAUT',))
        self.assertEqual([b.id for b in sheet.row_breaks.brk], [47])
        self.assertIsNone(manager._first_print_overflow_row(
            sheet, 1, 47, manager._report_height_scale(sheet)))
        self.assertEqual(sheet.row_dimensions[20].height, 2)
        for r in (1, 5, 14, 35, 38, 39, 46, 47):
            self.assertEqual(sheet.row_dimensions[r].height, 20)
        self.assertEqual([sheet.cell(r, 2).value for r in range(39, 47)], sorted(originals))

    def test_final_limit_does_not_copy_empty_mt_header_into_blank_page(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.cell(1, 6, '월간 용역 진도 보고서')
        sheet.cell(5, 6, '3.0 비파괴검사현황')
        sheet.cell(49, 2, '1.2.2 MT')
        sheet.cell(50, 6, 'Line No.')
        sheet.cell(52, 2, '해당 월 검사실적 없음')
        sheet.cell(53, 2, 'TOTAL')
        sheet.row_breaks.append(Break(id=70))
        manager._enforce_final_ndt_body_limit(sheet)
        self.assertEqual([b.id for b in sheet.row_breaks.brk], [70])
        self.assertEqual(sum('Line No.' == sheet.cell(r, 6).value
                             for r in range(1, sheet.max_row + 1)), 1)

    def test_final_limit_moves_whole_table_without_duplicate_heading(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.cell(1, 6, '월간 용역 진도 보고서')
        sheet.cell(5, 6, '3.0 비파괴검사현황')
        sheet.cell(55, 2, '1.2.2 MT')
        sheet.cell(56, 6, 'Line No.')
        sheet.cell(58, 2, '업체')
        sheet.cell(58, 3, 1)
        sheet.cell(59, 2, 'TOTAL')
        sheet.row_breaks.append(Break(id=59))
        manager._enforce_final_ndt_body_limit(sheet)
        self.assertEqual(sum('Line No.' == sheet.cell(r, 6).value
                             for r in range(1, sheet.max_row + 1)), 1)
        self.assertEqual(sheet.cell(65, 3).value, 1)
        self.assertEqual([b.id for b in sheet.row_breaks.brk], [54, 66])

    def test_final_limit_uses_print_height_and_survives_save(self):
        import io
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.cell(1, 6, '월간 용역 진도 보고서')
        sheet.cell(5, 6, '3.0 비파괴검사현황')
        sheet.cell(9, 2, '2.1 PAUT')
        sheet.cell(10, 6, 'Line No.')
        sheet.merge_cells('F10:G11')
        for index in range(58):
            sheet.cell(12 + index, 2, index + 1)
            sheet.row_dimensions[12 + index].height = 9
        sheet.cell(70, 2, 'TOTAL')
        sheet.row_breaks.append(Break(id=70))
        sheet.cell(71, 6, '4.0 용접결함, 불량률 현황')
        manager._enforce_final_ndt_body_limit(sheet)
        buffer = io.BytesIO()
        sheet.parent.save(buffer)
        buffer.seek(0)
        result = openpyxl.load_workbook(buffer).active
        # Short records fit beyond the old fixed 58-row boundary before the
        # continuation is needed.
        self.assertGreater(min(b.id for b in result.row_breaks.brk), 58)
        page_start = 1
        for end in sorted(b.id for b in result.row_breaks.brk):
            self.assertIsNone(manager._first_print_overflow_row(
                result, page_start, end - page_start + 1,
                manager._report_height_scale(result)))
            page_start = end + 1
        records = [result.cell(r, 2).value for r in range(1, result.max_row + 1)
                   if isinstance(result.cell(r, 2).value, int)]
        self.assertEqual(records, list(range(1, 59)))
        self.assertEqual(sum(result.cell(r, 2).value == 'TOTAL'
                             for r in range(1, result.max_row + 1)), 1)
        count = result.max_row
        manager._enforce_final_ndt_body_limit(result)
        self.assertEqual(result.max_row, count)

    def test_defect_table_continuations_do_not_unmerge_next_table_header(self):
        sheet = openpyxl.Workbook().active
        # The middle block is a continuation without its own TOTAL. The old
        # search used the next table's TOTAL and consumed that table's header.
        for header in (700, 715, 729):
            sheet.cell(header + 1, 2, 'ID')
            for col in (16, 17, 18):
                sheet.cell(header, col, f'HEADER {col}')
                sheet.merge_cells(start_row=header, end_row=header + 1,
                                  start_column=col, end_column=col)
            sheet.cell(header + 2, 16, header)
        for total in (704, 733):
            sheet.cell(total, 2, '총계')
            sheet.cell(total, 16, 123)
        manager = MonthlyReportManager('unused.xlsx')
        manager._widen_weld_defect_table(sheet)
        self.assertEqual(sheet.cell(717, 16).value, 715)
        self.assertEqual(sheet.cell(731, 16).value, 729)
        self.assertEqual(sheet.cell(733, 16).value, 123)
        merges = {str(m) for m in sheet.merged_cells.ranges}
        self.assertIn('P729:Q730', merges)
        self.assertIn('P715:Q716', merges)

    def test_summary_compaction_only_reclaims_blank_table_gaps(self):
        manager = MonthlyReportManager('unused.xlsx')
        manager.table_markers.update(qty=10, paut_1=20, mt_1=40,
                                     rt_1=60, pt_1=80, paut_2=90)
        sheet = openpyxl.Workbook().active
        for r in range(1, 91):
            sheet.row_dimensions[r].height = 15
        for header in (20, 40, 60, 80):
            sheet.cell(header, 2, '표 머리글')
            sheet.cell(header + 2, 2, '검사 데이터')
            sheet.cell(header + 3, 2, 'TOTAL')
        sheet.cell(25, 6, '보존할 소제목')
        sheet.merge_cells('B26:C26')
        manager._compact_ndt_summary_page(sheet)
        self.assertLess(sheet.row_dimensions[24].height, 15)
        for r in (20, 22, 23, 25, 26, 80, 82, 83, 90):
            self.assertEqual(sheet.row_dimensions[r].height, 15)
        self.assertEqual(sheet.cell(22, 2).value, '검사 데이터')
        self.assertIn('B26:C26', {str(m) for m in sheet.merged_cells.ranges})

    def test_safety_training_folder_follows_history_data_drive(self):
        manager = MonthlyReportManager('unused.xlsx')
        manager._history_path = (
            r'G:\내 드라이브\07_Antigravity\PMI_Data\central'
            r'\daily_work_history.json'
        )

        with patch.object(os.path, 'isdir', side_effect=lambda path: path.startswith('G:')):
            folder = manager._resolve_safety_training_dir()

        self.assertTrue(folder.startswith('G:'))
        self.assertTrue(folder.endswith('교육일지'))

    def test_detail_continuations_fill_available_height_and_preserve_records(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.page_setup.scale = 100
        sheet.cell(1, 6, '월간 용역 진도 보고서')
        sheet.cell(5, 6, '3.0 비파괴검사현황')
        sheet.cell(30, 2, 'PAUT')
        sheet.cell(31, 2, '순번')
        for r in range(1, 33):
            sheet.row_dimensions[r].height = 12
        for index in range(160):
            sheet.cell(33 + index, 2, index + 1)
            sheet.row_dimensions[33 + index].height = 9
        sheet.cell(193, 2, 'TOTAL')
        sheet.row_breaks.append(Break(id=193))
        start, remaining = 33, 160
        counts = []
        while remaining:
            split = manager._detail_continuation_overflow_row(sheet, start, remaining)
            if split is None:
                counts.append(remaining)
                break
            count = split - start
            self.assertGreater(count, 0)
            counts.append(count)
            inserted = manager._insert_table_continuation(sheet, split, 1, 31)
            sheet.row_breaks.append(Break(id=split - 1))
            remaining -= count
            start = split + inserted
            self.assertLess(len(counts), 10)
        records = [sheet.cell(r, 2).value for r in range(1, sheet.max_row + 1)
                   if isinstance(sheet.cell(r, 2).value, int)]
        self.assertEqual(records, list(range(1, 161)))
        self.assertGreater(counts[1], 47)
        self.assertLess(counts[0], counts[1])
        page_start = 1
        for end in sorted(b.id for b in sheet.row_breaks.brk):
            self.assertIsNone(manager._first_print_overflow_row(
                sheet, page_start, end - page_start + 1,
                manager._report_height_scale(sheet)))
            page_start = end + 1
        for r in range(1, sheet.max_row + 1):
            if isinstance(sheet.cell(r, 2).value, int):
                self.assertEqual(sheet.row_dimensions[r].height, 9)

    def test_print_height_allows_more_than_58_short_rows(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        sheet.page_setup.scale = 100
        for r in range(1, 150):
            sheet.row_dimensions[r].height = 5
        # Page 101:107 is the title, 108:111 is the table heading.
        sheet.row_breaks.append(Break(id=100))
        self.assertIsNone(manager._detail_page_overflow_row(sheet, 112, 46))
        self.assertIsNone(manager._detail_page_overflow_row(sheet, 112, 47))
        self.assertIsNone(manager._detail_page_overflow_row(sheet, 112, 48))

    def test_rewritten_table_discards_template_break_through_total_row(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        for row in (100, 124, 158, 180, 200):
            sheet.row_breaks.append(Break(id=row))

        manager._remove_stale_table_breaks(sheet, 112, 180)

        self.assertEqual([brk.id for brk in sheet.row_breaks.brk], [100, 200])

    def test_total_can_overflow_without_moving_last_record(self):
        manager = MonthlyReportManager('unused.xlsx')
        sheet = openpyxl.Workbook().active
        with patch.object(manager, '_paut_print_scale', return_value=100):
            for count in (1, 4):
                with patch.object(manager, '_first_print_overflow_row', return_value=20 + count):
                    self.assertEqual(manager._detail_page_overflow_row(sheet, 20, count),
                                     20 + count)

    def test_four_added_rows_consume_padding_without_moving_next_header(self):
        sheet = openpyxl.Workbook().active
        for row in range(1, 62):
            sheet.row_dimensions[row].height = 16.5
        sheet.cell(10, 2, 'TOTAL')
        sheet.cell(61, 6, 'NEXT HEADER')
        sheet.row_breaks.append(Break(id=60))
        manager = MonthlyReportManager('unused.xlsx')
        manager.table_markers = {'next': 61}
        manager._insert_rows_safely(sheet, 6, 4)
        manager.table_markers['next'] += 4
        consumed = manager._consume_detail_page_whitespace(sheet, 14, 66)
        self.assertEqual(consumed, 66)
        self.assertEqual(sheet.cell(14, 2).value, 'TOTAL')
        self.assertEqual(sheet.cell(61, 6).value, 'NEXT HEADER')
        self.assertEqual(manager.table_markers['next'], 61)
        self.assertEqual([b.id for b in sheet.row_breaks.brk], [60])

    def test_only_unused_height_is_consumed_and_data_is_preserved(self):
        sheet = openpyxl.Workbook().active
        sheet.cell(10, 2, 'TOTAL')
        sheet.cell(12, 2, 'FOLLOWING TABLE')
        sheet.row_dimensions[13].height = 10
        sheet.row_dimensions[14].height = 20
        sheet.row_breaks.append(Break(id=14))
        manager = MonthlyReportManager('unused.xlsx')
        consumed = manager._consume_detail_page_whitespace(sheet, 10, 16.5)
        self.assertEqual(consumed, 16.5)
        self.assertEqual(sheet.row_dimensions[14].height, 3.5)
        self.assertEqual(sheet.cell(12, 2).value, 'FOLLOWING TABLE')
        self.assertEqual([b.id for b in sheet.row_breaks.brk], [14])
        consumed = manager._consume_detail_page_whitespace(sheet, 10, 100)
        self.assertEqual(consumed, 13.5)
        self.assertEqual(sheet.cell(12, 2).value, 'FOLLOWING TABLE')

    def test_when_page_is_full_new_detail_rows_overflow(self):
        sheet = openpyxl.Workbook().active
        for row in range(1, 70):
            sheet.row_dimensions[row].height = 16.5
        manager = MonthlyReportManager('unused.xlsx')
        self.assertIsNone(manager._first_print_overflow_row(sheet, 10, 30))
        self.assertIsNotNone(manager._first_print_overflow_row(sheet, 10, 60))

    def test_paut_tail_uses_actual_width_fit_instead_of_full_size(self):
        sheet = openpyxl.Workbook().active
        sheet.print_area = 'B1:W60'
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.scale = None
        for dimension in sheet.column_dimensions.values():
            dimension.width = 7
        from openpyxl.utils import get_column_letter
        for col in range(2, 24):
            sheet.column_dimensions[get_column_letter(col)].width = 7
        for row in range(1, 11):
            sheet.row_dimensions[row].height = 16
        for row in range(11, 48):
            sheet.row_dimensions[row].height = 16.5
        manager = MonthlyReportManager('unused.xlsx')
        self.assertIsNotNone(manager._first_print_overflow_row(sheet, 11, 37))
        scale = manager._paut_print_scale(sheet)
        self.assertLess(scale, 100)
        self.assertIsNone(manager._first_print_overflow_row(sheet, 11, 37, scale))

    def test_moves_header_boundary_after_inserted_padding(self):
        sheet = openpyxl.Workbook().active
        for start in (1, 66, 112):
            sheet.cell(start, 6, '월 간 용 역 진 도 보 고 서')
            sheet.cell(start + 4, 6, '3.0 비파괴검사현황')
        sheet.cell(111, 6, 'PAUT 마지막 검사행')
        for end in (65, 111, 176):
            sheet.row_breaks.append(Break(id=end))
        manager = MonthlyReportManager('unused.xlsx')
        manager._pad_ndt_continuation_pages(sheet)
        self.assertEqual(sorted(b.id for b in sheet.row_breaks.brk), [65, 130, 195])
        self.assertEqual(sheet.cell(111, 6).value, 'PAUT 마지막 검사행')
        self.assertEqual(sheet.cell(131, 6).value, '월 간 용 역 진 도 보 고 서')
        self.assertFalse(any(sheet.cell(r, 6).value for r in range(112, 131)))
        manager._align_report_breaks_to_repeated_headers(sheet)
        self.assertNotIn(111, [b.id for b in sheet.row_breaks.brk])
        self.assertIn(130, [b.id for b in sheet.row_breaks.brk])

    def test_pt_spacer_is_compacted_to_keep_document_header(self):
        sheet = openpyxl.Workbook().active
        sheet.cell(20, 2, 'RT TOTAL')
        sheet.cell(23, 2, '2.4 비파괴검사 세부현황(PT)')
        manager = MonthlyReportManager('unused.xlsx')

        manager._compact_ndt_method_spacers(sheet)

        self.assertEqual(sheet.row_dimensions[21].height, 1.0)
        self.assertEqual(sheet.row_dimensions[22].height, 1.0)
        self.assertFalse(sheet.row_dimensions[21].hidden)

    def test_only_weld_defect_trailing_rows_are_compacted(self):
        sheet = openpyxl.Workbook().active
        sheet.cell(300, 6, '4.0 용접결함, 불량률 현황')
        sheet.cell(310, 2, 'TABLE TOTAL')
        sheet.row_breaks.append(Break(id=315))
        manager = MonthlyReportManager('unused.xlsx')

        compacted = manager._compact_weld_defect_trailing_spacers(sheet)

        self.assertEqual(compacted, 5)
        for row in range(311, 316):
            self.assertEqual(sheet.row_dimensions[row].height, 1.0)
            self.assertFalse(sheet.row_dimensions[row].hidden)

    def test_formatted_blank_safety_rows_are_not_bulk_compacted(self):
        sheet = openpyxl.Workbook().active
        sheet.cell(300, 6, '월간 용역 진도 보고서')
        sheet.cell(310, 2, '방사선안전관리 표')
        sheet.cell(316, 6, '월간 용역 진도 보고서')
        manager = MonthlyReportManager('unused.xlsx')

        manager._collapse_blank_rows_before_report_headers(sheet)

        self.assertIsNone(sheet.row_dimensions[311].height)
        self.assertIsNone(sheet.row_dimensions[315].height)

    def test_pt_overflow_gets_repeated_document_header(self):
        sheet = openpyxl.Workbook().active
        sheet.cell(100, 6, '월간 용역 진도 보고서')
        sheet.cell(104, 6, '3.0 비파괴검사현황')
        sheet.cell(120, 6, '2.4 비파괴검사 세부현황(PT)')
        sheet.cell(124, 2, 'TOTAL')
        sheet.row_breaks.append(Break(id=99))
        sheet.row_breaks.append(Break(id=124))
        manager = MonthlyReportManager('unused.xlsx')
        manager.table_markers = {'next': 125}

        with patch.object(manager, '_first_print_overflow_row', return_value=118):
            with patch.object(manager, '_paut_print_scale', return_value=88):
                inserted = manager._ensure_pt_continuation_header(sheet)

        self.assertEqual(inserted, 7)
        self.assertEqual(sheet.cell(120, 6).value, '월간 용역 진도 보고서')
        self.assertEqual(sheet.cell(124, 6).value, '3.0 비파괴검사현황')
        self.assertEqual(sheet.cell(127, 6).value, '2.4 비파괴검사 세부현황(PT)')
        self.assertIn(119, [int(brk.id) for brk in sheet.row_breaks.brk])
        self.assertEqual(manager.table_markers['next'], 132)

if __name__ == '__main__':
    unittest.main()
