import json
import datetime as dt
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from urllib.parse import unquote

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from site_apps.central.src.attribute_export import (
    contained_path, export_workbook, issues, load_history, load_report_data,
    merge_rows, normalize, point_images, restore_connections, write_connections, read_connections,
    create_submission_folders, auto_connect, submission_plan,
)


class AttributeExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.record = normalize({
            '공사명': '=HYPERLINK("evil")', '시공사': '시공사', '용역사': '검사사',
            '라인번호': 'LINE-01', '용접부번호': 'S01', '도면번호': 'DWG-01',
            '검사일자': '2026-10-17', '용접부특성': 'O', '보고서 매수': '2',
            '결과': '합격',
            '보고서번호': 'RT-2026-001',
        }, method='RT')
        self.record['id'] = 'record-1'
        (self.root / 'RT' / '보고서').mkdir(parents=True)
        (self.root / 'RT' / '필름영상화').mkdir()
        for field, path in [('보고서', 'RT/보고서/한글 #1%.pdf'),
                            ('필름영상화', 'RT/필름영상화/영상_1-2.jpg')]:
            (self.root / path).write_bytes(b'fixture')
            self.record[field] = path
        self.record['RT영상'] = {'1-2': [self.record['필름영상화']]}

    def tearDown(self):
        self.temp.cleanup()

    def test_history_and_report_schema(self):
        history = self.root / 'history.json'
        history.write_text(json.dumps({'2026-10-17': {'ndt_results': [
            {'검사방법': 'RT', '라인번호': 'LINE-01', 'Joint No.': 'S01', 'RT_OR': '1'},
            {'검사방법': 'UT', 'Joint No.': 'S02'}]}}), encoding='utf-8')
        rows = load_history(history)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['용접부특성'], 'O')
        connections = self.root / 'connections.json'
        connections.write_text(json.dumps({'version': 3, 'rows': []}), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '연결정보'):
            load_history(connections)
        path = self.root / 'RT_Data.json'
        path.write_text(json.dumps([{'Date': '2026/10/17', 'Joint': 'S01', 'Dwg': 'DWG-01'},
                                    {'Joint': 'S02', 'selected': False}]), encoding='utf-8')
        data = load_report_data(path, 'RT')
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]['도면번호'], 'DWG-01')
        self.assertEqual(data[0]['검사일자'], '2026-10-17')
        self.assertEqual(normalize({'Date': dt.datetime(2026, 10, 17)})['검사일자'], '2026-10-17')

    def test_unique_merge_conflict_and_ambiguity(self):
        report = {**self.record, 'id': 'report-1', '도면번호': 'OTHER', '확인사항': []}
        merged = merge_rows([self.record], [report])
        self.assertEqual(len(merged), 1)
        self.assertEqual(merged[0]['도면번호'], 'DWG-01')
        self.assertTrue(any('불일치' in s for s in merged[0]['확인사항']))
        unmatched = []
        self.assertEqual(len(merge_rows([self.record], [report, {**report, 'id': 'report-2'}], unmatched)), 1)
        self.assertEqual(len(unmatched), 2)
        missing = {**report, '라인번호': ''}
        self.assertEqual(len(merge_rows([self.record], [missing])), 1)
        self.assertEqual(merge_rows([], [report]), [])

    def test_submission_boundaries_and_missing_data(self):
        with self.assertRaises(ValueError):
            contained_path(self.root, '../outside.pdf')
        self.assertEqual(issues(self.record, self.root), [])
        broken = {**self.record, '필름영상화': '', 'RT영상': {}, '검사일자': 'wrong'}
        self.assertTrue(any('포인트 영상 미연결' in s for s in issues(broken, self.root)))
        self.assertIn('검사일자 형식 확인', issues(broken, self.root))
        wrong_method = {**self.record, '보고서': 'PT/report.pdf'}
        (self.root / 'PT').mkdir()
        (self.root / 'PT' / 'report.pdf').write_bytes(b'fixture')
        self.assertIn('보고서 검사종류 폴더 불일치', issues(wrong_method, self.root))

    def test_export_links_survive_move_and_text_is_literal(self):
        from openpyxl import load_workbook
        destination = self.root / '중앙지사_속성정보.xlsx'
        export_workbook([self.record], self.root, destination)
        moved = self.root / 'relocated'
        moved.mkdir()
        shutil.copytree(self.root / 'RT', moved / 'RT')
        shutil.copy(destination, moved / destination.name)
        book = load_workbook(moved / destination.name)
        try:
            ws = book['속성정보']
            self.assertEqual(ws['A2'].data_type, 's')
            self.assertEqual(ws['A2'].value, self.record['공사명'])
            self.assertEqual(ws['L2'].value, 2)
            self.assertEqual(ws['I2'].value.strftime('%Y-%m-%d'), '2026-10-17')
            for cell in ('M2', 'N2'):
                target = ws[cell].hyperlink.target
                self.assertNotIn('file:', target)
                self.assertTrue((moved / unquote(target)).is_file())
            self.assertIn('%23', ws['M2'].hyperlink.target)
            self.assertIn('%25', ws['M2'].hyperlink.target)
            self.assertIn('보고서/한글 ', ws['M2'].hyperlink.target)
            self.assertNotIn('%EB', ws['M2'].hyperlink.target)
            self.assertEqual(ws['M2'].value, 'RT-2026-001')
            self.assertEqual(book['검사내역']['H2'].value, 'RT-2026-001')
            self.assertEqual(book['검사내역']['H2'].hyperlink.target, ws['M2'].hyperlink.target)
            self.assertEqual(book['제출파일목록'].max_row, 3)
            self.assertEqual(ws['P2'].value, '입력 및 연결 확인 완료')
        finally:
            book.close()
        with self.assertRaises(FileExistsError):
            export_workbook([self.record], self.root, destination)

    def test_rt_points_have_individual_file_links_and_duplicate_exposures(self):
        from openpyxl import load_workbook
        folder = self.root / 'RT' / '필름영상화' / 'LINE-01' / 'S01'
        folder.mkdir(parents=True)
        for filename in ('S01_1-2.tif', 'S01_2-3.jpg', 'S01_3-1.png', 'S01_1-2_재촬영.jpg'):
            (folder / filename).write_bytes(b'fixture')
        # A different weld's images must never leak into this record.
        (folder.parent / 'S02_2-3.jpg').write_bytes(b'other weld')
        row = {**self.record, 'RT영상': {}, '필름영상화': folder.relative_to(self.root).as_posix()}
        mapping = point_images(row, self.root)
        self.assertEqual(set(mapping), {'1-2', '2-3', '3-1'})
        self.assertEqual(len(mapping['1-2']), 2)
        output = self.root / 'points.xlsx'
        export_workbook([row], self.root, output)
        book = load_workbook(output)
        try:
            self.assertEqual(book.sheetnames[0], '검사내역')
            ws = book['검사내역']
            headers = {cell.value: cell.column for cell in ws[1]}
            for label in ('RT 1-2', 'RT 1-2 (2)', 'RT 2-3', 'RT 3-1'):
                cell = ws.cell(2, headers[label])
                target = self.root / unquote(cell.hyperlink.target)
                self.assertTrue(target.is_file())
                self.assertEqual(target.parent, folder)
            self.assertEqual(ws.cell(2, headers['검사결과']).value, '합격')
        finally:
            book.close()

    def test_explicit_point_mapping_and_missing_image(self):
        row = {**self.record, '필름영상화': '', 'RT영상': {'1-2': [self.record['필름영상화']], '2-3': ['RT/missing.tif']}}
        self.assertEqual(len(point_images(row, self.root)), 2)
        self.assertTrue(any('2-3 영상 파일 없음' in s for s in issues(row, self.root)))
        with self.assertRaises(ValueError):
            point_images({**row, 'RT영상': {'bad': ['RT/image.tif']}}, self.root)

    def test_report_number_is_not_the_preview_row_number(self):
        self.assertEqual(normalize({'No': 123})['보고서번호'], '')
        self.assertEqual(normalize({'Report No.': 'PT-001'})['보고서번호'], 'PT-001')
        self.assertIn('보고서번호 미입력', issues({**self.record, '보고서번호': ''}, self.root))

    def test_connections_restore_by_source_identity_not_position(self):
        source = {**self.record, '보고서': '', '보고서번호': '', 'RT영상': {}}
        source['원본값'] = {key: source.get(key, '') for key in
                          ('검사일자', '보고서종류', '라인번호', '용접부번호', '용접부특성', '검사결과', '보고서', '보고서번호')}
        saved = {**source, '보고서': self.record['보고서'], '보고서번호': 'RT-SAVED'}
        saved['id'] = 'old-file-and-position'
        changed = {**source, 'id': 'new-position', '검사결과': '불합격'}
        changed['원본값'] = {**source['원본값'], '검사결과': '불합격'}
        restored = restore_connections([changed], [saved, {**saved, '라인번호': 'OTHER', '원본값': {**source['원본값'], '라인번호': 'OTHER'}}])
        self.assertEqual(len(restored), 1)
        self.assertEqual(restored[0]['보고서번호'], 'RT-SAVED')
        self.assertEqual(restored[0]['보고서'], self.record['보고서'])
        self.assertEqual(restored[0]['검사결과'], '불합격')
        self.assertTrue(any('변경' in value for value in restored[0]['확인사항']))
        duplicated = restore_connections([source, {**source, 'id': 'duplicate'}], [saved])
        self.assertTrue(all(not row['보고서'] for row in duplicated))

    def test_connections_roundtrip_and_moved_submission(self):
        source = {**self.record, '보고서': '', '보고서번호': ''}
        source['원본값'] = dict(source)
        saved = {**source, '보고서': str(self.root / self.record['보고서']), '보고서번호': 'RT-RESTORED',
                 'RT영상': {'1-2': [self.record['필름영상화']]}}
        cache = self.root / '.connections.json'
        write_connections(cache, [saved], self.root)
        loaded = read_connections(cache)
        self.assertEqual(loaded[0]['보고서'], self.record['보고서'])
        self.assertEqual(restore_connections([source], loaded)[0]['보고서번호'], 'RT-RESTORED')
        self.assertEqual(len(restore_connections([], loaded)), 0)

    def test_central_points_show_missing_and_unmatched_reports_stay_separate(self):
        from openpyxl import load_workbook
        row = {**self.record, '촬영포인트': ['1-2', '2-3', '3-1']}
        self.assertEqual(normalize({'구간정보': '1-2, 2-3,3-1'}, method='RT')['촬영포인트'], row['촬영포인트'])
        self.assertEqual(normalize({'구간정보': 'unknown'}, method='RT')['촬영포인트'], [])
        unmatched = []
        records = merge_rows([row], [{**row, 'id': 'unmatched', '라인번호': 'OTHER'}], unmatched)
        output = self.root / 'missing-points.xlsx'
        export_workbook(records, self.root, output, unmatched)
        book = load_workbook(output)
        try:
            ws = book['검사내역']
            columns = {cell.value: cell.column for cell in ws[1]}
            self.assertEqual(ws.max_row, 2)
            self.assertEqual(ws.cell(2, columns['RT 3-1']).value, '미연결')
            self.assertIsNone(ws.cell(2, columns['RT 3-1']).hyperlink)
            self.assertEqual(ws.cell(2, columns['RT 영상 연결현황']).value, '1/3 포인트 연결')
            self.assertEqual(book['보고서 대조확인'].max_row, 2)
        finally:
            book.close()

    def test_shared_rt_folder_cannot_mix_other_welds(self):
        for joint, point in [('S01', '1-2'), ('S02', '2-3')]:
            folder = self.root / 'RT' / joint
            folder.mkdir()
            (folder / f'{point}.jpg').write_bytes(b'image')
        for directory in ('RT', 'RT/S01'):
            with self.assertRaises(ValueError):
                point_images({**self.record, 'RT영상': {}, '필름영상화': directory}, self.root)
        # A nested sibling or repeat-inspection folder is not swept into this record.
        own = self.root / 'RT' / 'LINE-01' / 'S01'
        (own / 'S02').mkdir(parents=True)
        (own / '1-2.jpg').write_bytes(b'own')
        (own / 'S02' / '2-3.jpg').write_bytes(b'other')
        mapping = point_images({**self.record, 'RT영상': {}, '필름영상화': own.relative_to(self.root).as_posix()}, self.root)
        self.assertEqual(set(mapping), {'1-2'})

    def test_conflict_warning_survives_save_restore_and_export(self):
        from openpyxl import load_workbook
        source = dict(self.record)
        source['원본값'] = dict(source)
        report = {**source, '검사결과': '불합격'}
        merged = merge_rows([source], [report])
        cache = self.root / 'warnings.json'
        write_connections(cache, merged, self.root)
        restored = restore_connections([source], read_connections(cache))
        self.assertTrue(any('검사결과 불일치' in value for value in restored[0]['확인사항']))
        output = self.root / 'warnings.xlsx'
        export_workbook(restored, self.root, output)
        book = load_workbook(output)
        try:
            self.assertIn('검사결과 불일치', book['속성정보']['P2'].value)
        finally:
            book.close()

    def test_duplicate_attachments_survive_repeated_autosaves(self):
        from openpyxl import load_workbook
        sources = []
        saved = []
        for index in range(2):
            source = {**self.record, 'id': str(index), '보고서번호': '', '보고서': '', 'RT영상': {}}
            source['원본값'] = dict(source)
            sources.append(source)
            saved.append({**source, '보고서번호': f'RT-SAVED-{index}'})
        cache = self.root / 'duplicates.json'
        write_connections(cache, saved, self.root)
        for _ in range(3):
            reopened = restore_connections(sources, read_connections(cache))
            self.assertTrue(all(not row['보고서번호'] for row in reopened))
            write_connections(cache, reopened, self.root)
        payload = json.loads(cache.read_text(encoding='utf-8'))
        self.assertEqual({row['보고서번호'] for row in payload['unresolved']}, {'RT-SAVED-0', 'RT-SAVED-1'})
        self.assertEqual(len(payload['unresolved']), 2)
        output = self.root / 'pending.xlsx'
        export_workbook(sources, self.root, output, pending=payload['unresolved'])
        book = load_workbook(output)
        try:
            self.assertEqual(book['연결 복원보류'].max_row, 3)
            self.assertEqual({book['연결 복원보류'].cell(i, 6).value for i in (2, 3)}, {'RT-SAVED-0', 'RT-SAVED-1'})
        finally:
            book.close()

    def test_zero_rt_quantities_do_not_define_weld_characteristic(self):
        for value in (0, '0', 0.0, '0.0', '', None, -1, 'invalid'):
            row = normalize({'RT_OR': value, 'RT_RE': value}, method='RT')
            self.assertEqual(row['용접부특성'], '')
            self.assertFalse(any('동시' in warning for warning in row['확인사항']))
        self.assertEqual(normalize({'RT_OR': '1', 'RT_RE': '0'}, method='RT')['용접부특성'], 'O')
        self.assertEqual(normalize({'RT_OR': 0, 'RT_RE': 2}, method='RT')['용접부특성'], '')
        self.assertEqual(normalize({'RT_OR': 3, 'RT_RE': 1}, method='RT')['용접부특성'], 'O')
        self.assertTrue(any('불량매수' in value for value in normalize({'RT_OR': 0, 'RT_RE': 2}, method='RT')['확인사항']))

    def test_generated_tree_auto_connect_and_rescan(self):
        row = {**self.record, '보고서': '', '보고서번호': '', '필름영상화': '', 'RT영상': {}}
        self.assertEqual(create_submission_folders([row], self.root), 1)
        relative = submission_plan([row])[0][1]
        directory = self.root / relative
        self.assertEqual(relative.as_posix(), 'RT/LINE-01/S01/2026-10-17_최초')
        report = directory / '보고서' / 'RT-NEW.pdf'
        report.write_bytes(b'report')
        (directory / '필름영상화' / '1-2.tif').write_bytes(b'image')
        linked = auto_connect([row], self.root)[0]
        self.assertEqual(linked['보고서번호'], 'RT-NEW')
        self.assertEqual(set(linked['RT영상']), {'1-2'})
        (directory / '필름영상화' / '2-3.tif').write_bytes(b'image')
        self.assertEqual(set(auto_connect([linked], self.root)[0]['RT영상']), {'1-2', '2-3'})
        create_submission_folders([row], self.root)
        self.assertEqual(report.read_bytes(), b'report')

    def test_paut_and_shared_reports_and_ambiguous_candidates(self):
        rt = {**self.record, '보고서': '', '보고서번호': '', 'RT영상': {}, '필름영상화': ''}
        paut = {**rt, 'id': 'paut', '보고서종류': 'PAUT', '보고서번호': 'SHARED-001'}
        create_submission_folders([rt, paut], self.root)
        (self.root / 'PAUT' / '보고서' / 'SHARED-001.pdf').write_bytes(b'shared')
        directory = self.root / submission_plan([paut])[0][1]
        (directory / '데이터' / 'scan.dat').write_bytes(b'data')
        linked = auto_connect([paut], self.root)[0]
        self.assertEqual(linked['보고서'], 'PAUT/보고서/SHARED-001.pdf')
        self.assertTrue(linked['데이터'].endswith('/데이터'))
        rt_directory = self.root / submission_plan([rt])[0][1]
        for name in ('A.pdf', 'B.pdf'):
            (rt_directory / '보고서' / name).write_bytes(b'candidate')
        linked = auto_connect([rt], self.root)[0]
        self.assertFalse(linked['보고서'])
        self.assertTrue(any('여러' in s for s in linked['확인사항']))

    def test_folder_collision_and_traversal_protection(self):
        first = {**self.record, '라인번호': '../line', 'id': 'first'}
        second = {**first, 'id': 'second'}
        plans = submission_plan([first, second])
        self.assertNotEqual(plans[0][1], plans[1][1])
        create_submission_folders([first, second], self.root)
        self.assertTrue(all((self.root / path).resolve().is_relative_to(self.root) for _, path in plans))


if __name__ == '__main__':
    unittest.main()
