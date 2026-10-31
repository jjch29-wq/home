"""Check photo/caption boundaries after workbook serialization."""
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from PIL import Image
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor
from openpyxl.utils.units import pixels_to_EMU

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from site_apps.central.src.services.monthly_report_manager import MonthlyReportManager


class PhotoLayoutTest(unittest.TestCase):
    def test_enlarged_logo_centers_artwork_without_white_asset_padding(self):
        from io import BytesIO
        from PIL import ImageDraw
        from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        sheet = openpyxl.Workbook().active
        sheet.merge_cells('F10:O13')
        sheet.cell(10, 6, '월간 용역 진도 보고서')
        artwork = Image.new('RGB', (120, 30), 'white')
        ImageDraw.Draw(artwork).rectangle((10, 5, 49, 24), fill='red')
        buffer = BytesIO()
        artwork.save(buffer, format='PNG')
        buffer.seek(0)
        picture = openpyxl.drawing.image.Image(buffer)
        picture.anchor = OneCellAnchor(_from=AnchorMarker(row=10, col=1),
            ext=XDRPositiveSize2D(pixels_to_EMU(120), pixels_to_EMU(30)))
        sheet.add_image(picture)
        manager = MonthlyReportManager('unused.xlsx')
        manager._center_enlarge_report_logos(sheet)
        result = sheet._images[0]
        self.assertEqual(result.anchor.ext.cx, pixels_to_EMU(60))
        self.assertEqual(result.anchor.ext.cy, pixels_to_EMU(30))
        self.assertEqual((result.width, result.height), (40, 20))
        self.assertGreater(result.anchor._from.colOff, 0)

    def test_other_pages_get_top_spacer_without_duplicating_ndt(self):
        from openpyxl.worksheet.pagebreak import Break
        sheet = openpyxl.Workbook().active
        for start, title in ((10, '목차'), (40, '3.0 비파괴검사현황'),
                             (80, '6.0 사진대지')):
            sheet.cell(start, 6, '월간 용역 진도 보고서')
            sheet.cell(start + 4, 6, title)
            sheet.merge_cells(start_row=start, start_column=6,
                              end_row=start + 3, end_column=15)
            sheet.row_breaks.append(Break(id=start - 1))
        image = openpyxl.drawing.image.Image(Image.new('RGB', (120, 30)))
        image.anchor = TwoCellAnchor()
        image.anchor._from.row = 79
        image.anchor.to.row = 82
        sheet.add_image(image)
        manager = MonthlyReportManager('unused.xlsx')
        manager.table_markers = {'photo': 89}
        self.assertEqual(manager._insert_other_report_top_spacers(sheet), 2)
        self.assertEqual(sheet.row_dimensions[10].height, 20)
        self.assertEqual(sheet.row_dimensions[81].height, 20)
        self.assertEqual(sheet.cell(41, 6).value, '월간 용역 진도 보고서')
        self.assertEqual(sheet.cell(45, 6).value, '3.0 비파괴검사현황')
        self.assertEqual([b.id for b in sheet.row_breaks.brk], [9, 40, 80])
        self.assertEqual(image.anchor._from.row, 81)
        self.assertEqual(image.anchor.to.row, 84)
        self.assertIn('F82:O85', {str(m) for m in sheet.merged_cells.ranges})
        self.assertEqual(manager.table_markers['photo'], 91)

    def test_photo_caption_gap_preserves_grid_position(self):
        sheet = openpyxl.Workbook().active
        sheet.cell(8, 2, '1. 사진대지')
        sheet.row_dimensions[8].height = 12
        sheet.row_dimensions[9].height = 15
        sheet.cell(10, 3, '1.1 위상배열초음파탐상검사(SEC.16)')
        manager = MonthlyReportManager('unused.xlsx')
        manager._space_photo_page_caption(sheet)
        self.assertEqual(sheet.row_dimensions[8].height, 17)
        self.assertEqual(sheet.row_dimensions[9].height, 10)
        self.assertEqual(sheet.cell(8, 2).alignment.vertical, 'bottom')
        self.assertEqual(sheet.cell(10, 3).value,
                         '1.1 위상배열초음파탐상검사(SEC.16)')

    def test_cumulative_logos_use_toc_anchor_and_size(self):
        from io import BytesIO
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for start in (20, 100, 350):
            sheet.cell(start, 6, '월간 용역 진도 보고서')
            sheet.merge_cells(start_row=start, start_column=6,
                              end_row=start + 3, end_column=15)
        sheet.cell(2, 6, '월간 용역 진도 보고서')
        sheet.cell(24, 6, '목 차')
        sheet.row_dimensions[20].height = 10
        sheet.row_dimensions[100].height = 25
        sheet.row_dimensions[350].height = 7
        for row, offset in ((20, 12000), (100, 90000), (350, 70000)):
            buffer = BytesIO()
            Image.new('RGB', (120, 30), 'red').save(buffer, format='PNG')
            buffer.seek(0)
            picture = openpyxl.drawing.image.Image(buffer)
            picture.anchor = TwoCellAnchor()
            picture.anchor._from.row = row
            picture.anchor._from.col = 1
            picture.anchor._from.rowOff = offset
            picture.anchor.to.row = row + 2
            picture.anchor.to.col = 4
            sheet.add_image(picture)
        manager = MonthlyReportManager('unused.xlsx')
        manager._normalize_all_report_header_logos(sheet, reference_toc=True)
        self.assertEqual(len(sheet._images), 3)
        expected_size = (sheet._images[0].anchor.ext.cx,
                         sheet._images[0].anchor.ext.cy)
        for picture, row in zip(sheet._images, (20, 100, 350)):
            marker = picture.anchor._from
            inset = marker.rowOff
            for r in range(row, marker.row + 1):
                inset += int(float(sheet.row_dimensions[r].height or 15) * 12700)
            self.assertEqual(inset, 12000 + 10 * 12700)
            self.assertLess(marker.rowOff, int(float(
                sheet.row_dimensions[marker.row + 1].height or 15) * 12700))
            self.assertEqual((picture.anchor.ext.cx, picture.anchor.ext.cy),
                             expected_size)
            self.assertEqual((picture.width, picture.height), (120, 30))
        manager._normalize_all_report_header_logos(sheet, reference_toc=True)
        self.assertEqual(sheet._images[0].anchor._from.rowOff, 12000)

    def test_monthly_logos_keep_early_page_positions(self):
        sheet = openpyxl.Workbook().active
        sheet.cell(20, 6, '월간 용역 진도 보고서')
        sheet.cell(24, 6, '목차')
        manager = MonthlyReportManager('unused.xlsx')
        manager._normalize_all_report_header_logos(sheet)
        self.assertEqual(sheet.cell(24, 6).value, '목차')

    def test_fixed_ndt_top_spacer_uses_twenty_points(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.cell(100, 2, '월간 용역 진도 보고서')
        sheet.cell(104, 6, '3.0 비파괴검사현황')
        manager = MonthlyReportManager('unused.xlsx')

        inserted = manager._insert_fixed_ndt_top_spacers(sheet)

        self.assertEqual(inserted, 1)
        self.assertEqual(sheet.row_dimensions[100].height, 20.0)
        self.assertIsNone(sheet.cell(100, 2).value)
        self.assertEqual(sheet.cell(101, 2).value, '월간 용역 진도 보고서')

    def test_top_rows_preserve_page_boundaries_and_reserved_height(self):
        from openpyxl.worksheet.pagebreak import Break
        sheet = openpyxl.Workbook().active
        manager = MonthlyReportManager('unused.xlsx')
        for start, section in ((20, '3.0'), (50, '3.0'), (80, '4.0')):
            sheet.cell(start, 6, '월간 용역 진도 보고서')
            sheet.cell(start + 4, 6, section + ' 비파괴검사현황')
            sheet.row_dimensions[start + 6].height = 45
            sheet.row_breaks.append(Break(id=start - 1))
            sheet.merge_cells(start_row=start, start_column=6,
                              end_row=start + 3, end_column=15)
        manager.table_markers = {'detail': 58, 'other': 84}
        self.assertEqual(manager._insert_fixed_ndt_top_spacers(
            sheet, reserved_height=20), 2)
        self.assertEqual([b.id for b in sheet.row_breaks.brk], [19, 50, 81])
        for spacer, header in ((20, 21), (51, 52)):
            self.assertIsNone(sheet.cell(spacer, 6).value)
            self.assertEqual(sheet.row_dimensions[spacer].height, 20)
            self.assertEqual(sheet.row_dimensions[header + 6].height, 25)
            self.assertEqual(sheet.cell(header, 6).value, '월간 용역 진도 보고서')
            self.assertIn(f'F{header}:O{header + 3}',
                          {str(m) for m in sheet.merged_cells.ranges})
        self.assertEqual(sheet.row_dimensions[88].height, 45)
        self.assertEqual(manager.table_markers, {'detail': 60, 'other': 86})

    def test_sec16_marker_is_retained_after_photo_layout(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source = folder / 'photo.png'
            Image.new('RGB', (320, 240), 'steelblue').save(source)
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.cell(610, 3, '1.1 위상배열초음파탐상검사(SEC.16)')
            for row in range(600, 686):
                sheet.row_dimensions[row].height = 2.7
            sheet.print_area = 'B600:W685'
            manager = MonthlyReportManager('unused.xlsx')
            manager._history_path = str(folder / 'history.json')

            manager._populate_process_photo_pages(sheet, [dict(
                process='PAUT', file_path=str(source), date='2026-10-16',
                joint_no='S01', description='검사 공정사진', location='Sec.03',
            )])

            self.assertEqual(len(sheet._images), 1)
            self.assertIn('SEC.16', str(sheet.cell(610, 3).value).upper())

    def test_both_corners_stay_inside_frame_above_caption(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source = folder / 'photo.png'
            Image.new('RGB', (480, 640), 'steelblue').save(source)
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.cell(610, 3, '1.1 위상배열초음파탐상검사(SEC.16)')
            for row in range(600, 686):
                sheet.row_dimensions[row].height = 2.7
            sheet.print_area = 'B600:W685'
            manager = MonthlyReportManager('unused.xlsx')
            manager._history_path = str(folder / 'history.json')
            photos = [dict(process='PAUT', file_path=str(source),
                           date='2026-10-16', joint_no=f'S{i:02}',
                           description='검사 공정사진', location='Sec.16')
                      for i in range(18)]
            manager._populate_process_photo_pages(sheet, photos)
            output = folder / 'report.xlsx'
            workbook.save(output)
            checked = openpyxl.load_workbook(output).active
            self.assertEqual(len(checked._images), 18)
            self.assertIn('SEC.16', str(checked.cell(610, 3).value).upper())
            inset = pixels_to_EMU(6)
            for i, picture in enumerate(checked._images):
                anchor = picture.anchor
                self.assertIsInstance(anchor, TwoCellAnchor)
                self.assertEqual(anchor._from.colOff, inset)
                self.assertEqual(anchor._from.rowOff, inset)
                self.assertEqual(anchor.to.colOff, -inset)
                self.assertEqual(anchor.to.rowOff, -inset)
                self.assertEqual(anchor.to.row - anchor._from.row, 16)
                caption_row = anchor.to.row + 1
                caption = checked.cell(caption_row, anchor._from.col + 1)
                self.assertIn('검사 공정사진', caption.value)
                self.assertFalse(caption.alignment.wrap_text)
                self.assertTrue(caption.alignment.shrink_to_fit)
                self.assertEqual(checked.row_dimensions[caption_row].height, 30)
                if i % 2:
                    left = checked._images[i - 1].anchor
                    self.assertEqual(left._from.row, anchor._from.row)
                    self.assertEqual(left.to.row, anchor.to.row)
                    self.assertEqual(left._from.rowOff, anchor._from.rowOff)
                    self.assertEqual(left.to.rowOff, anchor.to.rowOff)


if __name__ == '__main__':
    unittest.main()
