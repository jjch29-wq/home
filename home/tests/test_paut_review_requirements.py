import importlib.util
from pathlib import Path
import re
import sys
import tempfile
import unittest

import openpyxl

SOURCE = Path(__file__).resolve().parents[1] / 'src' / 'PAUT_검사설정표_생성기.py'
spec = importlib.util.spec_from_file_location('paut_review_generator', SOURCE)
app = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = app
spec.loader.exec_module(app)


class ReviewRequirementsTests(unittest.TestCase):
    def review_values(self, book):
        self.assertEqual(book.sheetnames, ['PAUT Setup'])
        sheet = book['PAUT Setup']
        self.assertEqual(sheet.page_setup.fitToHeight, 1)
        return {sheet.cell(row, label).value: sheet.cell(row, value).value
                for row in range(1, sheet.max_row + 1)
                for label, value in ((1, 4), (7, 9))}

    def test_five_thicknesses_keep_frequency_and_required_width(self):
        with tempfile.TemporaryDirectory() as folder:
            for thickness, width in [(15.87, 15.87), (22.23, 22.23), (25.48, 25),
                                     (28.55, 25), (42.88, 25), (200, 25), (201, 50)]:
                with self.subTest(thickness=thickness):
                    settings = app.ScanSettings(probe_model='2.25L64-A2', frequency_mhz=2.25,
                                                thickness_mm=thickness)
                    path = Path(folder) / f'{thickness}.xlsx'
                    app.create_workbook(settings, str(path))
                    book = openpyxl.load_workbook(path)
                    try:
                        self.assertEqual(book['PAUT Setup']['B14'].value, 2.25)
                        self.assertEqual(book['PAUT Setup']['B14'].number_format, '0.00')
                        review = self.review_values(book)
                        self.assertTrue(review['Required extension on EACH side'].startswith(f'{width:g} mm'))
                        haz = review['Actual HAZ width (reduced-volume alternative only)']
                        self.assertIn('NOT ESTABLISHED', haz)
                        self.assertIn('based on WT', haz)
                        self.assertNotIn('5 mm HAZ', haz)
                        self.assertIn('NOT DEMONSTRATED', review['Coverage verification'])
                        self.assertEqual(len(book['PAUT Setup']._images), 1)
                    finally:
                        book.close()

    def test_old_settings_accept_new_fields_and_roundtrip(self):
        settings = app.ScanSettings.from_dict({'probe_model': '5L64-A2'})
        self.assertEqual(settings.pulser_voltage, '')
        settings.pulser_voltage = '80 V; group A1 only; setup R1'
        restored = app.ScanSettings.from_dict(app.asdict(settings))
        self.assertEqual(restored.pulser_voltage, settings.pulser_voltage)

    def test_wrong_nominal_frequency_is_rejected(self):
        with self.assertRaisesRegex(ValueError, '2.25 MHz'):
            app.ScanSettings(probe_model='2.25L64-A2', frequency_mhz=2).validate()

    def test_photo_presets_use_verified_x3_gain_specifications(self):
        self.assertEqual(
            app.COMMON_PHOTO_SETTINGS['dynamic_range'],
            '76 dB (OmniScan X3 instrument specification)',
        )
        gain = app.COMMON_PHOTO_SETTINGS['digital_gain']
        self.assertIn('Receiver gain 20.0 dB', gain)
        self.assertIn('Combined analog and digital amplification', gain)
        self.assertIn('Separate digital gain: N/A', gain)

    def test_photo_presets_switch_probe_model_with_thickness(self):
        for thickness in (15.87, 22.23, 25.48, 28.55):
            with self.subTest(thickness=thickness):
                preset = app.THICKNESS_PRESETS[thickness]
                self.assertEqual(preset['probe_model'], '5L64-A2')
                self.assertEqual(preset['frequency_mhz'], '5.0')
        thick_preset = app.THICKNESS_PRESETS[42.88]
        self.assertEqual(thick_preset['probe_model'], '2.25L64-A2')
        self.assertEqual(thick_preset['frequency_mhz'], '2.25')

    def test_22_23_preset_uses_band_pass_filter(self):
        self.assertTrue(
            app.THICKNESS_PRESETS[22.23]['filters_smoothing'].startswith('BP 5.25 MHz')
        )

    def test_photo_setup_applicability_uses_current_thickness_not_old_pdf(self):
        with tempfile.TemporaryDirectory() as folder:
            settings = app.ScanSettings(
                thickness_mm=15.87,
                instrument_identification='OmniScan X3 64 / MXU 5.18.1 / 설정 SL15.87_I15.oset',
            )
            for key, value in app.COMMON_PHOTO_SETTINGS.items():
                setattr(settings, key, value)
            for key, value in app.THICKNESS_PRESETS[15.87].items():
                if key not in {'setup_id', 'probe_model', 'frequency_mhz'}:
                    setattr(settings, key, value)
            path = Path(folder) / 'photo-15.87.xlsx'
            app.create_workbook(settings, str(path))
            book = openpyxl.load_workbook(path)
            try:
                review = self.review_values(book)
                self.assertIn('15.87 mm', review['Applicability'])
                self.assertNotIn('6.40 mm', review['Applicability'])
                self.assertIn('PHOTO-DOCUMENTED', review['Review status'])
                self.assertIn('screen photos', review['Setting evidence / approval'])
            finally:
                book.close()

    def test_print_fonts_are_large_enough_to_read(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'readable.xlsx'
            app.create_workbook(app.ScanSettings(), str(path))
            book = openpyxl.load_workbook(path)
            try:
                sheet = book['PAUT Setup']
                self.assertEqual(sheet['A1'].font.sz, 22)
                self.assertEqual(sheet['A4'].font.sz, 15)
                self.assertEqual(sheet['A5'].font.sz, 12)
                self.assertEqual(sheet['A8'].font.sz, 11)
                self.assertEqual(sheet.page_margins.left, 0.10)
                self.assertEqual(sheet.page_margins.right, 0.10)
                instrument_label = next(
                    cell for row in sheet.iter_rows() for cell in row
                    if cell.value == '1. Net digitizing frequency / points / compression'
                )
                self.assertEqual(instrument_label.font.sz, 12)
            finally:
                book.close()

    def test_photo_setup_workbook_contains_no_korean_text(self):
        with tempfile.TemporaryDirectory() as folder:
            preset = app.THICKNESS_PRESETS[15.87]
            settings = app.ScanSettings(
                thickness_mm=15.87,
                probe_model=preset['probe_model'],
                frequency_mhz=float(preset['frequency_mhz']),
                instrument_identification=f"OmniScan X3 64 / MXU 5.18.1 / Setup {preset['setup_id']}",
            )
            for key, value in app.COMMON_PHOTO_SETTINGS.items():
                setattr(settings, key, value)
            for key, value in preset.items():
                if key not in {'setup_id', 'probe_model', 'frequency_mhz'} and hasattr(settings, key):
                    setattr(settings, key, value)
            path = Path(folder) / 'english-only.xlsx'
            app.create_workbook(settings, str(path))
            book = openpyxl.load_workbook(path, data_only=False)
            try:
                korean_cells = [
                    (cell.coordinate, cell.value)
                    for sheet in book.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                    if isinstance(cell.value, str) and re.search(r'[가-힣]', cell.value)
                ]
                self.assertEqual(korean_cells, [])
            finally:
                book.close()

    def test_legacy_loaded_korean_values_are_translated_on_export(self):
        with tempfile.TemporaryDirectory() as folder:
            settings = app.ScanSettings(
                thickness_mm=15.87,
                instrument_identification='OmniScan X3 64 / 설정 SL15.87_I15.oset',
                prf='Acquisition rate 120.0 Hz / PRF 미확정',
            )
            path = Path(folder) / 'legacy-loaded.xlsx'
            app.create_workbook(settings, str(path))
            book = openpyxl.load_workbook(path, data_only=False)
            try:
                values = [
                    cell.value
                    for sheet in book.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                    if isinstance(cell.value, str)
                ]
                self.assertFalse(any(re.search(r'[가-힣]', value) for value in values))
                self.assertTrue(any('PRF not confirmed' in value for value in values))
                self.assertTrue(any('Setup SL15.87_I15.oset' in value for value in values))
            finally:
                book.close()

    def test_required_examination_extension_follows_rlng_section_2_2(self):
        expected = {
            15.87: 15.87,
            22.23: 22.23,
            25.48: 25.0,
            28.55: 25.0,
            42.88: 25.0,
            200.0: 25.0,
            201.0: 50.0,
        }
        for thickness, extension in expected.items():
            with self.subTest(thickness=thickness):
                self.assertEqual(
                    app.required_examination_extension_mm(thickness), extension
                )

    def test_filled_settings_preserve_values_without_claiming_approval(self):
        with tempfile.TemporaryDirectory() as folder:
            settings = app.ScanSettings()
            for key, _ in app.INSTRUMENT_FIELDS:
                setattr(settings, key, 'Example setting - not for examination')
            path = Path(folder) / 'filled.xlsx'
            app.create_workbook(settings, str(path))
            book = openpyxl.load_workbook(path)
            try:
                review = self.review_values(book)
                self.assertIn('SUBJECT TO PROCEDURE APPROVAL', review['Review status'])
                self.assertEqual(review['1. Net digitizing frequency / points / compression'], settings.net_digitizing)
            finally:
                book.close()


if __name__ == '__main__':
    unittest.main()
